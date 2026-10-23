"""Unified resource-center operations over canonical artifact objects."""

from __future__ import annotations

import difflib
import hashlib
import html
import io
import json
import re
from pathlib import Path
from typing import Any, Literal

from pydantic import Field, model_validator
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.database import (
    DBArtifact,
    DBArtifactVersion,
    DBAssessmentAttempt,
    DBAssessmentGradingResult,
    DBAssessmentQuestion,
    DBAssessmentQuestionVersion,
    DBCourse,
    DBGenerationTask,
    DBKnowledgePoint,
    DBLearningEvent,
    DBSourceDocument,
    DBSourceRef,
    DBTaskStep,
    DBMisconceptionReview,
)
from app.m1_lifecycle import record_audit, set_artifact_archived
from app.m1_repository import M1Repository, add_artifact_version, learner_safe_artifact_content
from app.m2_contracts import StrictModel
from app.m5_capabilities import M5CapabilityError
from app.m5_assessment_core import (
    grading_result_payload,
    learner_attempt_payload,
    learner_correct_answer,
    learner_safe_explanation,
    learner_question_payload,
    misconception_review_payload,
)
from app.training_tasks import legacy_lab_guide, parse_training_task_content


class ResourceArchiveRequest(StrictModel):
    archived: bool
    expected_lock_version: int


class ResourceVersionCreateRequest(StrictModel):
    expected_lock_version: int
    source_version_id: str
    content: str
    content_format: Literal["markdown", "json", "text", "html", "python"] = "markdown"


class ResourceVersionRestoreRequest(StrictModel):
    expected_lock_version: int
    version_id: str


class HandoutOutlineItem(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)


class HandoutConcept(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    explanation: str = Field(min_length=1, max_length=12_000)
    sourceRefIds: list[str] = Field(default_factory=list)


class HandoutFigure(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    kind: Literal["chart", "formula", "diagram"]
    content: str = Field(min_length=1, max_length=20_000)
    caption: str = Field(min_length=1, max_length=1_000)
    sourceRefIds: list[str] = Field(default_factory=list)


class HandoutExample(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    body: str = Field(min_length=1, max_length=12_000)
    sourceRefIds: list[str] = Field(default_factory=list)


class HandoutCodeBlock(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    language: str = Field(min_length=1, max_length=32)
    code: str = Field(min_length=1, max_length=30_000)
    explanation: str = Field(min_length=1, max_length=4_000)


class HandoutQuizItem(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    question: str = Field(min_length=1, max_length=2_000)
    options: list[str] = Field(min_length=2, max_length=8)
    correctIndex: int = Field(ge=0)
    explanation: str = Field(min_length=1, max_length=4_000)
    required: bool = True

    @model_validator(mode="after")
    def valid_answer_index(self) -> "HandoutQuizItem":
        if self.correctIndex >= len(self.options):
            raise ValueError("correctIndex must identify an available option")
        return self


class HandoutFootnote(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    sourceRefId: str = Field(min_length=1, max_length=128)
    label: str = Field(min_length=1, max_length=80)
    title: str = Field(min_length=1, max_length=256)
    location: str = Field(min_length=1, max_length=512)


class HandoutProfileAdaptation(StrictModel):
    summary: str = Field(min_length=1, max_length=2_000)
    difficulty: str = Field(min_length=1, max_length=64)
    scaffolding: list[str] = Field(min_length=1, max_length=12)


class HandoutContent(StrictModel):
    schemaVersion: Literal["f5.handout.v1"] = "f5.handout.v1"
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    learningGoal: str = Field(min_length=1, max_length=2_000)
    estimatedMinutes: int = Field(ge=1, le=480)
    outline: list[HandoutOutlineItem] = Field(min_length=1, max_length=30)
    concepts: list[HandoutConcept] = Field(min_length=1, max_length=40)
    figures: list[HandoutFigure] = Field(min_length=1, max_length=20)
    examples: list[HandoutExample] = Field(min_length=1, max_length=30)
    counterExamples: list[HandoutExample] = Field(min_length=1, max_length=30)
    codeBlocks: list[HandoutCodeBlock] = Field(min_length=1, max_length=20)
    miniQuiz: list[HandoutQuizItem] = Field(min_length=1, max_length=20)
    footnotes: list[HandoutFootnote] = Field(min_length=1, max_length=100)
    profileAdaptation: HandoutProfileAdaptation

    @model_validator(mode="after")
    def unique_ids_and_outline_targets(self) -> "HandoutContent":
        sections = (
            self.concepts,
            self.figures,
            self.examples,
            self.counterExamples,
            self.codeBlocks,
            self.miniQuiz,
            self.footnotes,
        )
        content_ids = [item.id for section in sections for item in section]
        if len(set(content_ids)) != len(content_ids):
            raise ValueError("handout section ids must be unique")
        outline_ids = [item.id for item in self.outline]
        if len(set(outline_ids)) != len(outline_ids):
            raise ValueError("handout outline ids must be unique")
        if not set(outline_ids).issubset(set(content_ids)):
            raise ValueError("handout outline must reference existing section ids")
        return self


class MindmapNode(StrictModel):
    nodeId: str = Field(min_length=1, max_length=96)
    parentNodeId: str | None = Field(default=None, max_length=96)
    label: str = Field(min_length=1, max_length=160)
    summary: str = Field(min_length=1, max_length=2_000)
    relation: str = Field(min_length=1, max_length=64)
    depth: int = Field(ge=0, le=20)
    sourceRefs: list[str] = Field(default_factory=list, max_length=40)
    masteryBand: Literal["unknown", "needs_work", "developing", "secure"] = "unknown"
    linkedResourceIds: list[str] = Field(default_factory=list, max_length=20)


class MindmapContent(StrictModel):
    schemaVersion: Literal["f5.mindmap.v1"] = "f5.mindmap.v1"
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    learningGoal: str = Field(min_length=1, max_length=2_000)
    rootNodeId: str = Field(min_length=1, max_length=96)
    nodes: list[MindmapNode] = Field(min_length=1, max_length=500)

    @model_validator(mode="after")
    def valid_connected_tree(self) -> "MindmapContent":
        by_id = {node.nodeId: node for node in self.nodes}
        if len(by_id) != len(self.nodes):
            raise ValueError("mind-map node ids must be unique")
        root = by_id.get(self.rootNodeId)
        if root is None:
            raise ValueError("mind-map root node must exist")
        roots = [node for node in self.nodes if node.parentNodeId is None]
        if roots != [root] or root.depth != 0:
            raise ValueError("mind-map must have exactly one depth-zero root")

        for node in self.nodes:
            if node.nodeId == self.rootNodeId:
                continue
            parent = by_id.get(node.parentNodeId or "")
            if parent is None:
                raise ValueError("mind-map node parent must exist")
            if parent.nodeId == node.nodeId:
                raise ValueError("mind-map node cannot parent itself")
            if node.depth != parent.depth + 1:
                raise ValueError("mind-map node depth must follow its parent")

        for node in self.nodes:
            seen: set[str] = set()
            current = node
            while current.nodeId != self.rootNodeId:
                if current.nodeId in seen:
                    raise ValueError("mind-map parent relation contains a cycle")
                seen.add(current.nodeId)
                parent = by_id.get(current.parentNodeId or "")
                if parent is None:
                    raise ValueError("mind-map node is disconnected from the root")
                current = parent
        return self


class VideoSegment(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    startSeconds: int = Field(ge=0, le=86_400)
    endSeconds: int = Field(ge=1, le=86_400)
    learningFocus: str = Field(min_length=1, max_length=1_000)

    @model_validator(mode="after")
    def valid_time_range(self) -> "VideoSegment":
        if self.endSeconds <= self.startSeconds:
            raise ValueError("video segment end must be after its start")
        return self


class VideoAlternative(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=256)
    url: str = Field(min_length=8, max_length=2_048)
    reason: str = Field(min_length=1, max_length=1_000)

    @model_validator(mode="after")
    def valid_https_url(self) -> "VideoAlternative":
        if not self.url.startswith("https://"):
            raise ValueError("video alternative must use HTTPS")
        return self


class VideoRecommendation(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=256)
    sourceKind: Literal["external_url", "local_ref"]
    availability: Literal["playable", "external_only", "unavailable"]
    url: str | None = Field(default=None, max_length=2_048)
    localRef: str | None = Field(default=None, max_length=1_024)
    durationSeconds: int | None = Field(default=None, ge=1, le=86_400)
    author: str = Field(min_length=1, max_length=256)
    language: str = Field(min_length=1, max_length=64)
    licenseStatus: Literal[
        "verified",
        "reference_only_unverified",
        "repository_license_required",
        "platform_link_only",
        "packaged_demo",
    ]
    licenseLabel: str = Field(min_length=1, max_length=2_000)
    licenseUrl: str | None = Field(default=None, max_length=2_048)
    reason: str = Field(min_length=1, max_length=2_000)
    segments: list[VideoSegment] = Field(default_factory=list, max_length=80)
    sourceRefIds: list[str] = Field(default_factory=list, max_length=100)
    alternatives: list[VideoAlternative] = Field(default_factory=list, max_length=12)

    @model_validator(mode="after")
    def valid_source_and_segments(self) -> "VideoRecommendation":
        if self.sourceKind == "external_url":
            if not self.url or not self.url.startswith("https://") or self.localRef:
                raise ValueError("external video recommendation requires one HTTPS URL")
            if self.availability == "playable":
                raise ValueError("external recommendations cannot claim packaged playback")
        elif not self.localRef or not self.localRef.startswith("/api/") or self.url:
            raise ValueError("local video recommendation requires one controlled API reference")
        elif self.availability != "playable":
            raise ValueError("local video reference must be marked playable")
        if self.licenseUrl and not self.licenseUrl.startswith("https://"):
            raise ValueError("video license URL must use HTTPS")
        segment_ids = [item.id for item in self.segments]
        if len(set(segment_ids)) != len(segment_ids):
            raise ValueError("video segment ids must be unique")
        alternative_ids = [item.id for item in self.alternatives]
        if len(set(alternative_ids)) != len(alternative_ids):
            raise ValueError("video alternative ids must be unique")
        if self.sourceKind == "local_ref" and (self.durationSeconds is None or not self.segments):
            raise ValueError("playable local video requires duration and verified segments")
        if self.sourceKind == "local_ref" and self.licenseStatus not in {"verified", "packaged_demo"}:
            raise ValueError("playable local video requires verified or packaged-demo permission")
        if self.durationSeconds is not None and any(
            segment.endSeconds > self.durationSeconds for segment in self.segments
        ):
            raise ValueError("video segment exceeds the recorded duration")
        ordered = sorted(self.segments, key=lambda item: (item.startSeconds, item.endSeconds))
        if any(left.endSeconds > right.startSeconds for left, right in zip(ordered, ordered[1:])):
            raise ValueError("video segments cannot overlap")
        return self


class VideoRecommendationContent(StrictModel):
    schemaVersion: Literal["f5.video-recommendation.v1"] = "f5.video-recommendation.v1"
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    learningGoal: str = Field(min_length=1, max_length=2_000)
    adaptationSummary: str = Field(min_length=1, max_length=2_000)
    recommendations: list[VideoRecommendation] = Field(default_factory=list, max_length=12)

    @model_validator(mode="after")
    def unique_recommendation_ids(self) -> "VideoRecommendationContent":
        ids = [item.id for item in self.recommendations]
        if len(set(ids)) != len(ids):
            raise ValueError("video recommendation ids must be unique")
        return self


class DeckSlide(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    layout: Literal["title", "content", "two_column", "closing"]
    title: str = Field(min_length=1, max_length=256)
    subtitle: str = Field(default="", max_length=512)
    body: str = Field(default="", max_length=8_000)
    bullets: list[str] = Field(default_factory=list, max_length=16)
    speakerNotes: str = Field(default="", max_length=8_000)
    sourceRefIds: list[str] = Field(default_factory=list, max_length=100)

    @model_validator(mode="after")
    def useful_slide_content(self) -> "DeckSlide":
        if self.layout != "title" and not self.body.strip() and not any(item.strip() for item in self.bullets):
            raise ValueError("content slides require body text or bullets")
        if any(not item.strip() or len(item) > 1_000 for item in self.bullets):
            raise ValueError("deck bullets must be non-empty and concise")
        return self


class DeckContent(StrictModel):
    schemaVersion: Literal["f5.deck.v1"] = "f5.deck.v1"
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    learningGoal: str = Field(min_length=1, max_length=2_000)
    adaptationSummary: str = Field(min_length=1, max_length=2_000)
    slides: list[DeckSlide] = Field(min_length=1, max_length=30)

    @model_validator(mode="after")
    def unique_slide_ids(self) -> "DeckContent":
        ids = [item.id for item in self.slides]
        if len(set(ids)) != len(ids):
            raise ValueError("deck slide ids must be unique")
        return self


class QuizQuestionRef(StrictModel):
    questionVersionId: str = Field(min_length=1, max_length=64)
    position: int = Field(ge=1, le=100)
    required: bool = True


class QuizAdaptation(StrictModel):
    learningGoal: str = Field(min_length=1, max_length=2_000)
    profileSummary: str = Field(min_length=1, max_length=4_000)
    masterySummary: str = Field(min_length=1, max_length=4_000)
    recentLearningSummary: str = Field(min_length=1, max_length=4_000)


class QuizResourceContent(StrictModel):
    schemaVersion: Literal["f6.quiz-resource.v1"] = "f6.quiz-resource.v1"
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    courseTitle: str = Field(min_length=1, max_length=256)
    knowledgePoint: dict[str, str]
    adaptation: QuizAdaptation
    semanticContext: dict[str, Any] = Field(default_factory=dict)
    sourceRefIds: list[str] = Field(min_length=1, max_length=50)
    questions: list[QuizQuestionRef] = Field(min_length=1, max_length=100)
    requiredQuestionCount: int = Field(ge=1, le=100)

    @model_validator(mode="after")
    def valid_question_set(self) -> "QuizResourceContent":
        ids = [item.questionVersionId for item in self.questions]
        positions = [item.position for item in self.questions]
        if len(set(ids)) != len(ids) or len(set(positions)) != len(positions):
            raise ValueError("quiz question ids and positions must be unique")
        if self.requiredQuestionCount > len(self.questions):
            raise ValueError("required quiz question count exceeds the question set")
        return self


class CodeSandboxAdaptation(StrictModel):
    learningGoal: str = Field(min_length=1, max_length=2_000)
    profileSummary: str = Field(min_length=1, max_length=4_000)
    masterySummary: str = Field(min_length=1, max_length=4_000)
    recentLearningSummary: str = Field(min_length=1, max_length=4_000)


class CodeSandboxCheckpoint(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    instruction: str = Field(min_length=1, max_length=2_000)
    completionHint: str = Field(min_length=1, max_length=1_000)


class CodeSandboxFile(StrictModel):
    path: str = Field(min_length=1, max_length=160)
    language: Literal["python"] = "python"
    content: str = Field(min_length=1, max_length=50_000)

    @model_validator(mode="after")
    def safe_relative_path(self) -> "CodeSandboxFile":
        normalized = self.path.replace("\\", "/")
        if normalized.startswith("/") or normalized.startswith("../") or "/../" in normalized:
            raise ValueError("code sandbox file path must remain inside the workspace")
        return self


class CodeSandboxRequiredCheck(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    label: str = Field(min_length=1, max_length=160)
    kind: Literal["exit_code_zero", "stdout_contains", "stderr_empty", "plot_count_at_least"]
    expectedText: str | None = Field(default=None, max_length=1_000)
    minimum: int | None = Field(default=None, ge=1, le=20)

    @model_validator(mode="after")
    def valid_expectation(self) -> "CodeSandboxRequiredCheck":
        if self.kind == "stdout_contains" and not self.expectedText:
            raise ValueError("stdout check requires expectedText")
        if self.kind == "plot_count_at_least" and self.minimum is None:
            raise ValueError("plot count check requires minimum")
        return self


class CodeSandboxRuntime(StrictModel):
    language: Literal["python"] = "python"
    mode: Literal["trusted_local"] = "trusted_local"
    timeoutSeconds: int = Field(default=10, ge=1, le=60)
    boundary: str = Field(min_length=1, max_length=2_000)


class CodeSandboxContent(StrictModel):
    schemaVersion: Literal["f6.code-sandbox.v1"] = "f6.code-sandbox.v1"
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    courseTitle: str = Field(min_length=1, max_length=256)
    knowledgePoint: dict[str, str]
    adaptation: CodeSandboxAdaptation
    checkpoints: list[CodeSandboxCheckpoint] = Field(min_length=1, max_length=12)
    files: list[CodeSandboxFile] = Field(min_length=1, max_length=8)
    entryFile: str = Field(min_length=1, max_length=160)
    standardInput: str = Field(default="", max_length=20_000)
    requiredChecks: list[CodeSandboxRequiredCheck] = Field(min_length=1, max_length=12)
    completionRule: str = Field(min_length=1, max_length=2_000)
    runtime: CodeSandboxRuntime
    sourceRefIds: list[str] = Field(min_length=1, max_length=100)
    linkedTrainingTaskId: str | None = Field(default=None, max_length=64)
    linkedModelId: str | None = Field(default=None, max_length=64)
    semanticContext: dict[str, Any] = Field(default_factory=dict)

    @model_validator(mode="after")
    def valid_resource_contract(self) -> "CodeSandboxContent":
        checkpoint_ids = [item.id for item in self.checkpoints]
        file_paths = [item.path for item in self.files]
        check_ids = [item.id for item in self.requiredChecks]
        if len(set(checkpoint_ids)) != len(checkpoint_ids):
            raise ValueError("code sandbox checkpoint ids must be unique")
        if len(set(file_paths)) != len(file_paths):
            raise ValueError("code sandbox file paths must be unique")
        if len(set(check_ids)) != len(check_ids):
            raise ValueError("code sandbox check ids must be unique")
        if self.entryFile not in set(file_paths):
            raise ValueError("code sandbox entry file must identify a declared file")
        if not self.knowledgePoint.get("id") or not self.knowledgePoint.get("title"):
            raise ValueError("code sandbox knowledge point requires id and title")
        return self


class InteractiveModelParameter(StrictModel):
    key: str = Field(min_length=1, max_length=80)
    label: str = Field(min_length=1, max_length=160)
    kind: Literal["float", "integer"]
    minimum: float
    maximum: float
    step: float = Field(gt=0)
    default: float
    unit: str = Field(default="", max_length=32)
    description: str = Field(min_length=1, max_length=1_000)

    @model_validator(mode="after")
    def valid_range(self) -> "InteractiveModelParameter":
        if self.maximum <= self.minimum:
            raise ValueError("interactive model parameter maximum must exceed minimum")
        if not self.minimum <= self.default <= self.maximum:
            raise ValueError("interactive model parameter default must stay within bounds")
        steps = (self.default - self.minimum) / self.step
        if abs(steps - round(steps)) > 1e-7:
            raise ValueError("interactive model parameter default must follow its step")
        if self.kind == "integer" and any(
            abs(value - round(value)) > 1e-7
            for value in (self.minimum, self.maximum, self.step, self.default)
        ):
            raise ValueError("integer interactive model parameters require integer bounds and default")
        return self


class InteractiveModelConstraint(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=160)
    description: str = Field(min_length=1, max_length=2_000)


class InteractiveModelDataset(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    sampleCount: int = Field(ge=1, le=1_000_000)
    sourceRefIds: list[str] = Field(min_length=1, max_length=100)


class InteractiveModelSeriesDefinition(StrictModel):
    key: str = Field(min_length=1, max_length=80)
    label: str = Field(min_length=1, max_length=160)
    unit: str = Field(default="", max_length=32)
    higherIsBetter: bool


class InteractiveModelChart(StrictModel):
    kind: Literal["line"] = "line"
    xAxisLabel: str = Field(min_length=1, max_length=160)
    yAxisLabel: str = Field(min_length=1, max_length=160)
    series: list[InteractiveModelSeriesDefinition] = Field(min_length=1, max_length=8)

    @model_validator(mode="after")
    def unique_series(self) -> "InteractiveModelChart":
        keys = [item.key for item in self.series]
        if len(set(keys)) != len(keys):
            raise ValueError("interactive model chart series keys must be unique")
        return self


class InteractiveModelObservationPrompt(StrictModel):
    id: str = Field(min_length=1, max_length=96)
    prompt: str = Field(min_length=1, max_length=4_000)
    minimumCharacters: int = Field(default=20, ge=1, le=2_000)
    completionRule: str = Field(min_length=1, max_length=2_000)


class InteractiveModelAdaptation(StrictModel):
    learningGoal: str = Field(min_length=1, max_length=2_000)
    profileSummary: str = Field(min_length=1, max_length=4_000)
    masterySummary: str = Field(min_length=1, max_length=4_000)
    recentLearningSummary: str = Field(min_length=1, max_length=4_000)


class InteractiveModelContent(StrictModel):
    schemaVersion: Literal["f6.interactive-model.v1"] = "f6.interactive-model.v1"
    templateId: str = Field(min_length=1, max_length=96)
    title: str = Field(min_length=1, max_length=256)
    summary: str = Field(min_length=1, max_length=4_000)
    courseTitle: str = Field(min_length=1, max_length=256)
    knowledgePoint: dict[str, str]
    adaptation: InteractiveModelAdaptation
    parameters: list[InteractiveModelParameter] = Field(min_length=1, max_length=12)
    constraints: list[InteractiveModelConstraint] = Field(min_length=1, max_length=12)
    dataset: InteractiveModelDataset
    chart: InteractiveModelChart
    observationPrompt: InteractiveModelObservationPrompt
    completionRule: str = Field(min_length=1, max_length=2_000)
    sourceRefIds: list[str] = Field(min_length=1, max_length=100)
    semanticContext: dict[str, Any] = Field(default_factory=dict)

    @model_validator(mode="after")
    def valid_resource_contract(self) -> "InteractiveModelContent":
        parameter_keys = [item.key for item in self.parameters]
        constraint_ids = [item.id for item in self.constraints]
        if len(set(parameter_keys)) != len(parameter_keys):
            raise ValueError("interactive model parameter keys must be unique")
        if len(set(constraint_ids)) != len(constraint_ids):
            raise ValueError("interactive model constraint ids must be unique")
        if set(self.dataset.sourceRefIds) - set(self.sourceRefIds):
            raise ValueError("interactive model dataset sources must be declared by the resource")
        if not self.knowledgePoint.get("id") or not self.knowledgePoint.get("title"):
            raise ValueError("interactive model knowledge point requires id and title")
        return self


RESOURCE_CATALOG_TYPES = (
    "handout",
    "mind_map",
    "quiz",
    "training_task",
    "code_sandbox",
    "video",
    "deck",
    "interactive_model",
    "feynman",
)

RESOURCE_TYPE_LABELS = {
    "handout": "个性化讲义",
    "mind_map": "思维导图",
    "quiz": "测验题",
    "training_task": "实训任务书",
    "code_sandbox": "代码沙盒",
    "video": "视频推荐",
    "deck": "PPT/网页课件",
    "interactive_model": "可交互模型",
    "feynman": "费曼复述",
}

RESOURCE_TYPE_ALIASES = {
    "lesson": "handout",
    "lecture_note": "handout",
    "lab_guide": "training_task",
    "tiered_quiz": "quiz",
    "code": "code_sandbox",
    "code_example": "code_sandbox",
    "notebook": "code_sandbox",
    "ppt": "deck",
    "web_deck": "deck",
    "interactive_skill_artifact": "interactive_model",
    "mindmap": "mind_map",
}

RESOURCE_ROUTES = {
    "handout": "/workspace/resources/handouts/{resource_id}",
    "mind_map": "/workspace/resources/mindmaps/{resource_id}",
    "quiz": "/workspace/resources/quizzes/{resource_id}",
    "training_task": "/workspace/resources/training-tasks/{resource_id}",
    "code_sandbox": "/workspace/code/{resource_id}",
    "video": "/workspace/resources/videos/{resource_id}",
    "deck": "/workspace/resources/decks/{resource_id}",
    "interactive_model": "/workspace/resources/models/{resource_id}",
    "feynman": "/workspace/feynman/{resource_id}",
}


def canonical_resource_type(artifact_type: str, metadata: dict[str, Any] | None = None) -> str:
    declared = str((metadata or {}).get("resource_type") or artifact_type).strip()
    return RESOURCE_TYPE_ALIASES.get(declared, declared)


def resource_route(resource_type: str, resource_id: str) -> str | None:
    template = RESOURCE_ROUTES.get(resource_type)
    return template.format(resource_id=resource_id) if template else None


def _metadata_values(metadata: dict[str, Any], key: str) -> list[str]:
    value = metadata.get(key)
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    return [text] if text else []


def _preview_text(content: str, limit: int = 220) -> str:
    try:
        structured = json.loads(content)
    except (json.JSONDecodeError, TypeError):
        structured = None
    if isinstance(structured, dict):
        content = " ".join(str(structured.get(key) or "") for key in ("title", "summary", "learningGoal"))
    clean = " ".join(content.replace("#", " ").split())
    return clean[:limit] + ("..." if len(clean) > limit else "")


def _markdown_title(content: str, fallback: str) -> str:
    for line in content.splitlines():
        match = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if match:
            return match.group(1).strip()[:256]
    return fallback


def _legacy_handout_projection(
    *, artifact: DBArtifact, version: DBArtifactVersion, source_refs: list[dict[str, Any]],
) -> dict[str, Any]:
    content = learner_safe_artifact_content(artifact.artifact_type, version.content)
    title = _markdown_title(content, artifact.title)
    clean = " ".join(line.strip(" #>*-") for line in content.splitlines() if line.strip())
    return {
        "schema_version": "f5.handout-projection.v1",
        "content_schema_version": None,
        "mode": "legacy_markdown",
        "content_format": version.content_format,
        "source_refs": source_refs,
        "structured": None,
        "legacy": {
            "title": title,
            "summary": clean[:320],
            "markdown": content,
        },
        "source_readiness": {
            "recorded": bool(source_refs),
            "available": bool(source_refs) and all(item.get("available") is True for item in source_refs),
            "publishable": False,
            "reason": "旧讲义可继续阅读和编辑，但需要补齐结构化章节与可用来源后才能作为完整讲义发布。",
        },
        "warnings": ["这是旧版 Markdown 讲义的兼容视图，系统不会补造缺失章节或来源。"],
    }


def parse_handout_content(content: str) -> HandoutContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError("handout_content_invalid", "structured handout content must be valid JSON") from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError("handout_content_invalid", "structured handout content must be a JSON object")
    try:
        return HandoutContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError("handout_content_incomplete", "structured handout is missing required learning sections") from exc


def parse_mindmap_content(content: str) -> MindmapContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError("mindmap_content_invalid", "structured mind-map content must be valid JSON") from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError("mindmap_content_invalid", "structured mind-map content must be a JSON object")
    try:
        return MindmapContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError("mindmap_content_incomplete", "structured mind-map is not a valid connected tree") from exc


def parse_video_recommendation_content(content: str) -> VideoRecommendationContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError(
            "video_content_invalid", "structured video recommendations must be valid JSON",
        ) from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError(
            "video_content_invalid", "structured video recommendations must be a JSON object",
        )
    try:
        return VideoRecommendationContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError(
            "video_content_incomplete", "video recommendations are missing valid source, license or segment data",
        ) from exc


def parse_deck_content(content: str) -> DeckContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError("deck_content_invalid", "structured deck content must be valid JSON") from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError("deck_content_invalid", "structured deck content must be a JSON object")
    try:
        return DeckContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError(
            "deck_content_incomplete", "structured deck is missing a valid slide tree or slide content",
        ) from exc


def parse_quiz_content(content: str) -> QuizResourceContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError("quiz_content_invalid", "structured quiz content must be valid JSON") from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError("quiz_content_invalid", "structured quiz content must be a JSON object")
    try:
        return QuizResourceContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError("quiz_content_incomplete", "structured quiz is missing required learning fields") from exc


def parse_code_sandbox_content(content: str) -> CodeSandboxContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError(
            "code_sandbox_content_invalid",
            "structured code sandbox content must be valid JSON",
        ) from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError(
            "code_sandbox_content_invalid",
            "structured code sandbox content must be a JSON object",
        )
    try:
        return CodeSandboxContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError(
            "code_sandbox_content_incomplete",
            "structured code sandbox is missing tasks, files, checks, sources or runtime boundaries",
        ) from exc


def parse_interactive_model_content(content: str) -> InteractiveModelContent:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as exc:
        raise M5CapabilityError(
            "interactive_model_content_invalid",
            "structured interactive model content must be valid JSON",
        ) from exc
    if not isinstance(payload, dict):
        raise M5CapabilityError(
            "interactive_model_content_invalid",
            "structured interactive model content must be a JSON object",
        )
    try:
        structured = InteractiveModelContent.model_validate(payload)
    except ValueError as exc:
        raise M5CapabilityError(
            "interactive_model_content_incomplete",
            "structured interactive model is missing parameters, sources or observation rules",
        ) from exc
    try:
        from app.interactive_lab import validate_resource_template_content

        validate_resource_template_content(structured)
    except ValueError as exc:
        raise M5CapabilityError(
            "interactive_model_template_mismatch",
            str(exc),
        ) from exc
    return structured


def _handout_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    if version.content_format != "json":
        return _legacy_handout_projection(artifact=artifact, version=version, source_refs=source_refs)
    try:
        structured = parse_handout_content(version.content)
    except M5CapabilityError as exc:
        warning = {
            "handout_content_invalid": "当前讲义格式无法解析，请在编辑区检查 JSON 格式后重新保存。",
            "handout_content_incomplete": "当前讲义缺少必需的学习章节，请补齐后重新保存。",
        }.get(exc.code, "当前讲义暂时无法按结构化内容展示，请检查并重新保存。")
        return {
            "schema_version": "f5.handout-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "content_format": version.content_format,
            "source_refs": source_refs,
            "structured": None,
            "legacy": None,
            "invalid": {
                "raw_content": learner_safe_artifact_content(
                    artifact.artifact_type, version.content
                ),
                "error_code": exc.code,
            },
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": bool(source_refs) and all(item.get("available") is True for item in source_refs),
                "publishable": False,
                "reason": "讲义结构不完整，修复全部必需章节后才能发布。",
            },
            "warnings": [warning],
        }
    source_ids = {str(item.get("source_ref_id") or "") for item in source_refs}
    cited_ids = {
        *(ref_id for item in structured.concepts for ref_id in item.sourceRefIds),
        *(ref_id for item in structured.figures for ref_id in item.sourceRefIds),
        *(ref_id for item in structured.examples for ref_id in item.sourceRefIds),
        *(ref_id for item in structured.counterExamples for ref_id in item.sourceRefIds),
        *(item.sourceRefId for item in structured.footnotes),
    }
    unknown_citations = sorted(cited_ids - source_ids)
    recorded = bool(source_refs)
    available = recorded and all(item.get("available") is True for item in source_refs)
    publishable = available and not unknown_citations
    reason = (
        "来源完整，可进入发布检查。" if publishable
        else "讲义引用了未登记来源，需修正脚注。" if unknown_citations
        else "来源已失效，请恢复来源后再发布。" if recorded
        else "未记录来源，当前讲义不能发布。"
    )
    return {
        "schema_version": "f5.handout-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "content_format": version.content_format,
        "source_refs": source_refs,
        "structured": structured.model_dump(),
        "legacy": None,
        "source_readiness": {
            "recorded": recorded,
            "available": available,
            "publishable": publishable,
            "reason": reason,
            "unknown_citation_ids": unknown_citations,
        },
        "warnings": [] if publishable else [reason],
    }


def _clean_legacy_mindmap_label(value: str) -> str:
    """Safely project the old Mermaid-like label syntax into plain text."""
    label = re.sub(r"^[-*+]\s+", "", str(value or "").strip())
    label = label.replace("**", "").replace("`", "")
    for opening, closing in (("((", "))"), ("{{", "}}"), ("(", ")"), ("[", "]")):
        match = re.match(
            rf"^[^({{\[\]\s]+\s*{re.escape(opening)}\s*(.*?)\s*{re.escape(closing)}$",
            label,
        )
        if match:
            label = match.group(1).strip()
            break
    return html.unescape(label.strip().strip('"'))[:160]


def _legacy_mindmap_hierarchy(content: str, fallback: str) -> list[dict[str, Any]]:
    """Parse old Markdown and Mermaid-like indentation without inventing evidence."""
    rows: list[dict[str, Any]] = []
    stack: list[tuple[int, str, bool]] = []
    lines = str(content or "").splitlines()
    mermaid_mode = any(line.strip().lower() == "mindmap" for line in lines)
    mermaid_base_indent: int | None = None
    markdown_root_heading = False
    for raw_line in lines:
        stripped = raw_line.strip()
        if not stripped or stripped.lower() == "mindmap" or stripped.startswith("%%"):
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        bullet = re.match(r"^([-*+])\s+(.+)$", stripped)
        if heading:
            depth = len(heading.group(1)) - 1
            label = heading.group(2).strip()
            if depth == 0:
                markdown_root_heading = True
        elif bullet:
            indent = max(0, len(raw_line) - len(raw_line.lstrip(" ")))
            depth = (1 + indent // 2) if markdown_root_heading else indent // 2
            label = bullet.group(2).strip()
        elif mermaid_mode and not stripped.startswith("```"):
            # The old renderer treated each indented Mermaid-like line as a node.
            indent = len(raw_line) - len(raw_line.lstrip(" "))
            if mermaid_base_indent is None:
                mermaid_base_indent = indent
            depth = max(0, indent - mermaid_base_indent) // 2
            label = stripped
        else:
            continue
        label = _clean_legacy_mindmap_label(label)
        if not label:
            continue
        while stack and stack[-1][0] >= depth:
            stack.pop()
        node_id = f"legacy-node-{len(rows) + 1}"
        parent_id = stack[-1][1] if stack else None
        rows.append({
            "nodeId": node_id,
            "parentNodeId": parent_id,
            "label": label,
            "depth": depth if parent_id is not None else 0,
        })
        stack.append((depth, node_id, bool(heading)))
    if rows:
        return rows
    return [{
        "nodeId": "legacy-node-1",
        "parentNodeId": None,
        "label": _clean_legacy_mindmap_label(fallback),
        "depth": 0,
    }]


def _legacy_mindmap_projection(
    *, artifact: DBArtifact, version: DBArtifactVersion, source_refs: list[dict[str, Any]],
) -> dict[str, Any]:
    content = learner_safe_artifact_content(artifact.artifact_type, version.content)
    return {
        "schema_version": "f5.mindmap-projection.v1",
        "content_schema_version": None,
        "mode": "legacy_markdown",
        "content_format": version.content_format,
        "source_refs": source_refs,
        "linked_resources": [],
        "structured": None,
        "legacy": {
            "title": _markdown_title(content, artifact.title),
            "markdown": content,
            "hierarchy": _legacy_mindmap_hierarchy(content, artifact.title),
            "render_profile": "legacy-compatible-v1",
        },
        "source_readiness": {
            "recorded": bool(source_refs),
            "available": bool(source_refs) and all(item.get("available") is True for item in source_refs),
            "publishable": False,
            "reason": "旧版导图可继续按层级列表阅读，但需要转换为结构化节点后才能使用完整画布、来源和资源深链。",
        },
        "warnings": ["这是旧版 Markdown 导图的兼容视图，系统不会补造节点关系、掌握区间或关联资源。"],
    }


def _linked_resource_projection(
    session: Any, artifact: DBArtifact, resource_ids: list[str],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for resource_id in dict.fromkeys(resource_ids):
        linked = session.get(DBArtifact, resource_id)
        if (
            linked is None
            or linked.owner_public_id != artifact.owner_public_id
            or linked.course_id not in {None, artifact.course_id}
        ):
            output.append({"id": resource_id, "available": False})
            continue
        version = session.get(DBArtifactVersion, linked.current_version_id) if linked.current_version_id else None
        metadata = dict(version.metadata_json or {}) if version else {}
        linked_type = canonical_resource_type(linked.artifact_type, metadata)
        output.append({
            "id": linked.id,
            "title": linked.title,
            "type": linked_type,
            "type_label": RESOURCE_TYPE_LABELS.get(linked_type, linked_type),
            "route": resource_route(linked_type, linked.id),
            "available": linked.archived_at is None and version is not None,
        })
    return output


def _canonical_binding(
    *, artifact: DBArtifact, version: DBArtifactVersion, protocol: dict[str, Any],
) -> dict[str, Any]:
    """Expose the exact artifact/version and frozen learning/source context."""
    metadata = dict(version.metadata_json or {})
    source_refs = list(protocol.get("source_refs") or [])
    return {
        "schema_version": "edumatrix.canonical-resource-binding.v1",
        "artifact_id": artifact.id,
        "artifact_version_id": version.id,
        "artifact_version_number": version.version_number,
        "resource_type": canonical_resource_type(artifact.artifact_type, metadata),
        "course_id": artifact.course_id,
        "knowledge_point_id": metadata.get("knowledge_point_id"),
        "chapter_id": metadata.get("chapter_id"),
        "profile_snapshot_id": metadata.get("profile_snapshot_id"),
        "source_ref_ids": [str(item.get("source_ref_id") or "") for item in source_refs],
        "source_refs": source_refs,
        "source_status": protocol.get("status"),
    }


def _linked_resource_bindings(
    session: Any, artifact: DBArtifact, resource_ids: list[str],
) -> list[dict[str, Any]]:
    """Resolve linked resource IDs to their current canonical versions without widening access."""
    output: list[dict[str, Any]] = []
    for resource_id in dict.fromkeys(resource_ids):
        linked = session.get(DBArtifact, resource_id)
        if (
            linked is None
            or linked.owner_public_id != artifact.owner_public_id
            or linked.course_id not in {None, artifact.course_id}
        ):
            output.append({"id": resource_id, "available": False})
            continue
        version = session.get(DBArtifactVersion, linked.current_version_id) if linked.current_version_id else None
        if version is None:
            output.append({"id": resource_id, "available": False})
            continue
        protocol = _resource_protocol(session, linked, version)
        metadata = dict(version.metadata_json or {})
        linked_type = canonical_resource_type(linked.artifact_type, metadata)
        output.append({
            "id": linked.id,
            "title": linked.title,
            "type": linked_type,
            "type_label": RESOURCE_TYPE_LABELS.get(linked_type, linked_type),
            "route": resource_route(linked_type, linked.id),
            "available": linked.archived_at is None and protocol.get("status") == "available",
            "artifact_id": linked.id,
            "artifact_version_id": version.id,
            "course_id": linked.course_id,
            "knowledge_point_id": metadata.get("knowledge_point_id"),
            "profile_snapshot_id": metadata.get("profile_snapshot_id"),
            "source_ref_ids": [str(item.get("source_ref_id") or "") for item in protocol.get("source_refs") or []],
            "source_status": protocol.get("status"),
        })
    return output


def _mindmap_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    binding = _canonical_binding(artifact=artifact, version=version, protocol=protocol)
    if version.content_format != "json":
        projection = _legacy_mindmap_projection(
            artifact=artifact, version=version, source_refs=source_refs,
        )
        projection["canonical_binding"] = binding
        return projection
    try:
        structured = parse_mindmap_content(version.content)
    except M5CapabilityError as exc:
        warning = {
            "mindmap_content_invalid": "当前思维导图格式无法解析，请检查资源内容后重试。",
            "mindmap_content_incomplete": "当前思维导图的节点层级或关系不完整，已切换为安全阅读状态。",
        }.get(exc.code, "当前思维导图暂时无法按结构化画布展示。")
        raw_content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f5.mindmap-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "content_format": version.content_format,
            "source_refs": source_refs,
            "canonical_binding": binding,
            "linked_resource_bindings": [],
            "linked_resources": [],
            "structured": None,
            "legacy": None,
            "invalid": {
                "raw_content": raw_content,
                "hierarchy": _legacy_mindmap_hierarchy(raw_content, artifact.title),
                "error_code": exc.code,
                "render_profile": "legacy-compatible-v1",
            },
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": bool(source_refs) and all(item.get("available") is True for item in source_refs),
                "publishable": False,
                "reason": "导图结构不完整，修复节点关系后才能使用完整画布。",
            },
            "warnings": [warning],
        }

    source_ids = {str(item.get("source_ref_id") or "") for item in source_refs}
    cited_ids = {ref_id for node in structured.nodes for ref_id in node.sourceRefs}
    unknown_citations = sorted(cited_ids - source_ids)
    linked_ids = [resource_id for node in structured.nodes for resource_id in node.linkedResourceIds]
    linked_resources = _linked_resource_projection(session, artifact, linked_ids)
    linked_bindings = _linked_resource_bindings(session, artifact, linked_ids)
    recorded = bool(source_refs)
    available = recorded and all(item.get("available") is True for item in source_refs)
    publishable = available and not unknown_citations
    reason = (
        "来源完整，可进入导图学习。" if publishable
        else "导图引用了未登记来源，需修正节点来源。" if unknown_citations
        else "来源已失效，请恢复来源后再继续。" if recorded
        else "未记录来源，当前导图只能作为未发布草稿查看。"
    )
    return {
        "schema_version": "f5.mindmap-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "content_format": version.content_format,
        "source_refs": source_refs,
        "canonical_binding": binding,
        "linked_resource_bindings": linked_bindings,
        "linked_resources": linked_resources,
        "structured": structured.model_dump(),
        "legacy": None,
        "source_readiness": {
            "recorded": recorded,
            "available": available,
            "publishable": publishable,
            "reason": reason,
            "unknown_citation_ids": unknown_citations,
        },
        "warnings": [] if publishable else [reason],
    }


def _quiz_attempt_projection(
    session: Any,
    *,
    attempt: DBAssessmentAttempt,
    question: DBAssessmentQuestion,
    version: DBAssessmentQuestionVersion,
) -> dict[str, Any]:
    grading = session.query(DBAssessmentGradingResult).filter_by(
        attempt_id=attempt.id,
    ).order_by(DBAssessmentGradingResult.created_at.desc(), DBAssessmentGradingResult.id.desc()).first()
    review = session.query(DBMisconceptionReview).filter_by(attempt_id=attempt.id).first()
    payload = learner_attempt_payload(attempt, grading)
    payload.update({
        "grading": grading_result_payload(grading),
        "answer_explanation": learner_safe_explanation(
            version.explanation,
            prompt=version.prompt,
            options=version.options,
            answer_key=version.answer_key,
            question_type=question.question_type,
        ),
        "correct_answer": learner_correct_answer(question.question_type, version.answer_key or {}),
        "misconception_review": misconception_review_payload(review),
    })
    return payload


def _quiz_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    source_refs = _resource_protocol(session, artifact, version)["source_refs"]
    if version.content_format != "json":
        return {
            "schema_version": "f6.quiz-projection.v1",
            "mode": "legacy_readonly",
            "content_format": version.content_format,
            "structured": None,
            "questions": [],
            "source_refs": source_refs,
            "warnings": ["旧测验题资源没有结构化题目版本，仍可查看历史，但不能伪装为可判分题组。"],
        }
    try:
        structured = parse_quiz_content(version.content)
    except M5CapabilityError as exc:
        return {
            "schema_version": "f6.quiz-projection.v1",
            "mode": "invalid_structured",
            "content_format": version.content_format,
            "structured": None,
            "questions": [],
            "source_refs": source_refs,
            "error": {"code": exc.code, "message": str(exc)},
            "raw_content": version.content,
            "warnings": ["当前版本结构不完整，已保留原内容且没有补造题目。"],
        }
    questions = []
    correct_question_ids: set[str] = set()
    attempted_question_ids: set[str] = set()
    for reference in sorted(structured.questions, key=lambda item: item.position):
        question_version = session.get(DBAssessmentQuestionVersion, reference.questionVersionId)
        question = session.get(DBAssessmentQuestion, question_version.question_id) if question_version else None
        if question_version is None or question is None or question.course_id != artifact.course_id:
            questions.append({
                "id": reference.questionVersionId,
                "position": reference.position,
                "required": reference.required,
                "status": "unavailable",
                "error": "题目版本不存在或不属于当前课程。",
                "attempts": [],
                "latest_attempt": None,
            })
            continue
        attempts = session.query(DBAssessmentAttempt).filter_by(
            owner_public_id=artifact.owner_public_id,
            question_version_id=question_version.id,
            mode="practice",
        ).order_by(DBAssessmentAttempt.submitted_at.desc(), DBAssessmentAttempt.id.desc()).all()
        attempt_payloads = [
            _quiz_attempt_projection(session, attempt=item, question=question, version=question_version)
            for item in attempts
        ]
        if attempt_payloads:
            attempted_question_ids.add(question_version.id)
        if any(item.get("correct") is True for item in attempt_payloads):
            correct_question_ids.add(question_version.id)
        questions.append({
            **learner_question_payload(question_version, question),
            "position": reference.position,
            "required": reference.required,
            "status": "available",
            "source_ref_ids": list(question_version.source_ref_ids or []),
            "explanation_available": bool(question_version.explanation),
            "attempts": attempt_payloads,
            "latest_attempt": attempt_payloads[0] if attempt_payloads else None,
        })
    required_ids = {item.questionVersionId for item in structured.questions if item.required}
    required_target = min(structured.requiredQuestionCount, len(required_ids))
    required_correct = len(required_ids & correct_question_ids)
    return {
        "schema_version": "f6.quiz-projection.v1",
        "mode": "structured",
        "content_format": version.content_format,
        "structured": structured.model_dump(mode="json"),
        "questions": questions,
        "source_refs": source_refs,
        "progress": {
            "question_count": len(structured.questions),
            "attempted_count": len(attempted_question_ids),
            "correct_count": len(correct_question_ids),
            "required_count": required_target,
            "required_correct_count": required_correct,
            "completed": required_target > 0 and required_correct >= required_target,
        },
        "warnings": [],
    }


def _training_task_link(
    session: Any,
    artifact: DBArtifact,
    resource_id: str | None,
    expected_type: str,
) -> dict[str, Any] | None:
    if not resource_id:
        return None
    linked = session.get(DBArtifact, resource_id)
    if (
        linked is None
        or linked.owner_public_id != artifact.owner_public_id
        or linked.course_id not in {None, artifact.course_id}
    ):
        return {"id": resource_id, "type": expected_type, "available": False}
    version = session.get(DBArtifactVersion, linked.current_version_id) if linked.current_version_id else None
    metadata = dict(version.metadata_json or {}) if version else {}
    linked_type = canonical_resource_type(linked.artifact_type, metadata)
    if linked_type != expected_type:
        return {"id": resource_id, "type": expected_type, "available": False}
    return {
        "id": linked.id,
        "title": linked.title,
        "type": linked_type,
        "type_label": RESOURCE_TYPE_LABELS[linked_type],
        "route": resource_route(linked_type, linked.id),
        "available": linked.archived_at is None and version is not None,
    }


def _training_task_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    if version.content_format != "json":
        content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f5.training-task-projection.v1",
            "content_schema_version": None,
            "mode": "legacy_document",
            "source_refs": source_refs,
            "structured": None,
            "legacy": {"title": _markdown_title(content, artifact.title), "content": content},
            "linked_resources": [],
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": bool(source_refs) and all(item.get("available") is True for item in source_refs),
                "publishable": False,
                "reason": "旧版任务说明可继续阅读，但需要补齐检查点、提交物和评分标准后才能作为完整实训任务书发布。",
            },
            "warnings": ["这是旧版任务说明的兼容视图，系统不会补造缺失的执行要求或关联工具。"],
        }
    try:
        structured = parse_training_task_content(version.content)
    except M5CapabilityError as exc:
        legacy = legacy_lab_guide(version.content)
        if legacy is not None:
            return {
                "schema_version": "f5.training-task-projection.v1",
                "content_schema_version": None,
                "mode": "legacy_lab_guide",
                "source_refs": source_refs,
                "structured": None,
                "legacy": legacy,
                "linked_resources": [],
                "source_readiness": {
                    "recorded": bool(source_refs),
                    "available": bool(source_refs) and all(item.get("available") is True for item in source_refs),
                    "publishable": False,
                    "reason": "旧版实验说明已保留，需要转换为实训任务书后才能使用评分表与关联工具。",
                },
                "warnings": ["当前展示的是旧版实验说明，原步骤与故障分支均已保留。"],
            }
        warning = {
            "training_task_content_invalid": "当前实训任务书格式无法解析，请检查资源内容后重试。",
            "training_task_content_incomplete": "当前实训任务书缺少环境、检查点、提交物或评分标准。",
        }.get(exc.code, "当前实训任务书暂时无法按结构化内容展示。")
        raw_content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f5.training-task-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "source_refs": source_refs,
            "structured": None,
            "legacy": None,
            "invalid": {"raw_content": raw_content, "error_code": exc.code},
            "linked_resources": [],
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": False,
                "publishable": False,
                "reason": "任务书结构不完整，修复全部必需执行章节后才能继续。",
            },
            "warnings": [warning],
        }

    source_ids = {str(item.get("source_ref_id") or "") for item in source_refs}
    cited_ids = {*structured.sourceRefs, *(ref for step in structured.steps for ref in step.sourceRefIds)}
    unknown_citations = sorted(cited_ids - source_ids)
    linked_resources = [
        item for item in (
            _training_task_link(session, artifact, structured.linkedSandboxId, "code_sandbox"),
            _training_task_link(session, artifact, structured.linkedModelId, "interactive_model"),
        ) if item is not None
    ]
    recorded = bool(source_refs)
    available = recorded and all(item.get("available") is True for item in source_refs)
    publishable = available and not unknown_citations
    reason = (
        "来源完整，可按任务书进入实训。" if publishable
        else "任务书引用了未登记来源，需修正步骤来源。" if unknown_citations
        else "来源已失效，请恢复来源后再开始实训。" if recorded
        else "未记录来源，当前任务书只能作为未发布草稿查看。"
    )
    return {
        "schema_version": "f5.training-task-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "source_refs": source_refs,
        "structured": structured.model_dump(),
        "legacy": None,
        "linked_resources": linked_resources,
        "source_readiness": {
            "recorded": recorded,
            "available": available,
            "publishable": publishable,
            "reason": reason,
            "unknown_citation_ids": unknown_citations,
        },
        "warnings": [] if publishable else [reason],
    }


def _source_readiness(
    source_refs: list[dict[str, Any]], cited_ids: set[str], *, ready: str, missing: str,
) -> dict[str, Any]:
    source_ids = {str(item.get("source_ref_id") or "") for item in source_refs}
    unknown = sorted(cited_ids - source_ids)
    recorded = bool(source_refs)
    available = recorded and all(item.get("available") is True for item in source_refs)
    publishable = available and not unknown
    reason = (
        ready if publishable
        else missing if unknown
        else "来源已失效，请恢复来源后再继续。" if recorded
        else "未记录课程来源，当前内容只能作为未发布草稿查看。"
    )
    return {
        "recorded": recorded,
        "available": available,
        "publishable": publishable,
        "reason": reason,
        "unknown_citation_ids": unknown,
    }


def _code_sandbox_links(
    session: Any, artifact: DBArtifact, structured: CodeSandboxContent,
) -> list[dict[str, Any]]:
    links: list[dict[str, Any]] = []
    seen: set[str] = set()

    def append(resource_id: str | None, resource_type: str) -> None:
        if not resource_id or resource_id in seen:
            return
        link = _training_task_link(session, artifact, resource_id, resource_type)
        if link is not None:
            links.append(link)
            seen.add(resource_id)

    append(structured.linkedTrainingTaskId, "training_task")
    append(structured.linkedModelId, "interactive_model")
    candidates = session.query(DBArtifact).filter_by(
        owner_public_id=artifact.owner_public_id,
        course_id=artifact.course_id,
    ).filter(DBArtifact.id != artifact.id).order_by(DBArtifact.created_at.desc()).all()
    for candidate in candidates:
        version = session.get(DBArtifactVersion, candidate.current_version_id) if candidate.current_version_id else None
        metadata = dict(version.metadata_json or {}) if version else {}
        if version is None or version.content_format != "json":
            continue
        if canonical_resource_type(candidate.artifact_type, metadata) != "training_task":
            continue
        try:
            task = parse_training_task_content(version.content)
        except M5CapabilityError:
            continue
        if task.linkedSandboxId != artifact.id:
            continue
        append(candidate.id, "training_task")
        append(task.linkedModelId, "interactive_model")
        break
    return links


def _code_sandbox_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    if version.content_format != "json":
        content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f6.code-sandbox-projection.v1",
            "content_schema_version": None,
            "mode": "legacy_readonly",
            "source_refs": source_refs,
            "structured": None,
            "legacy": {"title": _markdown_title(content, artifact.title), "content": content},
            "linked_resources": [],
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": False,
                "publishable": False,
                "reason": "旧版代码说明可继续阅读，但不能作为正式可运行任务。",
            },
            "warnings": ["这是旧版只读代码说明；系统不会补造任务检查点、运行历史或完成证据。"],
        }
    try:
        structured = parse_code_sandbox_content(version.content)
    except M5CapabilityError as exc:
        return {
            "schema_version": "f6.code-sandbox-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "source_refs": source_refs,
            "structured": None,
            "legacy": None,
            "invalid": {
                "raw_content": learner_safe_artifact_content(artifact.artifact_type, version.content),
                "error_code": exc.code,
            },
            "linked_resources": [],
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": False,
                "publishable": False,
                "reason": "代码沙盒结构不完整，当前版本不能保存、运行或完成。",
            },
            "warnings": ["请恢复有效版本或重新生成该资源；现有内容不会被覆盖。"],
        }
    readiness = _source_readiness(
        source_refs,
        set(structured.sourceRefIds),
        ready="课程来源完整，可保存代码并发起真实运行。",
        missing="代码沙盒引用了未登记课程来源，请修正来源后再运行。",
    )
    metadata_rule = " ".join(str((version.metadata_json or {}).get("completion_rule") or "").split())
    content_rule = " ".join(structured.completionRule.split())
    if not metadata_rule or metadata_rule != content_rule:
        readiness = {
            **readiness,
            "publishable": False,
            "reason": "代码沙盒完成规则与服务端资源版本不一致，请恢复有效版本。",
        }
    return {
        "schema_version": "f6.code-sandbox-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "source_refs": source_refs,
        "structured": structured.model_dump(),
        "legacy": None,
        "linked_resources": _code_sandbox_links(session, artifact, structured),
        "source_readiness": readiness,
        "warnings": [] if readiness["publishable"] else [readiness["reason"]],
    }


def _code_sandbox_completion(
    session: Any, artifact: DBArtifact, current_version: DBArtifactVersion,
) -> dict[str, Any]:
    events = session.query(DBLearningEvent).filter_by(
        owner_public_id=artifact.owner_public_id,
        event_type="code_tests_passed",
        object_type="artifact",
        object_id=artifact.id,
    ).order_by(DBLearningEvent.recorded_at.desc()).limit(50).all()
    for event in events:
        payload = dict(event.payload or {})
        if payload.get("resource_version_id") != current_version.id:
            continue
        return {
            "status": "completed",
            "event_id": event.id,
            "run_attempt_id": payload.get("run_id"),
            "evidence_status": event.evidence_status,
            "mastery_eligible": bool(event.mastery_eligible),
            "completed_at": event.recorded_at.isoformat() if event.recorded_at else None,
        }
    return {
        "status": "not_recorded",
        "event_id": None,
        "run_attempt_id": None,
        "evidence_status": None,
        "mastery_eligible": False,
        "completed_at": None,
    }


def _interactive_model_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    if version.content_format != "json":
        content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f6.interactive-model-projection.v1",
            "content_schema_version": None,
            "mode": "legacy_readonly",
            "source_refs": source_refs,
            "structured": None,
            "legacy": {"title": _markdown_title(content, artifact.title), "content": content},
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": False,
                "publishable": False,
                "reason": "旧版模型说明可继续阅读，但不能作为可运行实验。",
            },
            "warnings": ["这是旧版只读说明；系统不会补造参数、图表、观察提示或实验结果。"],
        }
    try:
        structured = parse_interactive_model_content(version.content)
    except M5CapabilityError as exc:
        template_mismatch = exc.code == "interactive_model_template_mismatch"
        reason = (
            str(exc)
            if template_mismatch
            else "可交互模型结构不完整，当前版本不能运行或提交观察。"
        )
        warning = (
            f"{str(exc)} 请选择其他有效版本或重新生成该资源。"
            if template_mismatch
            else "请恢复有效版本或重新生成该资源；现有内容不会被覆盖。"
        )
        return {
            "schema_version": "f6.interactive-model-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "source_refs": source_refs,
            "structured": None,
            "legacy": None,
            "invalid": {
                "raw_content": learner_safe_artifact_content(artifact.artifact_type, version.content),
                "error_code": exc.code,
            },
            "source_readiness": {
                "recorded": bool(source_refs),
                "available": False,
                "publishable": False,
                "reason": reason,
            },
            "warnings": [warning],
        }
    cited_ids = {*structured.sourceRefIds, *structured.dataset.sourceRefIds}
    readiness = _source_readiness(
        source_refs,
        cited_ids,
        ready="课程来源完整，可运行实验并提交观察。",
        missing="可交互模型引用了未登记课程来源，请修正来源后再运行。",
    )
    return {
        "schema_version": "f6.interactive-model-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "source_refs": source_refs,
        "structured": structured.model_dump(),
        "legacy": None,
        "source_readiness": readiness,
        "warnings": [] if readiness["publishable"] else [readiness["reason"]],
    }


def _video_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    if version.content_format != "json":
        content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f5.video-projection.v1",
            "content_schema_version": None,
            "mode": "legacy_document",
            "source_refs": source_refs,
            "structured": None,
            "legacy": {"title": _markdown_title(content, artifact.title), "content": content},
            "source_readiness": {
                "recorded": bool(source_refs), "available": False, "publishable": False,
                "reason": "旧视频说明可继续阅读，但缺少可校验的来源、许可状态和推荐定位。",
            },
            "warnings": ["这是旧版视频说明，系统不会把文本说明伪装成可播放视频。"],
        }
    try:
        structured = parse_video_recommendation_content(version.content)
    except M5CapabilityError as exc:
        return {
            "schema_version": "f5.video-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "source_refs": source_refs,
            "structured": None,
            "legacy": None,
            "invalid": {
                "raw_content": learner_safe_artifact_content(artifact.artifact_type, version.content),
                "error_code": exc.code,
            },
            "source_readiness": {
                "recorded": bool(source_refs), "available": False, "publishable": False,
                "reason": "视频推荐结构不完整，不能声称来源、许可或播放能力有效。",
            },
            "warnings": ["当前视频推荐无法解析，请恢复有效版本或重新生成该资源。"],
        }
    cited_ids = {
        ref_id for item in structured.recommendations for ref_id in item.sourceRefIds
    }
    readiness = _source_readiness(
        source_refs, cited_ids,
        ready="课程来源完整，视频推荐可用于学习。",
        missing="视频推荐引用了未登记课程来源，请修正来源定位。",
    )
    recommendations = structured.model_dump()["recommendations"]
    for recommendation in recommendations:
        recommendation["playback"] = {
            "mode": "local" if recommendation["sourceKind"] == "local_ref" else "external",
            "available": recommendation["availability"] == "playable",
            "external_only": recommendation["availability"] == "external_only",
            "reason": (
                "已登记为本地可播放素材。" if recommendation["availability"] == "playable"
                else "来源仅允许外链参考，请在原始页面观看。" if recommendation["availability"] == "external_only"
                else "当前链接不可用，请使用替代项或等待来源恢复。"
            ),
        }
    return {
        "schema_version": "f5.video-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "source_refs": source_refs,
        "structured": {**structured.model_dump(), "recommendations": recommendations},
        "legacy": None,
        "source_readiness": readiness,
        "warnings": [] if readiness["publishable"] else [readiness["reason"]],
    }


def _deck_projection(
    session: Any, artifact: DBArtifact, version: DBArtifactVersion,
) -> dict[str, Any]:
    protocol = _resource_protocol(session, artifact, version)
    source_refs = protocol["source_refs"]
    if version.content_format != "json":
        content = learner_safe_artifact_content(artifact.artifact_type, version.content)
        return {
            "schema_version": "f5.deck-projection.v1",
            "content_schema_version": None,
            "mode": "legacy_document",
            "source_refs": source_refs,
            "structured": None,
            "legacy": {"title": _markdown_title(content, artifact.title), "content": content},
            "source_readiness": {
                "recorded": bool(source_refs), "available": False, "publishable": False,
                "reason": "旧课件说明可继续阅读，但需要结构化页树后才能预览、编辑和导出。",
            },
            "warnings": ["这是旧版课件说明，系统不会补造缺失的幻灯片。"],
        }
    try:
        structured = parse_deck_content(version.content)
    except M5CapabilityError as exc:
        return {
            "schema_version": "f5.deck-projection.v1",
            "content_schema_version": None,
            "mode": "invalid_structured",
            "source_refs": source_refs,
            "structured": None,
            "legacy": None,
            "invalid": {
                "raw_content": learner_safe_artifact_content(artifact.artifact_type, version.content),
                "error_code": exc.code,
            },
            "source_readiness": {
                "recorded": bool(source_refs), "available": False, "publishable": False,
                "reason": "课件页树不完整，当前版本不能编辑或导出。",
            },
            "warnings": ["当前课件无法解析，请恢复有效版本；现有版本内容未被覆盖。"],
        }
    cited_ids = {ref_id for slide in structured.slides for ref_id in slide.sourceRefIds}
    readiness = _source_readiness(
        source_refs, cited_ids,
        ready="课程来源完整，课件可预览、编辑和导出。",
        missing="课件引用了未登记课程来源，请修正页面引用。",
    )
    payload = structured.model_dump()
    return {
        "schema_version": "f5.deck-projection.v1",
        "content_schema_version": structured.schemaVersion,
        "mode": "structured",
        "source_refs": source_refs,
        "structured": payload,
        "outline": [
            {"id": slide.id, "number": index, "title": slide.title, "layout": slide.layout}
            for index, slide in enumerate(structured.slides, start=1)
        ],
        "legacy": None,
        "source_readiness": readiness,
        "export_capabilities": {"html": True, "pptx": True},
        "warnings": [] if readiness["publishable"] else [readiness["reason"]],
    }


def search_resources(
    session: Any,
    *,
    owner_public_id: str,
    query: str = "",
    artifact_type: str | None = None,
    course_id: str | None = None,
    status: str | None = None,
    source_status: str | None = None,
    include_archived: bool = False,
    cursor: str | None = None,
    limit: int = 25,
) -> dict[str, Any]:
    limit = max(1, min(limit, 100))
    rows = session.query(DBArtifact).filter_by(owner_public_id=owner_public_id)
    if course_id:
        rows = rows.filter_by(course_id=course_id)
    if status:
        rows = rows.filter_by(status=status)
    if not include_archived:
        rows = rows.filter(DBArtifact.archived_at.is_(None))
    if cursor:
        rows = rows.filter(DBArtifact.id > cursor)
    candidates = rows.order_by(DBArtifact.id).all()
    query_text = query.strip().casefold()
    output: list[dict[str, Any]] = []
    type_counts = {resource_type: 0 for resource_type in RESOURCE_CATALOG_TYPES}
    for item in candidates:
        version = session.get(DBArtifactVersion, item.current_version_id) if item.current_version_id else None
        metadata = dict(version.metadata_json or {}) if version else {}
        canonical_type = canonical_resource_type(item.artifact_type, metadata)
        if canonical_type not in RESOURCE_CATALOG_TYPES:
            continue
        catalog_item = _catalog_item(session, item, version, canonical_type, metadata)
        searchable = " ".join([
            item.title,
            catalog_item["module"] or "",
            catalog_item["knowledge_point"]["title"] or "",
            *catalog_item["tags"],
        ]).casefold()
        if query_text and query_text not in searchable:
            continue
        if source_status == "available" and (
            catalog_item["source_count"] == 0 or catalog_item["source_unavailable_count"] > 0
        ):
            continue
        if source_status == "unavailable" and catalog_item["source_unavailable_count"] == 0:
            continue
        if source_status == "not_recorded" and catalog_item["source_count"] > 0:
            continue
        type_counts[canonical_type] += 1
        if artifact_type and canonical_type != artifact_type:
            continue
        output.append(catalog_item)
    has_more = len(output) > limit
    output = output[:limit]
    return {
        "items": output,
        "type_counts": type_counts,
        "types": [
            {"id": resource_type, "label": RESOURCE_TYPE_LABELS[resource_type], "count": type_counts[resource_type]}
            for resource_type in RESOURCE_CATALOG_TYPES
        ],
        "total": sum(type_counts.values()),
        "next_cursor": output[-1]["id"] if has_more and output else None,
        "has_more": has_more,
        "limit": limit,
    }


def _catalog_item(
    session: Any,
    item: DBArtifact,
    version: DBArtifactVersion | None,
    canonical_type: str,
    metadata: dict[str, Any],
) -> dict[str, Any]:
    protocol = _resource_protocol(session, item, version)
    course = session.get(DBCourse, item.course_id) if item.course_id else None
    knowledge_point_id = metadata.get("knowledge_point_id")
    knowledge_point = session.get(DBKnowledgePoint, knowledge_point_id) if knowledge_point_id else None
    source_refs = protocol["source_refs"]
    source_unavailable_count = sum(ref.get("available") is False for ref in source_refs)
    return {
            "id": item.id,
            "course_id": item.course_id,
            "artifact_type": item.artifact_type,
            "type": canonical_type,
            "type_label": RESOURCE_TYPE_LABELS[canonical_type],
            "title": item.title,
            "status": item.status,
            "current_version_id": item.current_version_id,
            "lock_version": item.lock_version,
            "archived": item.archived_at is not None,
            "updated_at": item.updated_at.isoformat() if item.updated_at else None,
            "route": resource_route(canonical_type, item.id),
            "course": {"id": item.course_id, "title": course.title if course else "未绑定课程"},
            "knowledge_point": {
                "id": knowledge_point_id,
                "title": knowledge_point.title if knowledge_point else str(metadata.get("knowledge_point_title") or ""),
            },
            "module": str(metadata.get("module") or metadata.get("chapter_title") or ""),
            "tags": _metadata_values(metadata, "tags"),
            "version_number": version.version_number if version else None,
            "source_count": len(source_refs),
            "source_unavailable_count": source_unavailable_count,
            "generation_state": protocol["generation_state"],
            "completion": {
                "status": str(metadata.get("completion_status") or "not_recorded"),
                "progress": metadata.get("completion_progress"),
                "rule": str(metadata.get("completion_rule") or ""),
            },
            "preview": {
                "text": _preview_text(
                    learner_safe_artifact_content(item.artifact_type, version.content)
                ) if version else "",
                "content_format": version.content_format if version else None,
            },
            "resource": protocol,
    }


def _resource_protocol(session: Any, artifact: DBArtifact, version: DBArtifactVersion | None) -> dict[str, Any]:
    metadata = dict(version.metadata_json or {}) if version else {}
    semantic_context = metadata.get("semantic_context")
    semantic_context = semantic_context if isinstance(semantic_context, dict) else {}
    source_refs = []
    for source_ref_id in metadata.get("source_ref_ids") or []:
        source_ref = session.get(DBSourceRef, source_ref_id)
        source = session.get(DBSourceDocument, source_ref.source_document_id) if source_ref else None
        if (
            source_ref is None
            or source is None
            or (artifact.course_id is not None and source.course_id != artifact.course_id)
        ):
            source_refs.append({"source_ref_id": source_ref_id, "available": False})
            continue
        source_refs.append({
            "source_ref_id": source_ref.id,
            "source_document_id": source.id,
            "title": source.title,
            "location_kind": source_ref.location_kind,
            "location_label": source_ref.location_label,
            "parser_version": source_ref.parser_version,
            "content_hash": source_ref.content_hash,
            "source_url": source.source_url,
            "license": source.license_text,
            "available": source.processing_status == "ready",
        })
    task = session.get(DBGenerationTask, artifact.creating_task_id) if artifact.creating_task_id else None
    step = (
        session.query(DBTaskStep)
        .filter_by(output_artifact_version_id=version.id)
        .order_by(DBTaskStep.created_at.desc())
        .first()
        if version else None
    )
    sources_available = bool(source_refs) and all(item["available"] for item in source_refs)
    return {
        "schema_version": "f2.resource.v1",
        "resource_id": artifact.id,
        "type": metadata.get("resource_type") or artifact.artifact_type,
        "course_id": artifact.course_id,
        "knowledge_point_id": metadata.get("knowledge_point_id"),
        "source_refs": source_refs,
        "profile_snapshot_id": metadata.get("profile_snapshot_id"),
        # The generated version already contains the public adaptation plan.
        # Expose it through the canonical resource protocol so every detail
        # workspace can render the same learner-facing explanation, including
        # resource types whose content schema predates profile fields.
        "profile_adaptation": semantic_context.get("profile_adaptation"),
        "profile_adaptation_plan": semantic_context.get("profile_adaptation_plan") or semantic_context.get("profile_context") or {},
        "profile_context": semantic_context.get("profile_context") or semantic_context.get("profile_adaptation_plan") or {},
        "mastery_summary": semantic_context.get("mastery_summary"),
        "recent_learning_summary": semantic_context.get("recent_learning_summary"),
        "version": {
            "id": version.id,
            "number": version.version_number,
            "content_hash": version.content_hash,
        } if version else None,
        "status": (
            "archived" if artifact.archived_at is not None
            else "available" if sources_available
            else "source_unavailable"
        ),
        "generation_state": (
            step.status if step
            else "succeeded" if task and task.status == "completed"
            else task.status if task
            else "legacy"
        ),
        "content_ref": f"artifact-version:{version.id}" if version else None,
    }


def _resource_history(session: Any, artifact: DBArtifact) -> dict[str, Any]:
    versions = (
        session.query(DBArtifactVersion)
        .filter_by(artifact_id=artifact.id)
        .order_by(DBArtifactVersion.version_number.desc())
        .all()
    )
    items = []
    for version in versions:
        metadata = dict(version.metadata_json or {})
        protocol = _resource_protocol(session, artifact, version)
        items.append({
            "id": version.id,
            "number": version.version_number,
            "is_current": version.id == artifact.current_version_id,
            "content_format": version.content_format,
            "content_hash": version.content_hash,
            "created_at": version.created_at.isoformat() if version.created_at else None,
            "created_by_public_id": version.created_by_public_id,
            "created_task_id": version.created_task_id,
            "revision_origin": metadata.get("revision_origin") or "generated",
            "revised_from_version_id": metadata.get("revised_from_version_id"),
            "restored_from_version_id": metadata.get("restored_from_version_id"),
            "restored_from_version_number": metadata.get("restored_from_version_number"),
            "source_refs": protocol["source_refs"],
            "source_status": protocol["status"],
        })
    return {
        "schema_version": "f5.resource-history.v1",
        "current_version_id": artifact.current_version_id,
        "items": items,
    }


def resource_detail(session: Any, artifact: DBArtifact) -> dict[str, Any]:
    payload = M1Repository(session).artifact_payload(artifact)
    payload["lock_version"] = artifact.lock_version
    payload["archived"] = artifact.archived_at is not None
    current = session.get(DBArtifactVersion, artifact.current_version_id) if artifact.current_version_id else None
    metadata = dict(current.metadata_json or {}) if current else {}
    canonical_type = canonical_resource_type(artifact.artifact_type, metadata)
    payload["resource"] = _resource_protocol(session, artifact, current)
    payload["history"] = _resource_history(session, artifact)
    if canonical_type in RESOURCE_CATALOG_TYPES:
        payload["catalog"] = _catalog_item(session, artifact, current, canonical_type, metadata)
    if canonical_type == "handout" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["handout"] = {
            "schema_version": "f5.handout-workspace.v1",
            "current_version_id": current.id,
            "current": _handout_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "projection": _handout_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "mind_map" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["mindmap"] = {
            "schema_version": "f5.mindmap-workspace.v1",
            "current_version_id": current.id,
            "current": _mindmap_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _mindmap_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "quiz" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["quiz"] = {
            "schema_version": "f6.quiz-workspace.v1",
            "current_version_id": current.id,
            "current": _quiz_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _quiz_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "training_task" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["training_task"] = {
            "schema_version": "f5.training-task-workspace.v1",
            "current_version_id": current.id,
            "current": _training_task_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _training_task_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "video" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["video"] = {
            "schema_version": "f5.video-workspace.v1",
            "current_version_id": current.id,
            "current": _video_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _video_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "deck" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["deck"] = {
            "schema_version": "f5.deck-workspace.v1",
            "current_version_id": current.id,
            "current": _deck_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _deck_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "code_sandbox" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["code_sandbox"] = {
            "schema_version": "f6.code-sandbox-workspace.v1",
            "current_version_id": current.id,
            "current": _code_sandbox_projection(session, artifact, current),
            "completion": _code_sandbox_completion(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _code_sandbox_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    if canonical_type == "interactive_model" and current is not None:
        versions = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id).order_by(
            DBArtifactVersion.version_number
        ).all()
        payload["interactive_model"] = {
            "schema_version": "f6.interactive-model-workspace.v1",
            "current_version_id": current.id,
            "current": _interactive_model_projection(session, artifact, current),
            "versions": [
                {
                    "id": version.id,
                    "number": version.version_number,
                    "content_format": version.content_format,
                    "content_hash": version.content_hash,
                    "projection": _interactive_model_projection(session, artifact, version),
                }
                for version in versions
            ],
        }
    return payload


def create_resource_version(
    session: Any,
    *,
    artifact: DBArtifact,
    actor_public_id: str,
    request: ResourceVersionCreateRequest,
) -> dict[str, Any]:
    if artifact.archived_at is not None:
        raise M5CapabilityError("resource_archived", "archived resource cannot create a new version")
    if artifact.lock_version != request.expected_lock_version:
        raise M5CapabilityError("resource_version_conflict", "resource lock version changed")
    content = request.content.strip()
    if not content:
        raise M5CapabilityError("resource_content_required", "resource version content cannot be empty")
    if len(content) > 200_000:
        raise M5CapabilityError("resource_content_too_large", "resource version content exceeds the editing limit")
    source_version = session.query(DBArtifactVersion).filter_by(
        id=request.source_version_id,
        artifact_id=artifact.id,
    ).first()
    if source_version is None:
        raise M5CapabilityError("resource_version_unknown", "resource source version does not exist")
    source_metadata = dict(source_version.metadata_json or {})
    source_type = canonical_resource_type(artifact.artifact_type, source_metadata)
    if source_type == "handout":
        if request.content_format == "json":
            parse_handout_content(content)
        elif source_version.content_format == "json":
            raise M5CapabilityError(
                "handout_content_format_required",
                "structured handout revisions must remain structured JSON",
            )
    if source_type == "mind_map":
        if request.content_format == "json":
            parse_mindmap_content(content)
        elif source_version.content_format == "json":
            raise M5CapabilityError(
                "mindmap_content_format_required",
                "structured mind map revisions must remain structured JSON",
            )
    if source_type == "training_task":
        if request.content_format == "json":
            parse_training_task_content(content)
        elif source_version.content_format == "json":
            raise M5CapabilityError(
                "training_task_content_format_required",
                "structured training task revisions must remain structured JSON",
            )
    if source_type == "video":
        if request.content_format != "json":
            raise M5CapabilityError(
                "video_content_format_required",
                "video recommendation revisions must use structured JSON",
            )
        parse_video_recommendation_content(content)
    if source_type == "deck":
        if request.content_format != "json":
            raise M5CapabilityError("deck_content_format_required", "deck revisions must use structured JSON")
        parse_deck_content(content)
    if source_type == "quiz":
        if request.content_format != "json":
            raise M5CapabilityError("quiz_content_format_required", "quiz revisions must use structured JSON")
        quiz_content = parse_quiz_content(content)
        for reference in quiz_content.questions:
            question_version = session.get(DBAssessmentQuestionVersion, reference.questionVersionId)
            question = session.get(DBAssessmentQuestion, question_version.question_id) if question_version else None
            if question_version is None or question is None or question.course_id != artifact.course_id:
                raise M5CapabilityError("quiz_question_version_invalid", "quiz question versions must belong to the resource course")
    if source_type == "code_sandbox":
        if request.content_format != "json":
            raise M5CapabilityError(
                "code_sandbox_content_format_required",
                "code sandbox revisions must use structured JSON",
            )
        structured = parse_code_sandbox_content(content)
        expected_rule = " ".join(str(source_metadata.get("completion_rule") or "").split())
        if not expected_rule or " ".join(structured.completionRule.split()) != expected_rule:
            raise M5CapabilityError(
                "code_sandbox_completion_rule_mismatch",
                "code sandbox revision must preserve the server-recorded completion rule",
            )
    if source_type == "interactive_model":
        if request.content_format != "json":
            raise M5CapabilityError(
                "interactive_model_content_format_required",
                "interactive model revisions must use structured JSON",
            )
        parse_interactive_model_content(content)
    metadata = dict(source_metadata)
    metadata.update({
        "revision_origin": "learning_workspace",
        "revised_from_version_id": source_version.id,
    })
    version = add_artifact_version(
        session,
        artifact=artifact,
        content=content,
        content_format=request.content_format,
        created_by_public_id=actor_public_id,
        metadata=metadata,
    )
    artifact.status = "draft"
    artifact.lock_version += 1
    record_audit(
        session,
        actor_public_id=actor_public_id,
        action="artifact.version.create",
        object_type="artifact_version",
        object_id=version.id,
        metadata={
            "artifact_id": artifact.id,
            "version_number": version.version_number,
            "source_version_id": source_version.id,
            "content_format": request.content_format,
        },
    )
    session.commit()
    return resource_detail(session, artifact)


def compare_resource_versions(
    session: Any,
    *,
    artifact: DBArtifact,
    left_version: int,
    right_version: int,
) -> dict[str, Any]:
    versions = session.query(DBArtifactVersion).filter(
        DBArtifactVersion.artifact_id == artifact.id,
        DBArtifactVersion.version_number.in_((left_version, right_version)),
    ).all()
    by_number = {item.version_number: item for item in versions}
    if left_version not in by_number or right_version not in by_number:
        raise M5CapabilityError("resource_version_unknown", "requested artifact version does not exist")
    left = by_number[left_version]
    right = by_number[right_version]
    diff = "\n".join(difflib.unified_diff(
        left.content.splitlines(), right.content.splitlines(),
        fromfile=f"v{left_version}", tofile=f"v{right_version}", lineterm="",
    ))
    return {
        "artifact_id": artifact.id,
        "left": {"version": left_version, "content_hash": left.content_hash},
        "right": {"version": right_version, "content_hash": right.content_hash},
        "changed": left.content_hash != right.content_hash,
        "diff": diff[:200000],
        "diff_truncated": len(diff) > 200000,
    }


def restore_resource_version(
    session: Any,
    *,
    artifact: DBArtifact,
    actor_public_id: str,
    request: ResourceVersionRestoreRequest,
) -> dict[str, Any]:
    if artifact.archived_at is not None:
        raise M5CapabilityError("resource_archived", "archived resource cannot restore a version")
    if artifact.lock_version != request.expected_lock_version:
        raise M5CapabilityError("resource_version_conflict", "resource lock version changed")
    version = session.query(DBArtifactVersion).filter_by(
        id=request.version_id,
        artifact_id=artifact.id,
    ).first()
    if version is None:
        raise M5CapabilityError("resource_version_unknown", "requested artifact version does not exist")
    if version.id == artifact.current_version_id:
        raise M5CapabilityError("resource_version_already_current", "requested artifact version is already current")
    previous_current_version_id = artifact.current_version_id
    metadata = dict(version.metadata_json or {})
    metadata.update({
        "revision_origin": "history_restore",
        "restored_from_version_id": version.id,
        "restored_from_version_number": version.version_number,
        "previous_current_version_id": previous_current_version_id,
    })
    restored = add_artifact_version(
        session,
        artifact=artifact,
        content=version.content,
        content_format=version.content_format,
        created_by_public_id=actor_public_id,
        metadata=metadata,
    )
    artifact.status = "draft"
    artifact.lock_version += 1
    record_audit(
        session,
        actor_public_id=actor_public_id,
        action="artifact.version.restore",
        object_type="artifact_version",
        object_id=restored.id,
        metadata={
            "artifact_id": artifact.id,
            "restored_from_version_id": version.id,
            "restored_from_version_number": version.version_number,
            "previous_current_version_id": previous_current_version_id,
        },
    )
    session.commit()
    return resource_detail(session, artifact)


def validate_resource_completion(
    session: Any,
    *,
    artifact: DBArtifact,
    resource_version_id: str,
    completion_rule: str,
) -> dict[str, Any]:
    if artifact.archived_at is not None:
        raise M5CapabilityError("resource_archived", "archived resource cannot be marked complete")
    version = session.query(DBArtifactVersion).filter_by(
        id=resource_version_id,
        artifact_id=artifact.id,
    ).first()
    if version is None:
        raise M5CapabilityError("resource_version_unknown", "resource completion version does not exist")
    if version.id != artifact.current_version_id:
        raise M5CapabilityError("resource_version_stale", "resource completion must reference the current version")
    metadata = dict(version.metadata_json or {})
    expected_rule = " ".join(str(metadata.get("completion_rule") or "").split())
    submitted_rule = " ".join(str(completion_rule or "").split())
    if not expected_rule:
        raise M5CapabilityError(
            "resource_completion_rule_unavailable",
            "resource has no server-recorded completion rule",
        )
    if submitted_rule != expected_rule:
        raise M5CapabilityError(
            "resource_completion_rule_mismatch",
            "resource completion rule does not match the current version",
        )
    protocol = _resource_protocol(session, artifact, version)
    if protocol["status"] != "available":
        raise M5CapabilityError(
            "resource_source_unavailable",
            "resource sources are unavailable and completion cannot be verified",
        )
    return {
        "resource_version_id": version.id,
        "completion_rule": expected_rule,
        "resource_type": canonical_resource_type(artifact.artifact_type, metadata),
        "knowledge_point_id": metadata.get("knowledge_point_id"),
        "generation_job_id": metadata.get("generation_job_id"),
    }


def export_resource_version(
    session: Any,
    *,
    artifact: DBArtifact,
    version_number: int | None = None,
) -> dict[str, Any]:
    query = session.query(DBArtifactVersion).filter_by(artifact_id=artifact.id)
    version = (
        query.filter_by(version_number=version_number).first()
        if version_number is not None
        else query.filter_by(id=artifact.current_version_id).first()
    )
    if version is None:
        raise M5CapabilityError("resource_version_unknown", "artifact has no exportable version")
    actual_hash = hashlib.sha256(version.content.encode("utf-8")).hexdigest()
    if actual_hash != version.content_hash:
        raise M5CapabilityError("resource_hash_mismatch", "artifact content differs from its persisted hash")
    export_content = learner_safe_artifact_content(artifact.artifact_type, version.content)
    export_hash = hashlib.sha256(export_content.encode("utf-8")).hexdigest()
    return {
        "schema_version": "m5.resource-export.v1",
        "artifact": {
            "id": artifact.id,
            "type": artifact.artifact_type,
            "title": artifact.title,
            "course_id": artifact.course_id,
            "status": artifact.status,
        },
        "version": {
            "id": version.id,
            "number": version.version_number,
            "content_format": version.content_format,
            "persisted_content_hash": version.content_hash,
            "export_content_hash": export_hash,
            "metadata": version.metadata_json or {},
            "content": export_content,
        },
    }


def _deck_export_version(
    session: Any, *, artifact: DBArtifact, version_id: str | None,
) -> tuple[DBArtifactVersion, DeckContent]:
    version = session.query(DBArtifactVersion).filter_by(
        id=version_id or artifact.current_version_id,
        artifact_id=artifact.id,
    ).first()
    if version is None:
        raise M5CapabilityError("resource_version_unknown", "deck export version does not exist")
    metadata = dict(version.metadata_json or {})
    if canonical_resource_type(artifact.artifact_type, metadata) != "deck":
        raise M5CapabilityError("resource_type_mismatch", "resource is not a deck")
    if version.content_format != "json":
        raise M5CapabilityError("deck_export_unavailable", "legacy deck content cannot be exported as slides")
    return version, parse_deck_content(version.content)


def _safe_export_name(value: str, fallback: str) -> str:
    clean = re.sub(r'[\\/:*?"<>|]+', "-", value).strip(" .-")
    return clean[:80] or fallback


def render_deck_export(
    session: Any,
    *,
    artifact: DBArtifact,
    version_id: str | None,
    format_name: Literal["html", "pptx"],
) -> tuple[bytes, str, str]:
    version, deck = _deck_export_version(session, artifact=artifact, version_id=version_id)
    filename = f"{_safe_export_name(deck.title, 'EduMatrix-courseware')}-v{version.version_number}.{format_name}"
    if format_name == "html":
        sections = []
        for index, slide in enumerate(deck.slides, start=1):
            bullets = "".join(f"<li>{html.escape(item)}</li>" for item in slide.bullets)
            sources = " ".join(f"<code>{html.escape(item)}</code>" for item in slide.sourceRefIds)
            sections.append(
                f'<section class="slide slide--{html.escape(slide.layout)}" data-slide-id="{html.escape(slide.id)}">'
                f'<p class="number">{index:02d}</p><div class="content"><h1>{html.escape(slide.title)}</h1>'
                f'<h2>{html.escape(slide.subtitle)}</h2><p class="body">{html.escape(slide.body)}</p>'
                f'<ul>{bullets}</ul></div><footer><span>{sources}</span><span>EduMatrix · v{version.version_number}</span></footer>'
                f'<aside class="speaker-notes">{html.escape(slide.speakerNotes)}</aside></section>'
            )
        document = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(deck.title)}</title><style>
:root{{--page:#f5f7f7;--surface:#fff;--ink:#06141b;--primary:#11212d;--muted:#60727c;--line:#ccd0cf}}*{{box-sizing:border-box}}body{{margin:0;background:var(--page);color:var(--ink);font-family:"Source Han Sans SC","Noto Sans CJK SC","Microsoft YaHei",sans-serif;letter-spacing:0}}.slide{{width:min(94vw,160vh);aspect-ratio:16/9;margin:3vh auto;padding:6.5%;background:var(--surface);border:1px solid var(--line);position:relative;overflow:hidden;break-after:page}}.content{{max-width:88%}}h1{{margin:0 0 2%;font-size:clamp(30px,4vw,58px);line-height:1.18}}h2{{margin:0 0 4%;font-size:clamp(16px,1.8vw,26px);font-weight:500;color:var(--muted)}}.body,li{{font-size:clamp(14px,1.55vw,22px);line-height:1.55}}ul{{padding-left:1.2em}}.number{{position:absolute;right:5%;top:4%;font-size:14px;color:var(--muted)}}footer{{position:absolute;left:6.5%;right:6.5%;bottom:4%;display:flex;justify-content:space-between;gap:24px;color:var(--muted);font-size:14px}}code{{font-family:inherit}}.speaker-notes{{display:none}}@media print{{body{{background:#fff}}.slide{{width:100vw;height:100vh;margin:0;border:0}}}}</style></head><body>{''.join(sections)}</body></html>"""
        return document.encode("utf-8"), filename, "text/html; charset=utf-8"

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for index, item in enumerate(deck.slides, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        title = slide.shapes.add_textbox(Inches(0.82), Inches(0.72), Inches(11.4), Inches(0.82))
        paragraph = title.text_frame.paragraphs[0]
        paragraph.text = item.title
        paragraph.font.name = "Source Han Sans SC"
        paragraph.font.size = Pt(30)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(6, 20, 27)
        if item.subtitle:
            subtitle = slide.shapes.add_textbox(Inches(0.85), Inches(1.62), Inches(11.1), Inches(0.52))
            paragraph = subtitle.text_frame.paragraphs[0]
            paragraph.text = item.subtitle
            paragraph.font.name = "Source Han Sans SC"
            paragraph.font.size = Pt(16)
            paragraph.font.color.rgb = RGBColor(96, 114, 124)
        body = slide.shapes.add_textbox(Inches(0.9), Inches(2.35), Inches(11.45), Inches(3.75))
        frame = body.text_frame
        frame.clear()
        frame.word_wrap = True
        lines = ([item.body] if item.body else []) + [f"•  {bullet}" for bullet in item.bullets]
        for line_index, line in enumerate(lines or [""]):
            paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Source Han Sans SC"
            paragraph.font.size = Pt(19)
            paragraph.font.color.rgb = RGBColor(17, 33, 45)
            paragraph.space_after = Pt(10)
        footer = slide.shapes.add_textbox(Inches(0.88), Inches(6.72), Inches(11.55), Inches(0.3))
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = f"EduMatrix · v{version.version_number} · {index:02d}"
        paragraph.alignment = PP_ALIGN.RIGHT
        paragraph.font.name = "Source Han Sans SC"
        paragraph.font.size = Pt(10)
        paragraph.font.color.rgb = RGBColor(96, 114, 124)
        notes = slide.notes_slide.notes_text_frame
        source_note = f"\n来源：{', '.join(item.sourceRefIds)}" if item.sourceRefIds else ""
        notes.text = f"{item.speakerNotes}{source_note}".strip()
    buffer = io.BytesIO()
    presentation.save(buffer)
    return (
        buffer.getvalue(),
        filename,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


def locate_video_media(
    session: Any,
    *,
    artifact: DBArtifact,
    recommendation_id: str,
    version_id: str | None,
) -> tuple[Path, str]:
    version = session.query(DBArtifactVersion).filter_by(
        id=version_id or artifact.current_version_id,
        artifact_id=artifact.id,
    ).first()
    if version is None:
        raise M5CapabilityError("resource_version_unknown", "video resource version does not exist")
    metadata = dict(version.metadata_json or {})
    if canonical_resource_type(artifact.artifact_type, metadata) != "video":
        raise M5CapabilityError("resource_type_mismatch", "resource is not a video recommendation")
    if version.content_format != "json":
        raise M5CapabilityError("video_media_unavailable", "legacy video resource has no playable media")
    content = parse_video_recommendation_content(version.content)
    recommendation = next(
        (item for item in content.recommendations if item.id == recommendation_id), None,
    )
    if (
        recommendation is None
        or recommendation.sourceKind != "local_ref"
        or recommendation.availability != "playable"
    ):
        raise M5CapabilityError("video_media_unavailable", "recommendation has no packaged playable media")
    expected_ref = f"/api/v1/m5/resources/{artifact.id}/video-media/{recommendation.id}"
    if recommendation.localRef != expected_ref:
        raise M5CapabilityError("video_media_reference_invalid", "video media reference does not match its resource")
    files = metadata.get("local_video_files") or {}
    relative = str(files.get(recommendation.id) or "").replace("\\", "/").strip("/")
    if not relative:
        raise M5CapabilityError("video_media_unavailable", "packaged video file is not registered")
    data_root = (Path(__file__).resolve().parents[1] / "data").resolve()
    path = (data_root / relative).resolve()
    try:
        path.relative_to(data_root)
    except ValueError as exc:
        raise M5CapabilityError("video_media_path_invalid", "video file must stay inside the data directory") from exc
    if not path.is_file() or path.suffix.lower() not in {".mp4", ".webm", ".ogg"}:
        raise M5CapabilityError("video_media_unavailable", "registered video file is missing or unsupported")
    media_type = {
        ".mp4": "video/mp4", ".webm": "video/webm", ".ogg": "video/ogg",
    }[path.suffix.lower()]
    return path, media_type


def set_resource_archived(
    session: Any,
    *,
    artifact: DBArtifact,
    actor_public_id: str,
    archived: bool,
    expected_lock_version: int,
) -> dict[str, Any]:
    if artifact.lock_version != expected_lock_version:
        raise M5CapabilityError("resource_version_conflict", "resource lock version changed")
    if (artifact.archived_at is not None) == archived:
        return resource_detail(session, artifact)
    set_artifact_archived(
        session, artifact=artifact, actor_public_id=actor_public_id, archived=archived,
    )
    session.commit()
    return resource_detail(session, artifact)
