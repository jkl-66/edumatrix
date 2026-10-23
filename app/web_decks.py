"""Versioned browser slide decks with HTML, PPTX and optional PDF exports."""

from __future__ import annotations

import hashlib
import html
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.authorization import ObjectAccess, active_course_role, can_access_object
from app.database import DBCourse, DBUser, DBWebDeck, DBWebDeckVersion, DBWebSlide


THEME_ID = "evidence-light-v1"
THEME = {
    "id": THEME_ID,
    "name": "证据浅色",
    "background": "#F5F7F4",
    "surface": "#FFFFFF",
    "ink": "#1E2822",
    "muted": "#637067",
    "accent": "#355B42",
    "highlight": "#C65A3C",
    "font_family": "Microsoft YaHei",
    "aspect_ratio": "16:9",
}
LAYOUTS = {"title", "content", "two_column", "closing"}
MAX_SLIDES = 30


class WebDeckError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _fail(code: str, message: str) -> None:
    raise WebDeckError(code, message)


def _clean(value: Any, limit: int) -> str:
    return " ".join(str(value or "").split())[:limit]


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def capabilities() -> dict[str, Any]:
    executable = _libreoffice()
    return {
        "themes": [THEME],
        "layouts": sorted(LAYOUTS),
        "exports": {
            "html": {"available": True, "required": False, "engine": "edumatrix-html-v1"},
            "pptx": {"available": True, "required": True, "engine": "python-pptx"},
            "pdf": {
                "available": executable is not None,
                "required": False,
                "engine": "LibreOffice" if executable else "unavailable",
                "boundary": "PDF 为可选导出；不可用或转换失败时 PPTX 与 HTML 保持可下载。",
            },
        },
        "limits": {"minimum_slides": 1, "maximum_slides": MAX_SLIDES},
    }


def _libreoffice() -> str | None:
    configured = os.getenv("EDUMATRIX_LIBREOFFICE_PATH", "").strip()
    candidates = [
        configured,
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        shutil.which("soffice") or "",
    ]
    return next((item for item in candidates if item and Path(item).is_file()), None)


def _course(db: Any, actor: DBUser, course_id: str | None) -> DBCourse | None:
    if not course_id:
        return None
    row = db.get(DBCourse, course_id)
    if row is None:
        _fail("web_deck.course_unknown", "课程不存在")
    from course_catalog import course_is_selectable
    if not course_is_selectable(db, row):
        _fail("web_deck.course_unavailable", "课程已停止发布或没有可用学习资料")
    role = active_course_role(db, course_id=row.id, user_public_id=actor.public_id)
    allowed = can_access_object(
        actor,
        ObjectAccess("course", row.id, owner_public_id=row.owner_public_id, course_id=row.id, visibility=row.visibility),
        action="read",
        course_role=role,
    )
    if not allowed:
        _fail("web_deck.course_forbidden", "无权将演示文稿关联到该课程")
    return row


def _owned(db: Any, actor: DBUser, deck_id: str) -> DBWebDeck:
    row = db.get(DBWebDeck, deck_id)
    if row is None:
        _fail("web_deck.unknown", "演示文稿不存在")
    if row.owner_public_id != actor.public_id:
        _fail("web_deck.forbidden", "无权访问该演示文稿")
    return row


def _normalize_slides(db: Any, actor: DBUser, deck: DBWebDeck, slides: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not 1 <= len(slides) <= MAX_SLIDES:
        _fail("web_deck.slide_count_invalid", f"演示文稿必须包含 1 到 {MAX_SLIDES} 页")
    result = []
    seen: set[str] = set()
    for index, raw in enumerate(slides):
        if not isinstance(raw, dict):
            _fail("web_deck.slide_invalid", f"第 {index + 1} 页必须是结构化对象")
        slide_id = _clean(raw.get("id"), 64)
        if slide_id:
            identity = db.get(DBWebSlide, slide_id)
            if identity is None or identity.deck_id != deck.id:
                _fail("web_deck.slide_forbidden", "幻灯片身份不属于当前演示文稿")
        else:
            stable_key = f"slide-{db.query(DBWebSlide).filter_by(deck_id=deck.id).count() + index + 1:03d}"
            identity = DBWebSlide(deck_id=deck.id, stable_key=stable_key, created_by_public_id=actor.public_id)
            db.add(identity)
            db.flush()
            slide_id = identity.id
        if slide_id in seen:
            _fail("web_deck.slide_duplicate", "同一版本不能重复引用同一幻灯片")
        seen.add(slide_id)
        layout = _clean(raw.get("layout") or "content", 32)
        if layout not in LAYOUTS:
            _fail("web_deck.layout_invalid", f"第 {index + 1} 页布局不受支持")
        title = _clean(raw.get("title"), 160)
        if not title:
            _fail("web_deck.slide_title_required", f"第 {index + 1} 页标题不能为空")
        bullets = raw.get("bullets") or []
        if not isinstance(bullets, list) or len(bullets) > 8:
            _fail("web_deck.bullets_invalid", f"第 {index + 1} 页要点不能超过 8 条")
        result.append({
            "id": slide_id,
            "order": index,
            "layout": layout,
            "title": title,
            "subtitle": _clean(raw.get("subtitle"), 240),
            "body": _clean(raw.get("body"), 1200),
            "bullets": [_clean(item, 240) for item in bullets if _clean(item, 240)],
            "speaker_notes": _clean(raw.get("speaker_notes"), 2000),
            "accent": _clean(raw.get("accent") or THEME["accent"], 16),
        })
    return result


def _version_payload(row: DBWebDeckVersion, *, include_slides: bool = True) -> dict[str, Any]:
    return {
        "id": row.id,
        "version_number": row.version_number,
        "slides": row.slides_snapshot if include_slides else [],
        "slide_count": len(row.slides_snapshot or []),
        "theme": row.theme_snapshot,
        "source_version_id": row.source_version_id,
        "change_summary": row.change_summary,
        "content_hash": row.content_hash,
        "created_at": row.created_at,
    }


def _deck_payload(db: Any, row: DBWebDeck, *, include_versions: bool = True) -> dict[str, Any]:
    current = db.get(DBWebDeckVersion, row.current_version_id) if row.current_version_id else None
    versions = (
        db.query(DBWebDeckVersion).filter_by(deck_id=row.id)
        .order_by(DBWebDeckVersion.version_number.desc()).all()
    ) if include_versions else []
    max_slide = max(0, len(current.slides_snapshot or []) - 1) if current else 0
    return {
        "id": row.id, "title": row.title, "theme_id": row.theme_id, "course_id": row.course_id,
        "current_version_id": row.current_version_id,
        "current_slide": min(row.current_slide, max_slide),
        "max_slide": max_slide, "playback_status": row.playback_status,
        "status": row.status, "lock_version": row.lock_version,
        "current_version": _version_payload(current) if current else None,
        "versions": [_version_payload(item, include_slides=False) for item in versions],
        "export_capabilities": capabilities()["exports"],
        "created_at": row.created_at, "updated_at": row.updated_at,
        "boundary": "内容和讲者备注保存在不可变版本中；PDF 是 LibreOffice 可用时的可选导出。",
    }


def list_decks(db: Any, *, actor: DBUser) -> dict[str, Any]:
    rows = db.query(DBWebDeck).filter_by(owner_public_id=actor.public_id, status="active").order_by(DBWebDeck.updated_at.desc()).all()
    return {"items": [_deck_payload(db, row, include_versions=False) for row in rows], "count": len(rows)}


def create_deck(db: Any, *, actor: DBUser, title: str, course_id: str | None, slides: list[dict[str, Any]]) -> dict[str, Any]:
    clean_title = _clean(title, 256)
    if not clean_title:
        _fail("web_deck.title_required", "演示文稿标题不能为空")
    _course(db, actor, course_id)
    row = DBWebDeck(owner_public_id=actor.public_id, course_id=course_id, title=clean_title, theme_id=THEME_ID)
    db.add(row)
    db.flush()
    result = create_version(db, actor=actor, deck_id=row.id, slides=slides, expected_lock_version=1, change_summary="创建首版")
    result["created"] = True
    return result


def create_version(
    db: Any, *, actor: DBUser, deck_id: str, slides: list[dict[str, Any]],
    expected_lock_version: int, change_summary: str,
) -> dict[str, Any]:
    row = _owned(db, actor, deck_id)
    if row.lock_version != expected_lock_version:
        _fail("web_deck.conflict", "演示文稿已在其他页面更新，请刷新后重试")
    normalized = _normalize_slides(db, actor, row, slides)
    identity = {"theme": THEME, "slides": normalized}
    content_hash = hashlib.sha256(_canonical(identity).encode("utf-8")).hexdigest()
    existing = db.query(DBWebDeckVersion).filter_by(deck_id=row.id, content_hash=content_hash).one_or_none()
    if existing is not None:
        return {**_deck_payload(db, row), "version_created": False}
    number = db.query(DBWebDeckVersion).filter_by(deck_id=row.id).count() + 1
    version = DBWebDeckVersion(
        deck_id=row.id, version_number=number, slides_snapshot=normalized,
        theme_snapshot=THEME, change_summary=_clean(change_summary, 256) or f"保存 v{number}",
        content_hash=content_hash, created_by_public_id=actor.public_id,
    )
    db.add(version)
    db.flush()
    row.current_version_id = version.id
    row.current_slide = 0
    row.playback_status = "stopped"
    row.lock_version += 1
    db.flush()
    return {**_deck_payload(db, row), "version_created": True}


def get_deck(db: Any, *, actor: DBUser, deck_id: str) -> dict[str, Any]:
    return _deck_payload(db, _owned(db, actor, deck_id))


def restore_version(db: Any, *, actor: DBUser, deck_id: str, version_id: str, expected_lock_version: int) -> dict[str, Any]:
    row = _owned(db, actor, deck_id)
    if row.lock_version != expected_lock_version:
        _fail("web_deck.conflict", "演示文稿已在其他页面更新，请刷新后重试")
    version = db.get(DBWebDeckVersion, version_id)
    if version is None or version.deck_id != row.id:
        _fail("web_deck.version_unknown", "历史版本不存在")
    row.current_version_id = version.id
    row.current_slide = 0
    row.playback_status = "stopped"
    row.lock_version += 1
    db.flush()
    return _deck_payload(db, row)


def update_playback(
    db: Any, *, actor: DBUser, deck_id: str, current_slide: int,
    playback_status: str, expected_lock_version: int,
) -> dict[str, Any]:
    row = _owned(db, actor, deck_id)
    if row.lock_version != expected_lock_version:
        _fail("web_deck.conflict", "播放位置已变化，请刷新后重试")
    if playback_status not in {"playing", "stopped"}:
        _fail("web_deck.playback_invalid", "播放状态只能是 playing 或 stopped")
    version = db.get(DBWebDeckVersion, row.current_version_id)
    if version is None:
        _fail("web_deck.version_missing", "演示文稿没有可播放版本")
    max_slide = len(version.slides_snapshot) - 1
    if current_slide < 0 or current_slide > max_slide:
        _fail("web_deck.slide_out_of_range", f"当前页必须在 0 到 {max_slide} 之间")
    row.current_slide = current_slide
    row.playback_status = playback_status
    row.lock_version += 1
    db.flush()
    return _deck_payload(db, row)


def _export_version(db: Any, actor: DBUser, deck_id: str, version_id: str | None) -> tuple[DBWebDeck, DBWebDeckVersion]:
    deck = _owned(db, actor, deck_id)
    version = db.get(DBWebDeckVersion, version_id or deck.current_version_id)
    if version is None or version.deck_id != deck.id:
        _fail("web_deck.version_unknown", "导出版本不存在")
    return deck, version


def render_html(db: Any, *, actor: DBUser, deck_id: str, version_id: str | None = None) -> bytes:
    deck, version = _export_version(db, actor, deck_id, version_id)
    slides = []
    for index, slide in enumerate(version.slides_snapshot):
        bullets = "".join(f"<li>{html.escape(item)}</li>" for item in slide["bullets"])
        notes = html.escape(slide["speaker_notes"])
        slides.append(
            f'<section class="slide" id="slide-{index + 1}" data-slide-id="{html.escape(slide["id"])}">'
            f'<p class="number">{index + 1:02d}</p><h1>{html.escape(slide["title"])}</h1>'
            f'<h2>{html.escape(slide["subtitle"])}</h2><p class="body">{html.escape(slide["body"])}</p>'
            f'<ul>{bullets}</ul><aside class="notes">{notes}</aside></section>'
        )
    document = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(deck.title)}</title><style>
body{{margin:0;background:#1e2822;font-family:'Microsoft YaHei',sans-serif;color:#1e2822}}.slide{{box-sizing:border-box;width:min(100vw,160vh);aspect-ratio:16/9;margin:3vh auto;padding:7%;background:#f5f7f4;position:relative;overflow:hidden}}h1{{font-size:clamp(28px,4vw,56px);margin:0 0 2%;letter-spacing:0}}h2{{font-size:clamp(16px,2vw,28px);color:#355b42;font-weight:500}}.body,li{{font-size:clamp(14px,1.6vw,23px);line-height:1.55}}.number{{position:absolute;right:4%;top:3%;color:#637067}}.notes{{display:none}}@media print{{body{{background:white}}.slide{{break-after:page;margin:0;width:100vw;height:100vh}}}}</style></head><body>{''.join(slides)}</body></html>"""
    return document.encode("utf-8")


def render_pptx(db: Any, *, actor: DBUser, deck_id: str, version_id: str | None = None) -> bytes:
    deck, version = _export_version(db, actor, deck_id, version_id)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index, item in enumerate(version.slides_snapshot):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(245, 247, 244)
        title = slide.shapes.add_textbox(Inches(0.75), Inches(0.72), Inches(11.7), Inches(0.75))
        p = title.text_frame.paragraphs[0]
        p.text = item["title"]
        p.font.name = "Microsoft YaHei"; p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = RGBColor(30, 40, 34)
        if item["subtitle"]:
            sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.3), Inches(0.5))
            p = sub.text_frame.paragraphs[0]; p.text = item["subtitle"]; p.font.name = "Microsoft YaHei"; p.font.size = Pt(16); p.font.color.rgb = RGBColor(53, 91, 66)
        body = slide.shapes.add_textbox(Inches(0.85), Inches(2.25), Inches(11.55), Inches(3.9))
        frame = body.text_frame; frame.clear(); frame.word_wrap = True
        lines = ([item["body"]] if item["body"] else []) + [f"•  {bullet}" for bullet in item["bullets"]]
        for line_index, line in enumerate(lines or [""]):
            p = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            p.text = line; p.font.name = "Microsoft YaHei"; p.font.size = Pt(19); p.font.color.rgb = RGBColor(45, 56, 49); p.space_after = Pt(10)
        number = slide.shapes.add_textbox(Inches(11.85), Inches(6.75), Inches(0.7), Inches(0.3))
        p = number.text_frame.paragraphs[0]; p.text = f"{index + 1:02d}"; p.alignment = PP_ALIGN.RIGHT; p.font.size = Pt(10); p.font.color.rgb = RGBColor(99, 112, 103)
        notes = slide.notes_slide.notes_text_frame
        notes.text = item["speaker_notes"]
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def render_pdf(db: Any, *, actor: DBUser, deck_id: str, version_id: str | None = None) -> bytes:
    executable = _libreoffice()
    if executable is None:
        _fail("web_deck.pdf_unavailable", "未检测到 LibreOffice；请继续使用 PPTX 或 HTML 导出")
    pptx_bytes = render_pptx(db, actor=actor, deck_id=deck_id, version_id=version_id)
    with tempfile.TemporaryDirectory(prefix="edumatrix-deck-") as folder:
        source = Path(folder) / "deck.pptx"
        source.write_bytes(pptx_bytes)
        profile = Path(folder) / "lo-profile"
        result = subprocess.run(
            [executable, "--headless", f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf", "--outdir", folder, str(source)],
            capture_output=True, text=True, timeout=90,
        )
        target = Path(folder) / "deck.pdf"
        if result.returncode != 0 or not target.is_file() or target.stat().st_size == 0:
            _fail("web_deck.pdf_failed", "LibreOffice PDF 转换失败；PPTX 与 HTML 导出仍可用")
        return target.read_bytes()
