"""Real-browser acceptance for high-fidelity PPTX reading."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
import time
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from PIL import Image, ImageChops
from playwright.sync_api import Page, sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import e2e_course_ingestion_multiformat as ingestion
import e2e_f6_t03_code_sandbox as base


ROOT = Path(__file__).resolve().parents[1]
TASK_ID = "PPTX-HIGH-FIDELITY-READER-01"
OUTPUT_DIR = ROOT / "outputs" / "原系统融合重构执行" / "evidence" / "custom_tasks" / TASK_ID / "browser"
BUILT_IN_COURSE_ID = "course-embedded-machine-learning-engineering-v2"


def digest_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest().lower()


def digest(path: Path) -> str:
    return digest_bytes(path.read_bytes())


def clear_output() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for item in OUTPUT_DIR.iterdir():
        if item.name == "derived-previews":
            continue
        if item.is_file() or item.is_symlink():
            item.unlink()
        elif item.is_dir():
            for attempt in range(20):
                try:
                    shutil.rmtree(item)
                    break
                except FileNotFoundError:
                    if not item.exists():
                        break
                    if attempt == 19:
                        raise
                except PermissionError:
                    if attempt == 19:
                        raise
                time.sleep(0.25)


def official_state() -> dict[str, str]:
    paths = [
        ROOT / "edumatrix.db",
        ROOT / "outputs" / "原系统融合重构执行" / "evidence" / "context_recovery" / "active_manifest.json",
        ROOT / "outputs" / "原系统融合重构执行" / "evidence" / "context_recovery" / "current_checkpoint.json",
    ]
    return {str(path): digest(path) for path in paths if path.is_file()}


def built_in_fixture(db_path: Path) -> dict[str, object]:
    with sqlite3.connect(db_path) as db:
        row = db.execute(
            """
            SELECT id, course_id, filename, source_sha256, parser_metadata
            FROM source_documents
            WHERE course_id=? AND processing_status='ready' AND filename LIKE '3.3.5.%.pptx'
            ORDER BY id LIMIT 1
            """,
            (BUILT_IN_COURSE_ID,),
        ).fetchone()
        if row is None:
            raise RuntimeError("built-in multi-page PPTX fixture is missing")
        metadata = json.loads(row[4] or "{}")
        bundle = metadata.get("bundle_asset") if isinstance(metadata, dict) else {}
        asset_path = Path(str((bundle or {}).get("path") or ""))
        resolved = (asset_path if asset_path.is_absolute() else ROOT / asset_path).resolve()
        if not resolved.is_file():
            raise RuntimeError("built-in PPTX source asset is missing")
        return {
            "source_id": row[0],
            "course_id": row[1],
            "filename": row[2],
            "asset_path": str(resolved),
            "asset_sha256": str((bundle or {}).get("sha256") or row[3] or "").lower(),
            "asset_bytes": resolved.stat().st_size,
            "initial_learning_events": db.execute("SELECT COUNT(*) FROM learning_events").fetchone()[0],
        }


def set_text(shape, text: str, *, size: int, color: RGBColor, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    frame.margin_left = Inches(0.12)
    frame.margin_right = Inches(0.12)
    frame.margin_top = Inches(0.06)
    frame.margin_bottom = Inches(0.06)


def make_uploaded_pptx(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    first = presentation.slides.add_slide(blank)
    first.background.fill.solid()
    first.background.fill.fore_color.rgb = RGBColor(20, 39, 57)
    accent = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(50, 184, 164); accent.line.fill.background()
    title = first.shapes.add_textbox(Inches(0.7), Inches(0.65), Inches(8.7), Inches(0.8))
    set_text(title, "学习数据闭环", size=30, color=RGBColor(255, 255, 255), bold=True)
    subtitle = first.shapes.add_textbox(Inches(0.72), Inches(1.48), Inches(8.2), Inches(0.55))
    set_text(subtitle, "从学习行为到个性化反馈的实时演示", size=16, color=RGBColor(195, 210, 220))
    cards = [
        ("学习事件", "24", RGBColor(54, 130, 184)),
        ("掌握提升", "+18%", RGBColor(50, 184, 164)),
        ("今日完成", "6 / 8", RGBColor(239, 170, 70)),
    ]
    for index, (label, value, color) in enumerate(cards):
        left = Inches(0.72 + index * 3.0)
        card = first.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.45), Inches(2.55), Inches(2.05))
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(247, 249, 250); card.line.color.rgb = color
        value_box = first.shapes.add_textbox(left + Inches(0.15), Inches(2.78), Inches(2.25), Inches(0.72))
        set_text(value_box, value, size=27, color=color, bold=True, align=PP_ALIGN.CENTER)
        label_box = first.shapes.add_textbox(left + Inches(0.15), Inches(3.62), Inches(2.25), Inches(0.45))
        set_text(label_box, label, size=15, color=RGBColor(46, 61, 72), align=PP_ALIGN.CENTER)
    footer = first.shapes.add_textbox(Inches(0.72), Inches(6.58), Inches(8.5), Inches(0.42))
    set_text(footer, "EduMatrix · 高保真 PPTX 阅读验收", size=13, color=RGBColor(160, 181, 194))

    second = presentation.slides.add_slide(blank)
    second.background.fill.solid(); second.background.fill.fore_color.rgb = RGBColor(245, 247, 248)
    title = second.shapes.add_textbox(Inches(0.65), Inches(0.45), Inches(8.7), Inches(0.7))
    set_text(title, "反馈如何进入下一轮学习", size=27, color=RGBColor(20, 39, 57), bold=True)
    steps = [
        ("学习行为", "阅读 · 练习 · 对话", RGBColor(54, 130, 184)),
        ("状态更新", "掌握度 · 负荷 · 参与", RGBColor(50, 184, 164)),
        ("路径调整", "难度 · 节奏 · 复习", RGBColor(239, 170, 70)),
    ]
    for index, (heading, body, color) in enumerate(steps):
        left = Inches(0.6 + index * 3.15)
        circle = second.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.78), Inches(1.55), Inches(0.82), Inches(0.82))
        circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.fill.background()
        number = second.shapes.add_textbox(left + Inches(0.78), Inches(1.69), Inches(0.82), Inches(0.42))
        set_text(number, str(index + 1), size=18, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
        box = second.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.65), Inches(2.4), Inches(1.75))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(255, 255, 255); box.line.color.rgb = RGBColor(211, 220, 225)
        heading_box = second.shapes.add_textbox(left + Inches(0.15), Inches(2.92), Inches(2.1), Inches(0.5))
        set_text(heading_box, heading, size=18, color=RGBColor(20, 39, 57), bold=True, align=PP_ALIGN.CENTER)
        body_box = second.shapes.add_textbox(left + Inches(0.15), Inches(3.57), Inches(2.1), Inches(0.48))
        set_text(body_box, body, size=13, color=RGBColor(83, 99, 109), align=PP_ALIGN.CENTER)
        if index < 2:
            arrow = second.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                left + Inches(2.48), Inches(3.5), left + Inches(3.03), Inches(3.5),
            )
            arrow.line.color.rgb = RGBColor(116, 135, 146); arrow.line.width = Pt(2)
    observation = second.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(5.2), Inches(8.75), Inches(1.15))
    observation.fill.solid(); observation.fill.fore_color.rgb = RGBColor(231, 238, 242); observation.line.fill.background()
    observation_text = second.shapes.add_textbox(Inches(0.88), Inches(5.47), Inches(8.2), Inches(0.6))
    set_text(observation_text, "观察重点：反馈应改变后续安排，同时保留学习者可理解的解释。", size=16, color=RGBColor(46, 61, 72), bold=True)

    presentation.core_properties.title = "PPTX 高保真阅读验收"
    presentation.save(path)


def diagnostics(page: Page, report: dict[str, object]) -> None:
    def on_response(response) -> None:
        parsed = urlparse(response.url)
        if parsed.scheme not in {"blob", "data"} and parsed.hostname not in {"127.0.0.1", "localhost"}:
            report["external_requests"].append(response.url)
        if response.status >= 400:
            report["http_errors"].append({"status": response.status, "url": response.url})

    def on_request_failed(request) -> None:
        if request.url.startswith("blob:") and request.failure == "net::ERR_ABORTED":
            report["expected_request_aborts"].append({"url": request.url, "failure": request.failure})
            return
        report["failed_requests"].append({"url": request.url, "failure": request.failure})

    page.on("response", on_response)
    page.on("requestfailed", on_request_failed)
    page.on("pageerror", lambda error: report["page_errors"].append(str(error)))
    page.on("console", lambda message: report["console_errors"].append(message.text) if message.type == "error" else None)


def capture(page: Page, report: dict[str, object], filename: str) -> None:
    path = OUTPUT_DIR / filename
    page.screenshot(path=path, full_page=False)
    report["screenshots"].append({
        "file": filename,
        "width": page.viewport_size["width"],
        "height": page.viewport_size["height"],
        "bytes": path.stat().st_size,
        "sha256": digest(path),
    })


def create_course(page: Page, base_url: str, report: dict[str, object]) -> str:
    result = page.evaluate(
        r"""
        async ({baseUrl}) => {
          const token = localStorage.getItem('edumatrix_token') || ''
          const headers = { Authorization: `Bearer ${token}`, 'Content-Type': 'application/json' }
          const created = await fetch(`${baseUrl}/api/knowledge/courses`, {
            method: 'POST',
            headers,
            body: JSON.stringify({
              title: '高保真课件阅读验收课程',
              subject: '演示文稿阅读与学习反馈',
              description: '验证后续上传的 PPTX 可以保留原始页面布局，并继续支持课程搜索与学习工具。',
            }),
          })
          const createdBody = await created.json()
          if (![200, 201].includes(created.status)) {
            throw new Error(`course creation failed: ${created.status} ${JSON.stringify(createdBody)}`)
          }

          const courseId = createdBody.course.id
          const form = new FormData()
          form.append('files', new File([
            '# 课程导读\n\n本课程用于验证演示文稿上传、阅读与学习工具。\n\n' +
            '## 学习目标\n\n能够查看原始幻灯片版式，按页浏览，并在解析阅读版中搜索内容。\n\n' +
            '## 使用方式\n\n上传课件后进入阅读器，通过缩略图、翻页、缩放和搜索完成学习。\n'
          ], '课程导读.md', { type: 'text/markdown' }))
          const admitted = await fetch(
            `${baseUrl}/api/knowledge/upload-batch?student_id=default&course_id=${encodeURIComponent(courseId)}`,
            { method: 'POST', headers: { Authorization: `Bearer ${token}` }, body: form },
          )
          const admittedBody = await admitted.json()
          if (admitted.status !== 200 || admittedBody.succeeded !== 1) {
            throw new Error(`course bootstrap failed: ${admitted.status} ${JSON.stringify(admittedBody)}`)
          }
          return {
            courseId,
            createStatus: created.status,
            bootstrapStatus: admitted.status,
            bootstrapReady: admittedBody.succeeded,
          }
        }
        """,
        {"baseUrl": base_url},
    )
    report["checks"]["new_course_bootstrap"] = result
    return str(result["courseId"])


def upload_pptx(page: Page, base_url: str, course_id: str, path: Path) -> dict[str, object]:
    page.goto(f"{base_url}/knowledge/courses/{course_id}", wait_until="networkidle", timeout=60_000)
    page.get_by_test_id("add-course-material").wait_for(state="visible", timeout=30_000)
    with page.expect_response(
        lambda response: "/api/knowledge/upload-batch" in response.url and response.request.method == "POST",
        timeout=180_000,
    ) as upload:
        page.locator("input[type=file]").set_input_files(str(path))
    response = upload.value
    payload = response.json()
    if response.status != 200 or payload.get("succeeded") != 1:
        raise AssertionError(f"PPTX upload failed: {response.status} {payload}")
    page.get_by_text("1 份资料已进入课程", exact=False).wait_for(state="visible", timeout=60_000)
    document = next(item["document"] for item in payload["results"] if item.get("status") == "ready")
    return document


def open_reader(page: Page, base_url: str, source_id: str, course_id: str) -> None:
    page.goto(
        f"{base_url}/knowledge/documents/{source_id}?course_id={course_id}",
        wait_until="domcontentloaded",
        timeout=180_000,
    )
    page.locator(".reader-page").wait_for(state="visible", timeout=60_000)
    page.get_by_text("正在打开原文", exact=True).wait_for(state="hidden", timeout=180_000)
    page.locator(".pptx-reader").wait_for(state="visible", timeout=60_000)


def wait_for_slide_images(page: Page, expected: int) -> None:
    page.locator(".pptx-slide-canvas__page").wait_for(state="visible", timeout=180_000)
    page.wait_for_function(
        """
        expected => {
          const images = [...document.querySelectorAll('.pptx-thumbnail img[src^="blob:"]')]
          return images.length === expected && images.every(image => image.complete && image.naturalWidth > 0)
        }
        """,
        arg=expected,
        timeout=180_000,
    )


def visual_contract(page: Page) -> dict[str, object]:
    return page.evaluate(
        """
        () => {
          const root = document.querySelector('.pptx-reader')
          const visible = element => {
            if (!element) return false
            const style = getComputedStyle(element)
            const box = element.getBoundingClientRect()
            return style.display !== 'none' && style.visibility !== 'hidden' && box.width > 0 && box.height > 0
          }
          const textElements = [...(root?.querySelectorAll('button, strong, span, p, summary') || [])]
            .filter(element => visible(element) && element.textContent.trim())
          const smallText = textElements
            .filter(element => Number.parseFloat(getComputedStyle(element).fontSize) < 14)
            .map(element => element.textContent.trim().slice(0, 80))
          const clippedText = textElements.filter(element => {
            const style = getComputedStyle(element)
            if (style.textOverflow === 'ellipsis' || ['auto', 'scroll'].includes(style.overflowX) || ['auto', 'scroll'].includes(style.overflowY)) return false
            return element.scrollWidth > element.clientWidth + 1 || element.scrollHeight > element.clientHeight + 4
          }).map(element => element.textContent.trim().slice(0, 80))
          const image = root?.querySelector('.pptx-slide-canvas__page')
          const box = image?.getBoundingClientRect()
          const naturalRatio = image?.naturalHeight ? image.naturalWidth / image.naturalHeight : 0
          const renderedRatio = box?.height ? box.width / box.height : 0
          return {
            rootVisible: visible(root),
            pageOverflow: document.documentElement.scrollWidth > document.documentElement.clientWidth + 1,
            smallText,
            clippedText,
            thumbnailCount: root?.querySelectorAll('.pptx-thumbnail').length || 0,
            blobThumbnailCount: root?.querySelectorAll('.pptx-thumbnail img[src^="blob:"]').length || 0,
            projectionInMainCanvas: root?.querySelectorAll('.pptx-slide-canvas__text').length || 0,
            slideImageVisible: visible(image),
            slideNaturalWidth: image?.naturalWidth || 0,
            slideNaturalHeight: image?.naturalHeight || 0,
            slideNaturalRatio: naturalRatio,
            slideRenderedRatio: renderedRatio,
            ratioDelta: Math.abs(naturalRatio - renderedRatio),
            parsedReadingVisible: visible(root?.querySelector('.pptx-parsed-fallback')),
          }
        }
        """
    )


def pixel_difference(first: bytes, second: bytes) -> float:
    first_image = Image.open(BytesIO(first)).convert("RGB")
    second_image = Image.open(BytesIO(second)).convert("RGB")
    if first_image.size != second_image.size:
        second_image = second_image.resize(first_image.size)
    mask = ImageChops.difference(first_image, second_image).convert("L").point(lambda value: 255 if value else 0)
    return mask.histogram()[255] / (first_image.width * first_image.height)


def browser_export(page: Page, base_url: str, source_id: str) -> dict[str, object]:
    return page.evaluate(
        """
        async ({url}) => {
          const token = localStorage.getItem('edumatrix_token') || ''
          const response = await fetch(url, { headers: { Authorization: `Bearer ${token}` } })
          const bytes = await response.arrayBuffer()
          const hash = await crypto.subtle.digest('SHA-256', bytes)
          return {
            status: response.status,
            contentType: response.headers.get('content-type') || '',
            bytes: bytes.byteLength,
            prefix: String.fromCharCode(...new Uint8Array(bytes.slice(0, 2))),
            sha256: [...new Uint8Array(hash)].map(value => value.toString(16).padStart(2, '0')).join(''),
          }
        }
        """,
        {"url": f"{base_url}/api/v1/document-reader/documents/{source_id}/export"},
    )


def exercise_reader(
    page: Page,
    report: dict[str, object],
    base_url: str,
    *,
    source_id: str,
    course_id: str,
    expected_slides: int,
    prefix: str,
    expected_ratio: float,
) -> None:
    open_reader(page, base_url, source_id, course_id)
    wait_for_slide_images(page, expected_slides)
    page.locator(".pptx-reader").scroll_into_view_if_needed()
    audit = visual_contract(page)
    assert audit["rootVisible"] and audit["slideImageVisible"], audit
    assert audit["thumbnailCount"] == expected_slides == audit["blobThumbnailCount"], audit
    assert audit["projectionInMainCanvas"] == 0, audit
    assert not audit["pageOverflow"] and not audit["smallText"] and not audit["clippedText"], audit
    assert abs(audit["slideNaturalRatio"] - expected_ratio) < 0.02, audit
    assert audit["ratioDelta"] < 0.02, audit
    report["visual"][f"{prefix}_initial"] = audit
    capture(page, report, f"{prefix}-01-original-layout-1440x900.png")

    image = page.locator(".pptx-slide-canvas__page")
    first = image.screenshot(path=OUTPUT_DIR / f"{prefix}-slide-01.png")
    page.get_by_label("下一张幻灯片").click()
    page.get_by_text(f"第 2 / {expected_slides} 张", exact=True).wait_for(state="visible")
    second = image.screenshot(path=OUTPUT_DIR / f"{prefix}-slide-02.png")
    changed = pixel_difference(first, second)
    assert digest_bytes(first) != digest_bytes(second) and changed > 0.01, changed
    report["checks"][f"{prefix}_slide_difference"] = {
        "first_sha256": digest_bytes(first),
        "second_sha256": digest_bytes(second),
        "changed_pixel_ratio": changed,
    }
    capture(page, report, f"{prefix}-02-next-page-1440x900.png")

    before = image.bounding_box()
    page.get_by_role("button", name="放大", exact=True).click()
    page.wait_for_timeout(250)
    after = image.bounding_box()
    assert before and after and after["width"] > before["width"] * 1.08, (before, after)
    report["checks"][f"{prefix}_zoom"] = {"before": before["width"], "after": after["width"]}
    capture(page, report, f"{prefix}-03-zoom-110-1440x900.png")
    page.get_by_role("button", name="缩小", exact=True).click()

    fullscreen = page.get_by_role("button", name="全屏查看幻灯片", exact=True)
    assert fullscreen.is_enabled()
    fullscreen.click()
    page.wait_for_function("document.fullscreenElement !== null", timeout=10_000)
    fullscreen_box = image.bounding_box()
    viewport = page.viewport_size
    assert fullscreen_box and viewport
    assert fullscreen_box["x"] >= -1 and fullscreen_box["y"] >= -1, fullscreen_box
    assert fullscreen_box["x"] + fullscreen_box["width"] <= viewport["width"] + 1, fullscreen_box
    assert fullscreen_box["y"] + fullscreen_box["height"] <= viewport["height"] + 1, fullscreen_box
    report["checks"][f"{prefix}_fullscreen"] = {
        "image_box": fullscreen_box,
        "viewport": viewport,
        "fully_visible": True,
    }
    capture(page, report, f"{prefix}-04-fullscreen-1440x900.png")
    page.get_by_role("button", name="退出全屏", exact=True).click()
    page.wait_for_function("document.fullscreenElement === null", timeout=10_000)

    page.get_by_label("在当前文档中搜索").fill("反馈" if prefix == "uploaded" else "Confusion Matrix")
    page.locator(".pptx-parsed-fallback").scroll_into_view_if_needed()
    page.wait_for_timeout(250)
    assert page.locator(".pptx-parsed-fallback .course-document-preview").is_visible()
    assert "无结果" not in page.locator(".reader-search").inner_text()
    assert page.locator(".pptx-parsed-fallback .course-search-hit--current").count() == 1
    report["checks"][f"{prefix}_parsed_search"] = True
    capture(page, report, f"{prefix}-05-parsed-search-1440x900.png")


def exercise_failure_fallback(page: Page, report: dict[str, object], base_url: str, document: dict[str, object], course_id: str) -> None:
    open_reader(page, base_url, str(document["source_document_id"]), course_id)
    notice = page.get_by_text("原版幻灯片暂时无法生成", exact=False)
    notice.wait_for(state="visible", timeout=60_000)
    notice.scroll_into_view_if_needed()
    page.wait_for_timeout(200)
    assert page.locator(".pptx-slide-canvas__page").count() == 0
    assert page.locator(".pptx-slide-canvas__text").is_visible()
    assert page.locator(".pptx-parsed-fallback .course-document-preview").is_visible()
    report["checks"]["failure_fallback"] = {
        "reader_visible": True,
        "projection_visible": True,
        "parsed_reading_visible": True,
        "notice": notice.inner_text(),
    }
    capture(page, report, "failure-01-readable-fallback-1440x900.png")


def database_state(db_path: Path, course_id: str, initial_events: int) -> dict[str, object]:
    with sqlite3.connect(db_path, timeout=20) as db:
        return {
            "uploaded_ready_pptx": db.execute(
                "SELECT COUNT(*) FROM source_documents WHERE course_id=? AND file_format='pptx' AND processing_status='ready'",
                (course_id,),
            ).fetchone()[0],
            "learning_events_before": initial_events,
            "learning_events_after": db.execute("SELECT COUNT(*) FROM learning_events").fetchone()[0],
            "integrity": db.execute("PRAGMA integrity_check").fetchone()[0],
            "foreign_key_violations": len(db.execute("PRAGMA foreign_key_check").fetchall()),
        }


def run() -> dict[str, object]:
    clear_output()
    base.OUTPUT_DIR = OUTPUT_DIR
    base.PORT_START = int(os.getenv("EDUMATRIX_PPTX_HIGH_FIDELITY_PORT", "9018"))
    db_path = OUTPUT_DIR / "pptx_high_fidelity.sqlite3"
    fixture_path = OUTPUT_DIR / "uploaded-4x3-two-page.pptx"
    make_uploaded_pptx(fixture_path)
    before = official_state()
    model_server, _model_thread = ingestion._start_model_stub()
    port = base.free_port()
    env = base.runtime_env(db_path, port)
    env.update({
        "EDUMATRIX_AUTH_SECRET_KEY": "pptx-high-fidelity-isolated-secret",
        "EDUMATRIX_LLM_PROVIDER": "openai",
        "EDUMATRIX_LLM_MODEL": "deepseek-v4-flash",
        "EDUMATRIX_LLM_API_KEY": "isolated-browser-model",
        "EDUMATRIX_LLM_ENDPOINT": f"http://127.0.0.1:{model_server.server_port}/chat/completions",
        "EDUMATRIX_LLM_THINKING_MODE": "max",
        "EDUMATRIX_LLM_REASONING_EFFORT": "max",
        "EDUMATRIX_UPLOAD_DIR": str(OUTPUT_DIR / "uploads"),
        "EDUMATRIX_PPTX_PREVIEW_DIR": str(OUTPUT_DIR / "derived-previews"),
        "EDUMATRIX_MAX_UPLOAD_BYTES": str(20 * 1024 * 1024),
    })
    ingestion.seed_isolated_database(env)
    built_in = built_in_fixture(db_path)
    report: dict[str, object] = {
        "schema_version": "pptx-high-fidelity-reader-01.browser.v1",
        "task_id": TASK_ID,
        "status": "running",
        "viewports": [[1440, 900], [1366, 768]],
        "screenshots": [],
        "checks": {},
        "visual": {},
        "http_errors": [],
        "failed_requests": [],
        "expected_request_aborts": [],
        "page_errors": [],
        "console_errors": [],
        "external_requests": [],
        "built_in": built_in,
    }
    process = None
    second_process = None
    uploaded: dict[str, object] = {}
    course_id = ""
    try:
        process = base.start_service(env, port)
        base_url = f"http://127.0.0.1:{port}"
        report["base_url"] = base_url
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            context = browser.new_context(viewport={"width": 1440, "height": 900}, locale="zh-CN")
            page = context.new_page(); diagnostics(page, report); token = ingestion.login(page, base_url)
            exercise_reader(
                page, report, base_url,
                source_id=str(built_in["source_id"]), course_id=str(built_in["course_id"]),
                expected_slides=18, prefix="builtin", expected_ratio=16 / 9,
            )
            built_in_export = browser_export(page, base_url, str(built_in["source_id"]))
            assert built_in_export["status"] == 200 and built_in_export["prefix"] == "PK"
            assert "presentationml.presentation" in str(built_in_export["contentType"])
            assert built_in_export["sha256"] == built_in["asset_sha256"]
            report["checks"]["built_in_original_export"] = built_in_export

            course_id = create_course(page, base_url, report)
            uploaded = upload_pptx(page, base_url, course_id, fixture_path)
            report["uploaded"] = {"course_id": course_id, **uploaded}
            capture(page, report, "uploaded-00-course-material-ready-1440x900.png")
            exercise_reader(
                page, report, base_url,
                source_id=str(uploaded["source_document_id"]), course_id=course_id,
                expected_slides=2, prefix="uploaded", expected_ratio=4 / 3,
            )
            uploaded_export = browser_export(page, base_url, str(uploaded["source_document_id"]))
            assert uploaded_export["status"] == 200 and uploaded_export["prefix"] == "PK"
            assert uploaded_export["sha256"] == digest(fixture_path)
            report["checks"]["uploaded_original_export"] = uploaded_export
            context.close()

            compact = browser.new_context(viewport={"width": 1366, "height": 768}, locale="zh-CN")
            page = compact.new_page(); diagnostics(page, report); ingestion.login(page, base_url)
            open_reader(page, base_url, str(uploaded["source_document_id"]), course_id)
            wait_for_slide_images(page, 2)
            audit = visual_contract(page)
            assert not audit["pageOverflow"] and not audit["smallText"] and not audit["clippedText"], audit
            assert audit["ratioDelta"] < 0.02 and audit["blobThumbnailCount"] == 2, audit
            report["visual"]["uploaded_1366x768"] = audit
            capture(page, report, "uploaded-06-original-layout-1366x768.png")
            compact.close()
            browser.close()

        base.stop_service(process, port); process = None

        blocker = OUTPUT_DIR / "blocked-preview-root"
        blocker.write_text("intentional conversion-root fixture", encoding="utf-8")
        second_port = base.free_port()
        failure_env = dict(env)
        failure_env["EDUMATRIX_PORT"] = str(second_port)
        failure_env["EDUMATRIX_PPTX_PREVIEW_DIR"] = str(blocker)
        second_process = base.start_service(failure_env, second_port)
        failure_url = f"http://127.0.0.1:{second_port}"
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            context = browser.new_context(viewport={"width": 1440, "height": 900}, locale="zh-CN")
            page = context.new_page(); diagnostics(page, report); ingestion.login(page, failure_url)
            failure_document = {"source_document_id": built_in["source_id"]}
            exercise_failure_fallback(
                page,
                report,
                failure_url,
                failure_document,
                str(built_in["course_id"]),
            )
            export = browser_export(page, failure_url, str(built_in["source_id"]))
            assert export["status"] == 200 and export["sha256"] == built_in["asset_sha256"]
            report["checks"]["failure_path_original_export"] = export
            context.close(); browser.close()

        report["database"] = database_state(db_path, course_id, int(built_in["initial_learning_events"]))
        assert report["database"]["uploaded_ready_pptx"] == 1
        assert report["database"]["learning_events_after"] == report["database"]["learning_events_before"]
        assert report["database"]["integrity"] == "ok"
        assert report["database"]["foreign_key_violations"] == 0
        assert digest(Path(str(built_in["asset_path"]))) == built_in["asset_sha256"]
        assert not report["http_errors"], report["http_errors"]
        assert not report["failed_requests"], report["failed_requests"]
        assert not report["page_errors"], report["page_errors"]
        assert not report["console_errors"], report["console_errors"]
        assert not report["external_requests"], report["external_requests"]
        report["status"] = "passed"
        return report
    except Exception as exc:
        report["status"] = "failed"
        report["failure"] = {"type": type(exc).__name__, "message": str(exc)}
        raise
    finally:
        base.stop_service(process, port)
        if second_process is not None:
            base.stop_service(second_process, int(second_process.args and failure_env["EDUMATRIX_PORT"]))
        model_server.shutdown()
        model_server.server_close()
        if db_path.exists():
            ingestion.stabilize_evidence_database(db_path)
        report["finished_at"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
        report["official_state_before"] = before
        report["official_state_after"] = official_state()
        report["official_state_unchanged"] = report["official_state_before"] == report["official_state_after"]
        (OUTPUT_DIR / "browser_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        files = [
            {"file": path.name, "bytes": path.stat().st_size, "sha256": digest(path)}
            for path in sorted(OUTPUT_DIR.iterdir())
            if path.is_file() and path.name not in {"browser_report.json", "manifest.json"} and not path.name.endswith(("-wal", "-shm"))
        ]
        (OUTPUT_DIR / "manifest.json").write_text(
            json.dumps({"schema_version": "pptx-high-fidelity-reader-01.evidence.v1", "task_id": TASK_ID, "files": files}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2))
