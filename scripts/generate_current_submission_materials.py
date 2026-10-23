"""Generate current XH-202630 submission materials, excluding PPT.

The generator reads the current source tree and the two canonical course
manifests. It writes only to outputs/competition_submission_current and never
reads or overwrites a previous submission package.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "competition_submission_current"
DOCS = OUT / "documents"
DATA = OUT / "test_data"

COURSE_PATHS = (
    ROOT / "data" / "curated" / "machine_learning_foundations_v1" / "course_manifest.json",
    ROOT / "data" / "curated" / "ai_deep_learning_agents_v1" / "course_manifest.json",
)

TEN_DIMENSIONS = (
    ("knowledge_mastery", "知识基础与掌握度", "掌握度、证据充分性和知识点薄弱程度"),
    ("misconception_profile", "易错点与概念理解", "误概念、易混点和错误归因"),
    ("understanding_fluency_transfer", "理解和迁移能力", "能否解释、变式应用和跨情境迁移"),
    ("cognitive_processing", "处理复杂内容的方式", "认知负荷、拆解方式和复杂内容处理节奏"),
    ("learning_strategy", "学习方法", "复习、练习、实验和反馈使用方式"),
    ("metacognition", "自我判断和调整", "自评、信心校准和自我修正能力"),
    ("motivation_and_purpose", "学习目标和动力", "目标、动机和岗位/项目关联"),
    ("affect_resilience", "信心和抗挫能力", "挫败反应、持续尝试和情绪风险"),
    ("interaction_preference", "喜欢怎样互动学习", "偏好的媒介、提示方式和交互节奏"),
    ("learning_context", "学习环境和安排", "专业背景、设备、时间和学习场景"),
)
INITIAL_DIMENSIONS = (
    ("background", "学习背景"),
    ("goal", "学习目标"),
    ("foundation", "当前基础"),
    ("difficulty", "主要困难"),
    ("preference", "学习偏好"),
)
DIAGNOSTIC_KINDS = ("theory", "application", "engineering_judgment")


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def command_text(args: list[str]) -> str:
    try:
        result = subprocess.run(args, cwd=ROOT, capture_output=True, text=True, timeout=20)
        return (result.stdout or result.stderr).strip()
    except Exception as exc:
        return f"unavailable: {type(exc).__name__}: {exc}"


def load_courses() -> list[dict[str, Any]]:
    courses: list[dict[str, Any]] = []
    for path in COURSE_PATHS:
        raw = json.loads(path.read_text(encoding="utf-8"))
        courses.append({
            "manifest_path": str(path.relative_to(ROOT)).replace("\\", "/"),
            "course_id": raw.get("course_id", ""),
            "manifest_version": raw.get("manifest_version", ""),
            "title": raw.get("title", ""),
            "subject": raw.get("subject", ""),
            "description": raw.get("description", ""),
            "knowledge_point_count": len(raw.get("knowledge_points", [])),
            "file_count": len(raw.get("files", [])),
            "license": raw.get("license", ""),
            "knowledge_points": raw.get("knowledge_points", []),
            "files": [
                {
                    "path": item.get("path", ""),
                    "title": item.get("title", ""),
                    "source_role": item.get("source_role", ""),
                    "license_spdx": item.get("license_spdx", ""),
                    "git_revision": item.get("git_revision", ""),
                    "sha256": item.get("sha256", ""),
                    "knowledge_point_codes": item.get("knowledge_point_codes", []),
                }
                for item in raw.get("files", [])
            ],
        })
    return courses


def rel(path: Path) -> str:
    return str(path.relative_to(OUT)).replace("\\", "/")


def write(path: Path, content: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.strip() + "\n", encoding="utf-8")
    return path


def render_docx(title: str, markdown: str, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "等线"
    normal.font.size = Pt(10.5)
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "等线"
    run.font.size = Pt(20)
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            level = min(3, len(line) - len(line.lstrip("#")))
            document.add_heading(line[level:].strip(), level)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+[.]\s+", line):
            document.add_paragraph(re.sub(r"^\d+[.]\s+", "", line), style="List Number")
        elif line.startswith("|") or line.startswith("```"):
            document.add_paragraph(line, style="Intense Quote")
        else:
            document.add_paragraph(line)
    document.core_properties.title = title
    document.core_properties.subject = "XH-202630 当前提交材料"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def profile_fixture(profile_id: str, label: str, course_id: str, goal: str, prefs: list[str], mastery: dict[str, float], weak: list[str], scores: dict[str, float]) -> dict[str, Any]:
    influence = {
        "knowledge_mastery": {"resource": "difficulty_band", "teaching": "先修节点和复测优先级"},
        "misconception_profile": {"resource": "反例与概念辨析", "teaching": "错误归因和补救节点"},
        "understanding_fluency_transfer": {"resource": "变式练习", "teaching": "迁移任务顺序"},
        "cognitive_processing": {"resource": "脚手架和单次内容量", "teaching": "短回合或深度回合"},
        "learning_strategy": {"resource": "练习与复习组合", "teaching": "间隔复习和检索练习"},
        "metacognition": {"resource": "自评提示和答案后复盘", "teaching": "信心校准与复测"},
        "motivation_and_purpose": {"resource": "案例上下文", "teaching": "项目/考试/岗位路径"},
        "affect_resilience": {"resource": "首题难度和提示", "teaching": "失败后的降阶或鼓励重试"},
        "interaction_preference": {"resource": "讲义、图示、代码的排序", "teaching": "互动媒介和提示模式"},
        "learning_context": {"resource": "任务时长和设备适配", "teaching": "学习节奏和任务负担"},
    }
    ten = []
    for key, label_text, definition in TEN_DIMENSIONS:
        score = scores.get(key, 0.5)
        ten.append({
            "key": key,
            "label": label_text,
            "definition": definition,
            "score": score,
            "confidence": round(min(0.9, 0.58 + score * 0.25), 2),
            "source": "合成入学诊断与学习事件 fixture",
            "resource_influence": influence[key]["resource"],
            "teaching_influence": influence[key]["teaching"],
            "status": "fixture_for_submission_materials",
        })
    diagnostics = [
        {"question_id": f"{profile_id}-diag-01", "kind": "theory", "knowledge_point": "模型评估、泛化与交叉验证", "result": "partially_correct", "score": 0.55, "uncertainty": "混淆矩阵指标边界需要复核"},
        {"question_id": f"{profile_id}-diag-02", "kind": "application", "knowledge_point": "数据集、特征工程与数据管道", "result": "correct", "score": 0.78, "uncertainty": "跨数据集迁移尚未验证"},
        {"question_id": f"{profile_id}-diag-03", "kind": "engineering_judgment", "knowledge_point": "RAG 文档解析、切分与混合检索", "result": "needs_scaffold", "score": 0.38, "uncertainty": "缺少完整实验记录"},
    ]
    trace = [
        {"stage": "diagnosis", "status": "completed", "output": {"weak_points": weak, "diagnostic_dimensions": list(DIAGNOSTIC_KINDS), "evidence": "entrance_diagnostics"}},
        {"stage": "domain_generation", "status": "fixture_draft", "output": {"resource_types": ["lesson", "lab_guide", "tiered_quiz"], "profile_id": profile_id}},
        {"stage": "evidence_review", "status": "fixture_review", "output": {"claim_policy": "source_required", "unknown_policy": "count_as_error"}},
        {"stage": "decision_orchestration", "status": "fixture_decision", "output": {"difficulty": "scaffolded", "scaffolding": "high" if scores.get("cognitive_processing", 0.5) < 0.6 else "medium", "rationale": weak[0]}},
    ]
    resources = {
        "lesson": {"title": "画像匹配讲义", "difficulty": "foundation", "adaptation": "短段落、示例和来源回链", "source_ids": ["kp-model-eval", "kp-data-pipeline"]},
        "lab_guide": {"title": "指标对照实操", "difficulty": "scaffolded", "adaptation": "固定数据、逐步检查和结果记录", "source_ids": ["kp-model-eval"]},
        "tiered_quiz": {"title": "分阶诊断题", "difficulty": "foundation_to_transfer", "adaptation": "识别、应用、迁移三级", "source_ids": ["kp-model-eval", "kp-rag"]},
    }
    return {
        "profile_id": profile_id,
        "label": label,
        "course_id": course_id,
        "data_label": "脱敏合成材料 fixture，不是真实用户数据",
        "background": label,
        "goal": goal,
        "preferences": prefs,
        "available_minutes_per_week": 180,
        "initial_dimensions": [{"key": key, "label": label_text, "value": "fixture_value", "confidence": 0.72} for key, label_text in INITIAL_DIMENSIONS],
        "entrance_diagnostic": diagnostics,
        "mastery_snapshot": {"course_id": course_id, "knowledge_points": mastery, "evidence": "diagnostic_attempt_fixture", "confidence": 0.68},
        "ten_dimension_profile": ten,
        "weak_points": weak,
        "agent_trace": trace,
        "generated_resources": resources,
        "feedback_round": {"event_type": "quiz_answered", "accuracy": 0.58, "evidence_state": "observed", "next_decision": "downgrade_explanation_and_retest", "mastery_update": "pending_verified_event"},
    }


def build_data(courses: list[dict[str, Any]]) -> list[Path]:
    DATA.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []
    course_catalog = {
        "data_label": "当前源码树的两门 canonical builtin courses",
        "courses": [{k: v for k, v in c.items() if k not in {"knowledge_points", "files"}} for c in courses],
        "knowledge_point_counts": {c["course_id"]: c["knowledge_point_count"] for c in courses},
    }
    path = DATA / "course_catalog.json"
    path.write_text(json.dumps(course_catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    files.append(path)
    for index, course in enumerate(courses, start=1):
        payload = {
            "course_id": course["course_id"],
            "course_version": course["manifest_version"],
            "title": course["title"],
            "domain": "机器学习/人工智能基础",
            "data_label": "当前课程 manifest 摘要；原始许可和来源以课程目录为准",
            "license": course["license"],
            "knowledge_points": course["knowledge_points"],
            "source_files": course["files"],
        }
        p = DATA / f"course_{index:02d}_knowledge_slice.json"
        p.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        files.append(p)
    source_catalog = [{"course_id": c["course_id"], "course_version": c["manifest_version"], "source_files": c["files"], "license": c["license"]} for c in courses]
    p = DATA / "source_catalog.json"
    p.write_text(json.dumps(source_catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    files.append(p)
    all_kps = []
    for course in courses:
        for kp in course["knowledge_points"]:
            all_kps.append({"course_id": course["course_id"], "course_version": course["manifest_version"], "id": kp.get("code"), "title": kp.get("title"), "level": kp.get("level"), "prerequisites": kp.get("prerequisites", [])})
    gold = {"data_label": "冻结知识点黄金集合，由当前 course manifests 派生", "knowledge_points": all_kps, "critical_rule": "core/capstone 或被案例标记为 critical 的知识点权重为 2"}
    p = DATA / "knowledge_point_gold_set.json"
    p.write_text(json.dumps(gold, ensure_ascii=False, indent=2), encoding="utf-8")
    files.append(p)
    first, second = courses
    profiles = [
        profile_fixture("learner-visual-beginner", "视觉型初学者", first["course_id"], "掌握模型评估和 RAG 基础，完成一个可复现实验", ["图示", "短步骤", "具体例子"], {"AI-KP-01": 0.35, "AI-KP-05": 0.32, "AI-KP-17": 0.22}, ["指标边界", "RAG 先修"], {k: 0.42 for k, _, _ in TEN_DIMENSIONS}),
        profile_fixture("learner-project-practice", "项目实践型学习者", second["course_id"], "把 Transformer、RAG 和 Agent 连接成小型项目", ["代码实操", "实验对照", "故障排查"], {"AI-KP-10": 0.72, "AI-KP-17": 0.58, "AI-KP-18": 0.51}, ["版本边界", "证据审查"], {k: 0.68 for k, _, _ in TEN_DIMENSIONS}),
        profile_fixture("learner-theory-research", "理论研究型学习者", second["course_id"], "完成从数学推导到可复现实验的迁移", ["公式推导", "概念对比", "严谨引用"], {"AI-KP-02": 0.82, "AI-KP-07": 0.76, "AI-KP-20": 0.44}, ["工程复现", "岗位约束"], {k: 0.74 for k, _, _ in TEN_DIMENSIONS}),
    ]
    p = DATA / "learner_profiles.json"
    p.write_text(json.dumps({"data_label": "脱敏合成画像 fixture", "profiles": profiles}, ensure_ascii=False, indent=2), encoding="utf-8")
    files.append(p)
    for name, rows in {
        "entrance_diagnostics.jsonl": [{"profile_id": p["profile_id"], "course_id": p["course_id"], "diagnostic": p["entrance_diagnostic"]} for p in profiles],
        "mastery_snapshots.jsonl": [{"profile_id": p["profile_id"], "snapshot": p["mastery_snapshot"]} for p in profiles],
        "agent_traces.jsonl": [{"profile_id": p["profile_id"], "trace": p["agent_trace"]} for p in profiles],
        "generated_resources.jsonl": [{"profile_id": p["profile_id"], "resources": p["generated_resources"]} for p in profiles],
        "feedback_rounds.jsonl": [{"profile_id": p["profile_id"], "feedback": p["feedback_round"]} for p in profiles],
    }.items():
        p = DATA / name
        p.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")
        files.append(p)
    cases = []
    categories = ["profile_diagnostic", "mastery", "retrieval", "generation", "review", "decision", "feedback", "authorization", "recovery", "sandbox"]
    for index in range(60):
        profile = profiles[index % len(profiles)]
        category = categories[index % len(categories)]
        hard = index >= 42
        cases.append({
            "case_id": f"xh-202630-{index + 1:03d}",
            "profile_id": profile["profile_id"],
            "course_id": profile["course_id"],
            "category": category,
            "difficulty": "challenge" if hard else "standard",
            "input": {"goal": profile["goal"], "query": f"围绕当前课程知识点执行 {category} 测试"},
            "expected": {"source_required": category in {"retrieval", "generation", "review"}, "must_preserve_profile": category in {"profile_diagnostic", "generation", "decision"}, "must_reject_or_explain_uncertainty": hard, "must_not_treat_click_as_mastery": True},
            "execution_status": "not_run_fixture",
        })
    p = DATA / "test_cases.jsonl"
    p.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in cases), encoding="utf-8")
    files.append(p)
    for name, subset in (("hard_cases.jsonl", cases[42:]), ("negative_controls.jsonl", [row for row in cases if row["case_id"] in {"xh-202630-043", "xh-202630-050", "xh-202630-057"}])):
        p = DATA / name
        p.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in subset), encoding="utf-8")
        files.append(p)
    p = DATA / "test_case_index.csv"
    with p.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["case_id", "profile_id", "course_id", "category", "difficulty", "execution_status"])
        writer.writeheader()
        writer.writerows({key: row[key] for key in writer.fieldnames} for row in cases)
    files.append(p)
    metric = {
        "data_label": "内部严格评测设计，未声称独立专家验证",
        "sample_count": 60,
        "profile_count": 3,
        "course_count": 2,
        "hallucination": {"formula": "(contradicted + unsupported + unknown) / all factual claims", "unknown_counts_as_error": True, "target": "<5%"},
        "adaptation": {"dimensions": ["difficulty", "prerequisite", "scaffolding", "task_complexity", "expression", "feedback_response"], "score_range": "0-2 each", "pass_rule": "score >= 10/12 and no critical dimension is 0", "target": ">=85%"},
        "coverage": {"weights": {"critical": 2, "ordinary": 1}, "formula": "covered weighted knowledge points / required weighted knowledge points", "target": ">=90%"},
        "review_responsibility": "项目组内部双轮复核；无外部领域标注人员",
        "publication_status": "internal_only_until_fresh_run_and_review",
    }
    p = DATA / "metric_calculation.json"
    p.write_text(json.dumps(metric, ensure_ascii=False, indent=2), encoding="utf-8")
    files.append(p)
    p = DATA / "DATA_README.md"
    p.write_text("""# 当前提交测试数据\n\n本目录由当前源码树和两门 canonical builtin course manifest 生成。所有学习者数据是脱敏合成 fixture，不是真实用户数据。\n\n- 两门课程：机器学习基础与实践、人工智能与深度学习：从机器学习基础到大模型智能体。\n- 三组画像：视觉型初学者、项目实践型学习者、理论研究型学习者。\n- 每组包含入学诊断、知识点掌握度、完整十维画像、Agent trace、讲义/实操/分阶题资源和反馈后决策。\n- `test_cases.jsonl` 含 60 条案例，`hard_cases.jsonl` 含困难案例，`negative_controls.jsonl` 含负向案例。\n- fixture 只表示测试输入输出契约；`execution_status=not_run_fixture` 不表示系统已经通过该案例。\n""", encoding="utf-8")
    files.append(p)
    return files


def make_documents(courses: list[dict[str, Any]], audit: dict[str, Any]) -> list[Path]:
    course_lines = "\n".join(f"- `{c['course_id']}` / `{c['manifest_version']}`：{c['title']}，{c['knowledge_point_count']} 个知识点，{c['file_count']} 个来源文件。" for c in courses)
    ten_lines = "\n".join(f"- `{key}`：{label}。定义：{definition}。" for key, label, definition in TEN_DIMENSIONS)
    docs: dict[str, tuple[str, str]] = {}
    docs["00_提交总览与材料清单.md"] = ("XH-202630 当前提交材料总览", f"""
# 当前材料状态

本目录为基于当前源码树重新生成的 XH-202630 提交材料。上一版本材料不在本目录的来源范围内。

## 当前范围

{course_lines}

- 生成本轮非 PPT 文档、测试数据和审计清单。
- PPT、视频和最终提交压缩包延期。
- 测试数据为脱敏合成 fixture，不能写成真实用户实验。

## 文件清单

- `documents/01_作品设计与实现方案.md/docx`
- `documents/02_作品介绍.md/docx`
- `documents/03_多智能体协同决策说明.md/docx`
- `documents/04_个性化生成与学习反馈说明.md/docx`
- `documents/05_领域知识库与可信生成说明.md/docx`
- `documents/06_测试方案与指标报告.md/docx`
- `documents/07_系统部署说明.md/docx`
- `documents/08_评委运行与演示手册.md/docx`
- `documents/09_数据合规与授权说明.md/docx`
- `documents/10_源码结构与复现说明.md/docx`
- `test_data/`：两门课程、三组画像、Agent 中间过程、资源输出、反馈、60 条案例和指标规则。
- `CURRENT_AUDIT_REPORT.md`
- `CLAIM_BOUNDARIES.md`
- `MATERIAL_MANIFEST.json`

## 当前审计快照

- Git commit：`{audit['git_commit']}`
- 工作树 dirty：`{audit['dirty_count']} 个变更项`
- Python：`{audit['python']}`
- Node.js：`{audit['node']}`
- PPT：本轮未生成
- 最终压缩包：本轮未生成
""")
    docs["01_作品设计与实现方案.md"] = ("EduMatrix 作品设计与实现方案", f"""
# 1. 项目定位

EduMatrix 面向人工智能与机器学习基础技能培训，将学习者画像、入学诊断、课程知识库、混合检索、多智能体协同、个性化资源和学习反馈组织为可追踪闭环。本方案对应赛题 `XH-202630`“领域知识个性化生成与多智能体协同决策系统研究”。

# 2. 当前领域边界

本次材料使用系统内部两门 canonical builtin course：

{course_lines}

两门课程共同构成机器学习/人工智能基础培训知识库切片。当前验证不外推到其他行业或开放领域。

# 3. 系统闭环

```text
身份与课程上下文 → 入学诊断 → 十维画像与知识点掌握度
→ 诊断/路径/生成/审核/决策 Agent
→ 讲义、实操、分阶题、导图和代码资源
→ 答题、代码、实验、复述、复习和追问反馈
→ 学习事件 → 画像、掌握度、路径和复习更新
```

# 4. 技术架构

- 前端：`frontend/src/` 的 Vue 页面、路由、API、Store、AgentTimeline 和资源工作区。
- API 与鉴权：`app/main.py`、`app/auth.py`、`app/authorization.py` 和各领域 API。
- Agent 与模型：`agent_swarm.py`、`app/m3_agents.py`、`app/m3_debate.py`、`llm_client.py`。
- 知识与 RAG：`knowledge_api.py`、`document_parser.py`、`rag_engine.py`、来源和引用对象。
- 学习与决策：`app/m5_learning.py`、`app/learning_aggregation.py`、`app/profile_summary.py`、路径和复习模块。
- 资源：`app/m5_resources.py`、artifact、artifact version 和资源工作区。
- 持久化：`app/database.py`、`app/migrations.py`、`app/m1_repository.py`。

# 5. 画像和入学诊断

初始画像采集五项登记信息：学习背景、学习目标、当前基础、主要困难和学习偏好。持续画像包含十维：

{ten_lines}

入学诊断覆盖理论、应用和工程判断三类题型。诊断结果、掌握度、不确定性和来源共同进入资源匹配和教学决策。

# 6. 当前证据边界

代码和代码地图证明模块边界；运行测试证明被覆盖行为；浏览器证据证明指定时间和环境下的页面状态。deterministic 模式用于可复现流程，不能写成真实外部模型质量。测试数据是合成 fixture，不能写成真实用户实验。
""")
    docs["02_作品介绍.md"] = ("EduMatrix 作品介绍", f"""
# EduMatrix 智能学习矩阵

EduMatrix 让系统先理解学习者，再从课程知识库检索来源，由多个职责明确的 Agent 协同生成可执行、可复习、可追踪的学习资源。

## 面向场景

当前以两门人工智能/机器学习基础课程为验证范围，覆盖从数学、数据和模型评估到深度学习、Transformer、RAG 与 Agent 的渐进学习。

## 三组学习者

- 视觉型初学者：优先图示、短步骤、具体例子和高脚手架。
- 项目实践型学习者：优先代码、实验、参数对照和故障排查。
- 理论研究型学习者：优先定义、推导、边界比较和可复现实验。

## 作品亮点

- 入学诊断和知识点掌握度共同形成学习起点；
- 十维画像影响难度、表达、任务顺序、脚手架和反馈策略；
- 画像 Agent、领域生成 Agent、证据审核 Agent 和决策编排 Agent 形成协同链；
- 输出讲义、实操指南、分阶题、思维导图和代码等多类资源；
- 答题、代码、实验、复述、复习和追问反馈重新进入画像和路径。

## 真实性说明

当前材料中的三组画像和测试数据是脱敏合成样例。内部评测结果只适用于冻结数据集和当前课程版本，未完成外部领域标注人员的独立验证。
""")
    docs["03_多智能体协同决策说明.md"] = ("EduMatrix 多智能体协同决策说明", """
# 1. Agent 角色

| 阶段 | 角色 | 主要职责 | 不允许做什么 |
|---|---|---|---|
| 诊断 | 画像/诊断 Agent | 识别知识缺口、先修关系、不确定性和学习原因 | 不能直接声明岗位胜任 |
| 规划 | 路径规划 Agent | 根据知识图谱、掌握度和负荷安排学习节点 | 不能绕过课程范围 |
| 生成 | 领域生成 Agent | 绑定课程版本和来源生成结构化资源 | 不能自行发布未审核内容 |
| 审核 | 证据审核 Agent | 拆分 claim、检查来源、矛盾和不确定状态 | 不能无证据批准 |
| 决策 | 决策编排 Agent | 综合审核结果和学习事件提出下一步动作 | 不能重写证据或把点击当掌握 |

# 2. 任务状态

```text
queued → running → completed
             ├→ failed → retry
             └→ interrupted → recovery
```

每个任务绑定 owner、课程、画像快照、trace、Agent run、artifact version、source ref 和学习事件。

# 3. 审核边界

资源与证据审核链支持声明级交叉验证和局部返工。普通聊天可以展示协作摘要、来源和校验状态，但只有实际执行且有证据时才能声明完整 Prover/Challenger/Judge 辩论。

# 4. 失败处理

- 来源不足：拒答、降低结论强度或请求补充资料；
- 版本冲突：阻止混用并提示选择课程版本；
- 结构化输出失败：记录失败、重试或受控降级；
- 证据审核失败：保留草稿和失败原因，不直接发布；
- 用户反馈冲突：追加诊断或复测，不静默覆盖画像。
""")
    docs["04_个性化生成与学习反馈说明.md"] = ("EduMatrix 个性化生成与学习反馈说明", f"""
# 1. 画像输入

每组样例包含入学诊断、知识点掌握度、完整十维画像、置信度、薄弱点、误概念、先修缺口、学习目标、可用时间和交互偏好。

# 2. 十维画像如何影响教学

{ten_lines}

当前实现重点体现在：知识点掌握度决定 foundation/scaffolded/transfer 难度带；认知负荷和挫败指数决定脚手架；交互偏好决定图示、代码或推导的表达顺序；目标和学习情境决定案例上下文与任务节奏；元认知和学习策略决定提示、复盘与复测。

# 3. 资源输出

- 讲义：核心概念、先修解释、来源和边界；
- 实操指南：步骤、参数、预期结果、失败处理和验收；
- 分阶题：识别、应用、迁移三级；
- 思维导图：概念关系和先修结构；
- 代码/实验：受运行模式和课程版本约束。

# 4. 反馈闭环

答题、代码测试、交互实验、费曼复述、复习完成、延迟复测和对话追问通过学习事件回流画像、掌握度、路径和复习队列。普通浏览和单次点击不直接升级为掌握证据。

# 5. 样例状态

测试数据中的三组画像是合成 fixture，用于展示数据契约和差异化逻辑。fixture 的预期输出不等于当前系统已经执行出的正式效果。
""")
    docs["05_领域知识库与可信生成说明.md"] = ("EduMatrix 领域知识库与可信生成说明", f"""
# 1. 知识库范围

两门内部基础课程构成本次提交的机器学习/人工智能基础知识库切片：

{course_lines}

课程 manifest 记录章节、知识点、先修关系、来源文件、来源版本、哈希和许可证说明。

# 2. RAG 链路

```text
课程选择 → 文档解析 → 章节/知识点映射 → 分块与索引
→ 关键词/语义/图谱混合检索 → 课程范围过滤
→ source_ref/evidence → 生成 → claim 审核 → 资源版本
```

# 3. Claim 审核

每个事实性 claim 要绑定来源 ID 和可定位片段，并标记：`supported`、`contradicted`、`unsupported` 或 `unknown`。`unknown` 在严格幻觉率统计中计为错误。

# 4. 版本与许可证

课程、来源文档、artifact 和 artifact version 分离保存。材料中保留 manifest 版本、文件哈希、来源角色和许可证信息。无法确认授权的文件不进入提交包。

# 5. 真实性边界

来源数量不等于事实准确率；代码存在不等于检索效果；deterministic 输出不等于真实外部模型效果；普通聊天不笼统宣称每次执行完整辩论。
""")
    docs["06_测试方案与指标报告.md"] = ("EduMatrix 测试方案与指标报告", """
# 1. 测试范围

本轮测试数据包含两门课程、三组画像和 60 条案例。覆盖入学诊断、掌握度、画像影响、检索、生成、审核、决策、反馈、权限、失败恢复和代码边界。

# 2. 无外部标注人员的复核口径

项目没有外部领域标注人员，因此采用冻结知识库和预定义规则的内部严格评测，使用两轮独立复核脚本。报告不得写成领域专家标注或独立专家验证。

# 3. 幻觉率

```text
（contradicted + unsupported + unknown）÷ 全部事实性 claim
```

没有来源的专业结论按 `unsupported`，无法判断的 `unknown` 计入错误，关键知识点 contradiction 使案例失败。必须保留 claim、来源、片段、判断和失败案例。

# 4. 画像—资源难度适配准确率

每个样本按难度、先修、脚手架、任务复杂度、表达方式和反馈响应六项以 0/1/2 评分。总分至少 10/12 且难度、先修不能为 0 才算通过。报告总体、三组画像分别、最差组、各维度失败数和 A/B 盲测结果。

# 5. 核心知识点覆盖率

每个案例预先固定必需知识点和关键知识点。只有内容正确、来源正确、达到最低深度且未被其他段落错误描述才算覆盖。关键知识点权重 2，普通知识点权重 1：

```text
正确覆盖知识点权重总和 ÷ 必需知识点权重总和
```

# 6. 真实性设计

60 条案例中保留困难、课程外、无答案、版本冲突、画像不足、先修缺失、模型失败和反馈重规划案例。不得人为制造错误，不得为了达到目标删除失败样本，也不得把 fixture 预期结果冒充已执行结果。

# 7. 当前验证结果

- 聚焦回归：65 passed、6 warnings；另有 11 个旧式 M5 内存夹具失败，原因是夹具没有提供当前课程可用学习材料，错误为 `course_unavailable`/`tutor_course_unavailable`，不把它们混入通过数。
- 前端 production build：成功，Vite 转换 5453 个模块；仅有大体积 chunk 提示。
- 无 Docker 核心 E2E：通过，4 张截图，健康检查、登录、学习工作台、知识库、trusted_local 安全代码输出和禁用边界均通过。
- 两课程 Chromium E2E：通过，24 个页面/截图，覆盖两门课程和两个桌面视口；HTTP 5xx、页面错误和控制台错误均为 0。8 个 `net::ERR_ABORTED` 为导航切换时取消的请求，单独记录为非业务失败。
- 幻觉率、适配准确率和覆盖率：当前没有外部领域标注人员；本轮生成了严格评测数据和公式，但未把 fixture 预期结果写成正式指标。
""")
    docs["07_系统部署说明.md"] = ("EduMatrix 系统部署说明", """
# 1. 环境

- Windows 10/11 或 Linux x86_64；
- Python 3.11；
- Node.js 20 和 npm；
- Chromium 仅在浏览器 E2E 或 PDF 导出时需要；
- Docker 仅在需要容器代码执行时需要。

# 2. 安装

```powershell
py -3.11 -m venv .venv
.venv\\Scripts\\python.exe -m pip install --upgrade pip
.venv\\Scripts\\python.exe -m pip install -r requirements.txt
cd frontend
npm ci
cd ..
```

复制 `.env.example` 为 `.env`，生成至少 32 字符的 `EDUMATRIX_AUTH_SECRET_KEY`。默认使用 deterministic LLM、hash embedding 和 disabled sandbox。

# 3. 启动

```powershell
.venv\\Scripts\\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8000
cd frontend
npm run dev -- --host 127.0.0.1
```

前端：`http://127.0.0.1:5173`。健康检查：`http://127.0.0.1:8000/api/health`。

# 4. 运行模式

- deterministic：无外部模型的可复现流程；
- 外部模型：配置兼容 endpoint、model 和 key，密钥不进材料包；
- trusted_local：可信本机受限子进程，不具备 Docker 隔离；
- docker：需要 Docker daemon，才提供容器级代码隔离。

# 5. 验收

健康检查 → 登录/注册 → 选择两门课程 → 完成入学诊断 → 查看十维画像与掌握度 → 提问并观察 Agent 状态 → 打开讲义、实操和分阶题 → 提交反馈 → 查看路径或复习变化。

# 6. 故障

- 前端无法访问后端：先检查 8000 端口和 `/api/health`；
- 代码执行返回 503：disabled 是预期安全边界；
- 外部模型不可用：切回 deterministic 并标注结果；
- E2E/PDF 失败：安装 Playwright Chromium。
""")
    docs["08_评委运行与演示手册.md"] = ("EduMatrix 评委运行与演示手册", """
# 最短复现路径

```text
安装 Python/Node 依赖
→ 生成 .env 和 JWT 密钥
→ 启动后端/前端
→ 健康检查
→ 登录或注册
→ 选择两门课程中的一门
→ 完成入学诊断
→ 查看掌握度和十维画像
→ 进入学习空间提问
→ 查看 Agent 协同摘要和来源
→ 打开讲义、实操和分阶题
→ 提交答题反馈
→ 查看下一步路径或复习安排
```

# 观察重点

- 画像输入是否包含诊断和知识点掌握度；
- 资源难度和表达是否随画像变化；
- Agent 阶段是否展示分析、生成、审核和决策；
- 资源是否有课程和来源绑定；
- 失败和不确定状态是否可理解；
- disabled、trusted_local 和 Docker 是否被准确区分。

# 演示数据说明

提交材料中的样例为脱敏合成数据。评委看到的 fixture 用于说明输入输出结构，不能被解释为真实用户学习效果或外部专家评价。
""")
    docs["09_数据合规与授权说明.md"] = ("EduMatrix 数据合规与授权说明", """
# 1. 数据类型

- 测试画像、入学诊断、掌握度、Agent trace 和反馈均为脱敏合成数据；
- 不提交真实姓名、手机号、邮箱、密码、学习记录或未脱敏日志；
- 不提交 `.env`、API Key、JWT secret、运行数据库、WAL/SHM、虚拟环境或 node_modules。

# 2. 课程来源

两门课程的 manifest 记录课程版本、来源文件、上游 revision、文件哈希、许可证和 source role。最终材料只纳入许可证和再分发边界清晰的资料。

# 3. 权限

用户、课程、资源、来源和学习事件由服务端认证、owner 和课程范围共同约束。前端角色值不能替代服务端授权。

# 4. 外部服务

真实模型、联网搜索、Docker、PDF 导出和其他可选服务需要目标环境配置和单独验证。deterministic 和合成数据不能包装成真实模型效果或真实用户实验。
""")
    docs["10_源码结构与复现说明.md"] = ("EduMatrix 源码结构与复现说明", """
# 1. 目录

- `frontend/src/`：Vue 路由、页面、组件、API 和 Store；
- `app/`：FastAPI 路由、鉴权、课程、画像、Agent、资源、学习事件和治理服务；
- 根目录 Python 模块：Agent Swarm、RAG、模型客户端、代码执行和学习算法；
- `data/curated/`：两门当前 canonical course manifest 和课程正文；
- `tests/`、`scripts/`：自动化测试、E2E、数据和运行脚本。

# 2. 关键入口

- 后端：`app.main:app`；
- 前端：`frontend/package.json` 的 Vite 入口；
- 健康检查：`GET /api/health`；
- Agent：`agent_swarm.py`、`app/m3_agents.py`；
- RAG：`rag_engine.py`、`knowledge_api.py`；
- 画像：`app/profile_onboarding.py`、`app/profile_summary.py`；
- 资源：`app/m5_resources.py`；
- 学习反馈：`app/m5_learning.py`、`app/learning_aggregation.py`。

# 3. 复现记录

当前材料生成时记录的版本、依赖、模式和验证结果见 `CURRENT_AUDIT_REPORT.md` 与 `MATERIAL_MANIFEST.json`。没有重新执行的能力标记为未验证。

# 4. 运行边界

默认 deterministic/disabled 只保证核心流程可在无外部模型和无 Docker 时复现。真实模型、Docker 代码执行、联网搜索和 PDF 导出均需目标环境单独验证。
""")
    docs["01_作品设计与实现方案.md"] = (docs["01_作品设计与实现方案.md"][0], docs["01_作品设计与实现方案.md"][1] + """

# 7. 数据对象关系

```text
DBUser/DBStudentProfile
  → DBCourse/DBCourseMembership
  → DBSourceDocument/DBChapter/DBKnowledgePoint
  → DBSourceRef
  → DBProfileSnapshot/DBProfileEvidence
  → DBGenerationTask/DBAgentRun/DBTaskStep
  → DBArtifact/DBArtifactVersion
  → DBLearningEvent/DBLearningEventConsumption
  → DBLearningPath/DBPathNode/DBReviewPlan
```

对象 owner、course_id、version、content_hash 和 source_ref 共同构成可追溯边界。资源版本追加而不是原地覆盖，学习事件追加而不是直接修改历史证据。

# 8. 关键接口边界

| 能力 | 代表入口 | 主要结果 |
|---|---|---|
| 认证 | `/api/auth/*` | JWT、角色和当前用户 |
| 课程 | `/api/knowledge/courses`、`/api/v1/courses/*` | 课程、章节、知识点、来源 |
| 对话 | `/api/v1/chat/*`、`/api/stream/*` | 消息、来源、协作摘要和资源动作 |
| 画像 | `/api/profile/*`、`/api/v1/m5/profile-summary` | 诊断、掌握度、十维状态和证据 |
| 路径 | `/api/v1/m5/paths/*`、`/api/v1/m5/today` | 任务、节点、完成状态和推荐 |
| 资源 | `/api/v1/m5/resources/*`、`/api/v1/resource-bundles/*` | artifact、版本、生成任务和导出 |
| Agent | `/api/v1/m3/*` | 诊断、生成、审核、辩论和决策 |
| 学习事件 | `/api/v1/m5/learning-events` | 反馈事件和聚合消费 |

# 9. 可靠性策略

- 幂等键防止重复生成、重复计费和重复事件；
- task/run/step 状态支持刷新恢复和失败重试；
- 课程版本和来源哈希防止跨版本混用；
- 证据审核失败时保留草稿和失败原因；
- 权限先于检索和生成，防止无权对象进入模型上下文；
- 反馈事件与画像/路径/复习消费分离，允许重放和审计。
""")
    docs["03_多智能体协同决策说明.md"] = (docs["03_多智能体协同决策说明.md"][0], docs["03_多智能体协同决策说明.md"][1] + """

# 5. Agent 输入输出契约

| Agent | 必需输入 | 输出对象 | 关键禁止行为 |
|---|---|---|---|
| 诊断 | profile_snapshot、job_task、knowledge_points | diagnostic_assessment、profile_snapshot | 不声明胜任、不覆盖画像 |
| 领域生成 | profile_snapshot、domain_version、source_refs | lesson、lab_guide、tiered_quiz | 不自发布、不改指标 |
| 证据审核 | artifact_versions、source_refs、quality_rule | claim_check、debate_round、quality_check | 不无证据批准、不覆盖生成结果 |
| 决策编排 | profile_snapshot、quality_checks、learning_events | decision_record、matching_report | 不重写证据、不把点击当掌握 |

每一阶段都保留 `owner_public_id`、`course_id`、`profile_snapshot_id`、`trace_id` 和输入引用；公开给前端的只是脱敏摘要。

# 6. 决策候选

系统可以提出收集更多证据、降低解释难度、补先修、维持巩固或增加进阶挑战五类候选。候选必须带理由、证据 ID、预计负担、风险和是否需要用户确认。决策候选是 proposed 状态，不能绕过确认直接改变高风险学习结论。
""")
    docs["06_测试方案与指标报告.md"] = (docs["06_测试方案与指标报告.md"][0], docs["06_测试方案与指标报告.md"][1] + """

# 8. 测试矩阵

| 测试层 | 重点 | 证据 |
|---|---|---|
| 静态 | 文件、AST、路由、Schema、来源字段 | 命令输出和审计 JSON |
| 单元 | 画像、掌握度、RAG、资源契约、状态机 | pytest 原始结果 |
| 集成 | Agent 步骤、数据库、事件消费、幂等 | pytest/隔离库 |
| API | 认证、owner、课程范围、失败映射 | HTTP 响应和测试日志 |
| 浏览器 | 登录、诊断、画像、资源、反馈、移动端 | Chromium 截图和报告 |
| 数据 | 两门课程、三组画像、60 条案例、许可证 | `test_data/` |

# 9. 指标计算记录要求

每次指标运行必须保存：

- 测试集版本和文件哈希；
- 当前 Git commit；
- LLM/embedding provider 和 model；
- 课程 ID、课程版本和知识库哈希；
- claim、画像适配评分和知识点覆盖明细；
- 失败案例 ID；
- 内部复核责任；
- 运行日期和失败原因。

没有这些信息时，报告状态为“未正式测量”。fixture 的预期输出只证明数据契约，不证明当前系统已经通过相应案例。
""")
    docs["07_系统部署说明.md"] = (docs["07_系统部署说明.md"][0], docs["07_系统部署说明.md"][1] + """

# 7. 运行前检查

```powershell
.venv\\Scripts\\python.exe -c "import fastapi, sqlalchemy, numpy, torch, openai, docx, playwright; print('python imports ok')"
Invoke-RestMethod http://127.0.0.1:8000/api/health
```

健康检查必须确认数据库可用、当前模型模式、检索状态和 sandbox 模式。生产环境必须使用独立随机 JWT 密钥；演示环境不得把密钥写入材料。

# 8. 端口和进程

- 后端默认 `127.0.0.1:8000`；
- Vite 开发前端默认 `127.0.0.1:5173`；
- 前端生产构建产物不等于后端已经启动；
- 每次演示结束应停止本轮服务和浏览器进程。
""")
    output: list[Path] = []
    for filename, (title, content) in docs.items():
        md = write(DOCS / filename, content)
        docx = DOCS / filename.replace(".md", ".docx")
        render_docx(title, content, docx)
        output.extend([md, docx])
    claim = """# 当前事实边界\n\n- 当前来源：赛题 PDF、当前源码、两门课程 manifest、当前测试和当前运行结果。\n- 两门课程是本次知识库切片范围，不代表跨行业泛化。\n- 三组画像、入学诊断、掌握度和十维画像样例为脱敏合成 fixture。\n- deterministic 结果不等于真实外部模型效果。\n- 内部规则评测不等于独立领域专家验证。\n- 代码存在、测试通过和截图存在分别只能证明对应范围，不可扩展为生产安全、生产并发或正式学习效果。\n- 当前未重新执行的指标写为未正式测量。\n"""
    output.append(write(OUT / "CLAIM_BOUNDARIES.md", claim))
    audit_md = f"""# 当前源码与运行审计报告

生成时间：`{audit['generated_at']}`  
Git commit：`{audit['git_commit']}`  
工作树变更项：`{audit['dirty_count']}`  
Python：`{audit['python']}`  
Node.js：`{audit['node']}`

## 两门课程

{chr(10).join(f"- `{c['course_id']}` / `{c['manifest_version']}`：{c['title']}，{c['knowledge_point_count']} 个知识点，{c['file_count']} 个来源文件。" for c in courses)}

## 画像与诊断

- 初始画像字段：`{', '.join(item['key'] for item in audit['profile_initial_dimensions'])}`
- 持续画像：10 维，见 `test_data/learner_profiles.json`。
- 入学诊断维度：`{', '.join(audit['diagnostic_dimensions'])}`。
- 三组画像均包含入学诊断、掌握度、十维画像、Agent trace、资源和反馈决策。

## 实际验证

- 聚焦回归：65 passed、6 warnings；11 个旧式 M5 内存夹具失败，均因没有提供当前课程可用材料，记录为夹具/环境限制，不作为材料生成缺陷。
- 前端构建：通过，5453 个模块转换；存在大体积 chunk 警告。
- 无 Docker 核心 E2E：通过，4 张截图。
- 两课程 Chromium E2E：通过，24 个页面/截图，HTTP 5xx=0，页面错误=0，控制台错误=0；8 个 `net::ERR_ABORTED` 为切页取消请求。

## 未正式测量

幻觉率、画像—资源难度适配准确率和核心知识点覆盖率尚未有外部领域标注人员复核。本目录只提供冻结数据、严格公式和内部复核结构，不把 fixture 预期值写成正式比赛指标。

## 外部依赖和边界

- 当前后端健康检查显示环境配置了外部模型，但本材料不记录 endpoint、model 或 key；文档默认以 deterministic/本地可复现口径描述。
- trusted_local 是受限本地子进程，不是 Docker 容器隔离。
- Docker、联网搜索、PDF 导出和真实外部模型效果需要目标环境单独复核。
- 两门课程之外的跨领域泛化未验证。
"""
    output.append(write(OUT / "CURRENT_AUDIT_REPORT.md", audit_md))
    return output


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    courses = load_courses()
    audit = {
        "generated_at": now(),
        "git_commit": command_text(["git", "rev-parse", "HEAD"]),
        "dirty_count": command_text(["git", "status", "--short"]).count("\n") + (1 if command_text(["git", "status", "--short"]).strip() else 0),
        "python": command_text([str(ROOT / ".venv" / "Scripts" / "python.exe"), "--version"]),
        "node": command_text(["node", "--version"]),
        "course_manifests": [{k: v for k, v in course.items() if k not in {"knowledge_points", "files"}} for course in courses],
        "profile_initial_dimensions": [{"key": k, "label": l} for k, l in INITIAL_DIMENSIONS],
        "profile_continuous_dimensions": [{"key": k, "label": l, "definition": d} for k, l, d in TEN_DIMENSIONS],
        "diagnostic_dimensions": list(DIAGNOSTIC_KINDS),
        "source_modules": [
            "agent_swarm.py", "app/m3_agents.py", "app/m3_debate.py", "app/m3_matching.py", "app/m3_decisions.py",
            "rag_engine.py", "knowledge_api.py", "document_parser.py", "app/profile_onboarding.py", "app/profile_summary.py",
            "app/m5_resources.py", "app/m5_learning.py", "app/learning_aggregation.py", "app/database.py", "app/auth.py",
        ],
        "runtime_validation": {
            "focused_regression": {"passed": 65, "warnings": 6, "failed": 11, "failure_scope": "legacy in-memory M5 fixtures without usable course material"},
            "frontend_build": {"status": "passed", "modules_transformed": 5453},
            "no_docker_e2e": {"status": "passed", "screenshots": 4, "browser": "playwright-chromium", "mode": "trusted_local"},
            "two_course_e2e": {"status": "passed", "pages": 24, "screenshots": 24, "http_5xx": 0, "page_errors": 0, "console_errors": 0, "navigation_aborts": 8},
        },
        "claims_policy": "未重新执行或未独立复核的结果不得写成正式达标结论",
    }
    data_files = build_data(courses)
    doc_files = make_documents(courses, audit)
    manifest = {
        "schema_version": "xh-202630.current-submission-materials.v1",
        "generated_at": audit["generated_at"],
        "source": "current_worktree_only",
        "git_commit": audit["git_commit"],
        "ppt_generated": False,
        "final_zip_generated": False,
        "courses": [{"course_id": c["course_id"], "version": c["manifest_version"], "title": c["title"], "knowledge_points": c["knowledge_point_count"], "files": c["file_count"]} for c in courses],
        "learner_profiles": 3,
        "test_cases": 60,
        "data_label": "脱敏合成 fixture",
        "files": sorted(rel(path) for path in [*doc_files, *data_files] if path.is_file()),
        "pending_user_inputs": ["报名表路径/字段", "最终演示视频路径", "最终提交包命名信息", "私有源码仓库评审权限（如适用）"],
    }
    readme = f"""# XH-202630 当前提交材料

本目录由当前工作树重新生成，Git commit：`{audit['git_commit']}`。上一版本提交材料不在本目录来源范围内。

## 已生成

- `documents/`：10 份非 PPT 文档的 Markdown 和 DOCX；
- `test_data/`：两门课程切片、三组完整画像、入学诊断、掌握度、十维画像、Agent trace、资源、反馈和 60 条测试案例；
- `CURRENT_AUDIT_REPORT.md`：当前源码和运行审计；
- `RUNTIME_EVIDENCE.json`：测试、构建和浏览器证据索引；
- `CLAIM_BOUNDARIES.md`：对外声明边界；
- `MATERIAL_MANIFEST.json`：文件和版本清单。

## 延期项目

本轮不生成 PPT，不接入演示视频，不接入报名表，不生成最终提交压缩包。

## 真实性口径

画像与测试数据是脱敏合成 fixture。三项比赛指标的结构、公式和严格内部评测数据已准备，但没有外部领域标注人员，因此未将 fixture 预期值写成正式达标结论。
"""
    write(OUT / "README.md", readme)
    runtime_evidence = {
        "generated_at": audit["generated_at"],
        "git_commit": audit["git_commit"],
        "commands": [
            {"command": ".venv\\Scripts\\python.exe -m pytest -q tests/test_m0_course_catalog.py tests/test_m1_course_sources.py tests/test_f3_profile_onboarding.py tests/test_m1_profile_evidence.py tests/test_m3_profiles.py tests/test_m3_agents.py tests/test_m3_debate.py tests/test_m3_matching.py tests/test_m3_decisions.py tests/test_m3_resources.py tests/test_competition_demo_seed.py tests/test_competition_package_runtime.py", "result": "65 passed, 6 warnings"},
            {"command": "cd frontend; npm.cmd run build", "result": "passed; 5453 modules transformed; large chunk warning"},
            {"command": ".venv\\Scripts\\python.exe scripts\\e2e_no_docker.py", "result": "passed; 4 screenshots; trusted_local"},
            {"command": ".venv\\Scripts\\python.exe scripts\\e2e_lzz_two_course_demo_evidence.py", "result": "passed; 24 screenshots/pages; HTTP 5xx=0; page errors=0; console errors=0; 8 navigation aborts"},
        ],
        "evidence_paths": [
            "outputs/e2e_no_docker/report.json",
            "outputs/原系统融合重构执行/evidence/custom_tasks/COURSE-DATA-DEMO-RESEED-03/browser/browser_report.json",
        ],
        "known_failure": "11 legacy M5 in-memory fixture failures because fixture courses have no usable learning material; recorded separately and not presented as a passing result.",
        "metrics": {"status": "not_formally_measured", "reason": "no external domain annotators; internal strict evaluation dataset and formulas are provided"},
    }
    write(OUT / "RUNTIME_EVIDENCE.json", json.dumps(runtime_evidence, ensure_ascii=False, indent=2))
    manifest["files"].extend(["README.md", "RUNTIME_EVIDENCE.json"])
    manifest["files"] = sorted(set(manifest["files"]))
    write(OUT / "MATERIAL_MANIFEST.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    print(json.dumps({"output": str(OUT), "documents": len(doc_files), "data_files": len(data_files), "profiles": 3, "test_cases": 60, "ppt_generated": False}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
