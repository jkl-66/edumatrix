from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base, DBCourse, DBLearningEvent, DBSourceDocument, DBSourceRef
from app.m1_repository import create_artifact_with_version
from app.m5_capabilities import M5CapabilityError
from app.m5_resources import (
    ResourceVersionCreateRequest,
    create_resource_version,
    parse_deck_content,
    parse_video_recommendation_content,
    render_deck_export,
    resource_detail,
)
from app.resource_bundles import _deterministic_demo_content


ROOT = Path(__file__).resolve().parents[1]


def _session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _source(session, *, course_id: str = "course-media"):
    course = DBCourse(id=course_id, version="1", title="机器学习工程", manifest_version="media-v1")
    document = DBSourceDocument(
        course_id=course.id,
        title="优化算法课程讲义",
        filename="optimization.md",
        file_format="markdown",
        content_hash="a" * 64,
        source_sha256="b" * 64,
        processed_sha256="c" * 64,
        source_url="https://example.edu/course/optimization",
        license_text="课程演示材料，允许在本课程中引用。",
        processing_status="ready",
    )
    session.add_all([course, document])
    session.flush()
    source_ref = DBSourceRef(
        source_document_id=document.id,
        parser_version="media-parser-v1",
        location_kind="section",
        location_label="第 2 章 梯度下降",
        content_hash="d" * 64,
    )
    session.add(source_ref)
    session.flush()
    return course, source_ref


def _video_content(source_ref_id: str) -> dict:
    return {
        "schemaVersion": "f5.video-recommendation.v1",
        "title": "梯度下降视频推荐",
        "summary": "优先展示来源和许可边界明确的视频入口。",
        "learningGoal": "用视频补充理解梯度下降的更新机制。",
        "adaptationSummary": "当前画像需要先看机制概览。",
        "recommendations": [
            {
                "id": "gradient-course",
                "title": "公开课程视频页",
                "sourceKind": "external_url",
                "availability": "external_only",
                "url": "https://ml-course.github.io/master/",
                "localRef": None,
                "durationSeconds": None,
                "author": "ML Course 页面维护者",
                "language": "English",
                "licenseStatus": "reference_only_unverified",
                "licenseLabel": "仅提供来源页链接，不下载、不镜像、不分发视频。",
                "licenseUrl": "https://ml-course.github.io/master/",
                "reason": "与当前知识点及课程实训链路相关。",
                "segments": [],
                "sourceRefIds": [source_ref_id],
                "alternatives": [
                    {
                        "id": "alternative-vu",
                        "title": "VU Machine Learning",
                        "url": "https://mlvu.github.io/",
                        "reason": "原来源不可访问时使用另一条已登记课程来源。",
                    }
                ],
            }
        ],
    }


def _deck_content(source_ref_id: str, title: str = "梯度下降课件") -> dict:
    return {
        "schemaVersion": "f5.deck.v1",
        "title": title,
        "summary": "从问题、机制到下一步学习的结构化课件。",
        "learningGoal": "解释梯度下降并辨认适用边界。",
        "adaptationSummary": "按当前画像减少公式密度，保留来源和实践入口。",
        "slides": [
            {
                "id": "title",
                "layout": "title",
                "title": "梯度下降",
                "subtitle": "从方向到步长",
                "body": "",
                "bullets": [],
                "speakerNotes": "冻结合成教学示例",
                "sourceRefIds": [source_ref_id],
            },
            {
                "id": "mechanism",
                "layout": "content",
                "title": "核心机制",
                "subtitle": "沿负梯度方向更新",
                "body": "梯度给出局部最陡上升方向，因此参数沿相反方向更新。",
                "bullets": ["确定方向", "控制步长", "检查收敛"],
                "speakerNotes": "结合课程来源解释更新式。",
                "sourceRefIds": [source_ref_id],
            },
        ],
    }


def _artifacts(session):
    course, source_ref = _source(session)
    video, video_version = create_artifact_with_version(
        session,
        owner_public_id="usr-media",
        course_id=course.id,
        artifact_type="video",
        title="梯度下降视频推荐",
        content=json.dumps(_video_content(source_ref.id), ensure_ascii=False),
        content_format="json",
        metadata={"resource_type": "video", "source_ref_ids": [source_ref.id]},
    )
    deck, deck_version = create_artifact_with_version(
        session,
        owner_public_id="usr-media",
        course_id=course.id,
        artifact_type="deck",
        title="梯度下降课件",
        content=json.dumps(_deck_content(source_ref.id), ensure_ascii=False),
        content_format="json",
        metadata={"resource_type": "deck", "source_ref_ids": [source_ref.id]},
    )
    session.commit()
    return video, video_version, deck, deck_version, source_ref


def test_video_and_deck_schemas_reject_ambiguous_or_duplicate_structure():
    video = _video_content("ref-media")
    parsed = parse_video_recommendation_content(json.dumps(video, ensure_ascii=False))
    assert parsed.recommendations[0].availability == "external_only"

    duplicate = _video_content("ref-media")
    duplicate["recommendations"].append(dict(duplicate["recommendations"][0]))
    with pytest.raises(M5CapabilityError) as exc_info:
        parse_video_recommendation_content(json.dumps(duplicate, ensure_ascii=False))
    assert exc_info.value.code == "video_content_incomplete"

    fake_playback = _video_content("ref-media")
    fake_playback["recommendations"][0]["availability"] = "playable"
    with pytest.raises(M5CapabilityError):
        parse_video_recommendation_content(json.dumps(fake_playback, ensure_ascii=False))

    deck = _deck_content("ref-media")
    deck["slides"][1]["id"] = deck["slides"][0]["id"]
    with pytest.raises(M5CapabilityError) as deck_error:
        parse_deck_content(json.dumps(deck, ensure_ascii=False))
    assert deck_error.value.code == "deck_content_incomplete"


def test_resource_detail_projects_traceable_video_and_artifact_deck_without_writes():
    session = _session()
    video, _, deck, _, source_ref = _artifacts(session)
    before_events = session.query(DBLearningEvent).count()

    video_detail = resource_detail(session, video)
    deck_detail = resource_detail(session, deck)

    recommendation = video_detail["video"]["current"]["structured"]["recommendations"][0]
    assert recommendation["playback"]["external_only"] is True
    assert recommendation["licenseStatus"] == "reference_only_unverified"
    assert recommendation["sourceRefIds"] == [source_ref.id]
    assert video_detail["video"]["current"]["source_refs"][0]["license"]
    assert deck_detail["deck"]["current"]["outline"][1]["title"] == "核心机制"
    assert deck_detail["deck"]["current"]["export_capabilities"] == {"html": True, "pptx": True}
    assert session.query(DBLearningEvent).count() == before_events


def test_invalid_and_empty_video_states_are_honest():
    session = _session()
    course, source_ref = _source(session)
    empty = _video_content(source_ref.id)
    empty["recommendations"] = []
    empty_artifact, _ = create_artifact_with_version(
        session,
        owner_public_id="usr-media",
        course_id=course.id,
        artifact_type="video",
        title="暂无可用视频",
        content=json.dumps(empty, ensure_ascii=False),
        content_format="json",
        metadata={"resource_type": "video", "source_ref_ids": [source_ref.id]},
    )
    invalid, _ = create_artifact_with_version(
        session,
        owner_public_id="usr-media",
        course_id=course.id,
        artifact_type="video",
        title="损坏视频推荐",
        content=json.dumps({"schemaVersion": "f5.video-recommendation.v1", "title": "损坏"}),
        content_format="json",
        metadata={"resource_type": "video", "source_ref_ids": [source_ref.id]},
    )
    session.commit()

    assert resource_detail(session, empty_artifact)["video"]["current"]["structured"]["recommendations"] == []
    projection = resource_detail(session, invalid)["video"]["current"]
    assert projection["mode"] == "invalid_structured"
    assert projection["warnings"] == ["当前视频推荐无法解析，请恢复有效版本或重新生成该资源。"]


def test_video_edit_creates_new_version_and_failed_deck_save_preserves_current_version():
    session = _session()
    video, video_version, deck, deck_version, source_ref = _artifacts(session)
    revised_video = _video_content(source_ref.id)
    revised_video["summary"] = "第二版保留来源和许可说明，并调整推荐理由。"
    video_detail = create_resource_version(
        session,
        artifact=video,
        actor_public_id="usr-media",
        request=ResourceVersionCreateRequest(
            expected_lock_version=video.lock_version,
            source_version_id=video_version.id,
            content=json.dumps(revised_video, ensure_ascii=False),
            content_format="json",
        ),
    )
    assert len(video_detail["video"]["versions"]) == 2
    assert video_detail["video"]["current"]["structured"]["summary"] == revised_video["summary"]
    assert video_detail["history"]["items"][0]["revised_from_version_id"] == video_version.id

    current_id = deck.current_version_id
    invalid = _deck_content(source_ref.id)
    invalid["slides"][1]["body"] = ""
    invalid["slides"][1]["bullets"] = []
    with pytest.raises(M5CapabilityError):
        create_resource_version(
            session,
            artifact=deck,
            actor_public_id="usr-media",
            request=ResourceVersionCreateRequest(
                expected_lock_version=deck.lock_version,
                source_version_id=deck_version.id,
                content=json.dumps(invalid, ensure_ascii=False),
                content_format="json",
            ),
        )
    assert deck.current_version_id == current_id
    assert resource_detail(session, deck)["deck"]["current"]["structured"]["title"] == "梯度下降课件"


def test_deck_edit_creates_new_artifact_version_and_exports_real_files():
    session = _session()
    _, _, deck, first, source_ref = _artifacts(session)
    revised = _deck_content(source_ref.id, title="梯度下降课件修订版")
    detail = create_resource_version(
        session,
        artifact=deck,
        actor_public_id="usr-media",
        request=ResourceVersionCreateRequest(
            expected_lock_version=deck.lock_version,
            source_version_id=first.id,
            content=json.dumps(revised, ensure_ascii=False),
            content_format="json",
        ),
    )
    assert len(detail["deck"]["versions"]) == 2
    assert detail["deck"]["current"]["structured"]["title"] == "梯度下降课件修订版"
    assert detail["deck"]["versions"][0]["projection"]["structured"]["title"] == "梯度下降课件"

    html_bytes, html_name, html_type = render_deck_export(
        session, artifact=deck, version_id=deck.current_version_id, format_name="html",
    )
    pptx_bytes, pptx_name, pptx_type = render_deck_export(
        session, artifact=deck, version_id=deck.current_version_id, format_name="pptx",
    )
    assert html_type.startswith("text/html") and html_name.endswith(".html")
    assert b"<!doctype html>" in html_bytes
    assert pptx_name.endswith(".pptx") and pptx_type.endswith("presentationml.presentation")
    assert pptx_bytes[:2] == b"PK"
    presentation = Presentation(io.BytesIO(pptx_bytes))
    assert len(presentation.slides) == 2
    assert "梯度下降" in presentation.slides[0].shapes[0].text
    assert "冻结合成教学示例" in presentation.slides[0].notes_slide.notes_text_frame.text


def test_deterministic_generation_uses_legal_video_catalog_and_structured_deck():
    manifest = {
        "knowledge_point_id": "kp-gradient",
        "context": {
            "knowledge_point_title": "梯度下降",
            "learning_goal": "解释并应用梯度下降",
            "profile_adaptation": "先看机制，再进入实践",
        },
    }
    sources = [{"source_ref_id": "ref-demo", "title": "课程讲义", "location_label": "第 2 章"}]
    video = parse_video_recommendation_content(_deterministic_demo_content(
        resource_type="video", item={"title": "梯度下降视频推荐"}, manifest=manifest,
        source_snapshots=sources,
    ))
    deck = parse_deck_content(_deterministic_demo_content(
        resource_type="deck", item={"title": "梯度下降课件"}, manifest=manifest,
        source_snapshots=sources,
    ))
    assert video.recommendations
    assert all(item.sourceKind == "external_url" for item in video.recommendations)
    assert all(item.availability == "external_only" for item in video.recommendations)
    assert len(deck.slides) == 5
    assert deck.slides[-1].title == "下一步学习"


def test_frontend_contract_uses_dedicated_artifact_workspaces_and_no_generation_ui():
    router = (ROOT / "frontend/src/router/index.js").read_text(encoding="utf-8")
    video = (ROOT / "frontend/src/views/VideoRecommendationWorkspace.vue").read_text(encoding="utf-8")
    deck = (ROOT / "frontend/src/views/WebDeckWorkspace.vue").read_text(encoding="utf-8")

    assert "name: 'VideoResource', component: VideoRecommendationWorkspace" in router
    assert "name: 'DeckResource', component: WebDeckWorkspace" in router
    for label in ("推荐理由", "来源与许可", "替代来源", "写学习笔记", "打开来源"):
        assert label in video
    for forbidden in ("生成视频", "视频练习", "创建实验", "createM5LearningEvent", "unlockExam"):
        assert forbidden not in video
    for label in ("课件大纲", "课件预览", "来源与版本", "保存为新版本", "导出 HTML", "导出 PPTX"):
        assert label in deck
    assert "getM5Resource" in deck
    assert "listWebDecks" not in deck
    assert "createM5LearningEvent" not in deck
