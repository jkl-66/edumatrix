from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import text
from sqlalchemy.exc import DatabaseError

from app.database import DBWebDeckVersion, SessionLocal
from app.main import app
from app.m1_object_service import OBJECT_SPECS
from tests.api_test_helpers import auth_headers


SLIDES = [
    {
        "layout": "title", "title": "梯度下降技能微课", "subtitle": "参数、轨迹与证据",
        "body": "", "bullets": ["服务端确定性计算", "浏览器交互播放"],
        "speaker_notes": "先说明这是冻结合成教学示例。", "accent": "#355B42",
    },
    {
        "layout": "content", "title": "学习率如何影响收敛", "subtitle": "稳定条件",
        "body": "学习率和曲率共同决定每次参数更新的步长。",
        "bullets": ["0 < learning_rate * 2 * curvature < 2", "观察损失是否持续下降"],
        "speaker_notes": "演示不同学习率版本，但不宣称适用于全部优化器。", "accent": "#C65A3C",
    },
]


def _create(client: TestClient, headers: dict[str, str]):
    response = client.post("/api/v1/web-decks", headers=headers, json={
        "title": "梯度下降技能微课", "course_id": None, "slides": SLIDES,
    })
    assert response.status_code == 201, response.text
    return response.json()


def test_stable_slide_ids_version_reorder_restore_and_immutability():
    with TestClient(app) as client:
        owner = auth_headers(client, "fc1_008_version_owner")
        deck = _create(client, owner)
        deck_id = deck["id"]
        first_version_id = deck["current_version"]["id"]
        first_ids = [item["id"] for item in deck["current_version"]["slides"]]
        assert len(set(first_ids)) == 2
        assert deck["lock_version"] == 2

        reordered = [dict(deck["current_version"]["slides"][1]), dict(deck["current_version"]["slides"][0])]
        reordered[0]["title"] = "学习率与收敛边界"
        response = client.post(f"/api/v1/web-decks/{deck_id}/versions", headers=owner, json={
            "slides": reordered, "expected_lock_version": 2, "change_summary": "调整顺序并强化边界说明",
        })
        assert response.status_code == 200, response.text
        changed = response.json()
        assert changed["version_created"] is True
        assert changed["current_version"]["version_number"] == 2
        assert [item["id"] for item in changed["current_version"]["slides"]] == first_ids[::-1]
        assert changed["lock_version"] == 3

        replay = client.post(f"/api/v1/web-decks/{deck_id}/versions", headers=owner, json={
            "slides": changed["current_version"]["slides"], "expected_lock_version": 3,
            "change_summary": "重复保存",
        })
        assert replay.status_code == 200
        assert replay.json()["version_created"] is False
        assert replay.json()["lock_version"] == 3

        restored = client.post(f"/api/v1/web-decks/{deck_id}/restore", headers=owner, json={
            "version_id": first_version_id, "expected_lock_version": 3,
        })
        assert restored.status_code == 200
        assert restored.json()["current_version_id"] == first_version_id
        assert [item["id"] for item in restored.json()["current_version"]["slides"]] == first_ids
        assert restored.json()["lock_version"] == 4

        stale = client.put(f"/api/v1/web-decks/{deck_id}/playback", headers=owner, json={
            "current_slide": 1, "playback_status": "playing", "expected_lock_version": 3,
        })
        assert stale.status_code == 409

    assert OBJECT_SPECS["web_deck"].model.__tablename__ == "web_decks"
    assert OBJECT_SPECS["web_slide"].model.__tablename__ == "web_slides"
    assert OBJECT_SPECS["web_deck_version"].model.__tablename__ == "web_deck_versions"
    with SessionLocal() as db:
        version = db.get(DBWebDeckVersion, first_version_id)
        version.content_hash = "0" * 64
        with pytest.raises(DatabaseError, match="immutable"):
            db.commit()
        db.rollback()
        with pytest.raises(DatabaseError, match="immutable"):
            db.execute(text("DELETE FROM web_deck_versions WHERE id=:id"), {"id": first_version_id})
            db.commit()
        db.rollback()


def test_exports_are_real_and_pptx_contains_speaker_notes():
    with TestClient(app) as client:
        owner = auth_headers(client, "fc1_008_export_owner")
        deck = _create(client, owner)
        deck_id = deck["id"]
        capabilities = client.get("/api/v1/web-decks/capabilities", headers=owner)
        assert capabilities.status_code == 200
        assert capabilities.json()["exports"]["pptx"]["available"] is True

        html = client.get(f"/api/v1/web-decks/{deck_id}/export/html", headers=owner)
        assert html.status_code == 200
        assert html.content.startswith(b"<!doctype html>")
        assert html.text.count('class="slide"') == 2
        assert "先说明这是冻结合成教学示例" in html.text

        pptx = client.get(f"/api/v1/web-decks/{deck_id}/export/pptx", headers=owner)
        assert pptx.status_code == 200
        assert pptx.content[:2] == b"PK"
        with zipfile.ZipFile(io.BytesIO(pptx.content)) as archive:
            assert "ppt/slides/slide1.xml" in archive.namelist()
            assert "ppt/notesSlides/notesSlide1.xml" in archive.namelist()
        presentation = Presentation(io.BytesIO(pptx.content))
        assert len(presentation.slides) == 2
        assert "梯度下降技能微课" in presentation.slides[0].shapes[0].text
        assert "冻结合成教学示例" in presentation.slides[0].notes_slide.notes_text_frame.text

        pdf = client.get(f"/api/v1/web-decks/{deck_id}/export/pdf", headers=owner)
        if capabilities.json()["exports"]["pdf"]["available"]:
            assert pdf.status_code == 200, pdf.text
            assert pdf.content.startswith(b"%PDF")
            assert len(pdf.content) > 1000
        else:
            assert pdf.status_code == 503


def test_playback_refresh_owner_boundary_and_strict_inputs():
    with TestClient(app) as client:
        owner = auth_headers(client, "fc1_008_playback_owner")
        outsider = auth_headers(client, "fc1_008_playback_outsider")
        deck = _create(client, owner)
        deck_id = deck["id"]
        played = client.put(f"/api/v1/web-decks/{deck_id}/playback", headers=owner, json={
            "current_slide": 1, "playback_status": "playing", "expected_lock_version": 2,
        })
        assert played.status_code == 200, played.text
        assert played.json()["current_slide"] == 1
        assert played.json()["playback_status"] == "playing"
        assert played.json()["lock_version"] == 3
        restored = client.get(f"/api/v1/web-decks/{deck_id}", headers=owner)
        assert restored.status_code == 200
        assert restored.json()["current_slide"] == 1

        assert client.get(f"/api/v1/web-decks/{deck_id}", headers=outsider).status_code == 403
        assert client.get(f"/api/v1/web-decks/{deck_id}/export/pptx", headers=outsider).status_code == 403
        assert client.get("/api/v1/web-decks", headers=outsider).json()["count"] == 0

        injected = client.post("/api/v1/web-decks", headers=owner, json={
            "title": "额外字段注入", "slides": SLIDES, "rendered_html": "<script>alert(1)</script>",
        })
        assert injected.status_code == 422
        invalid = client.post("/api/v1/web-decks", headers=owner, json={
            "title": "空演示", "slides": [],
        })
        assert invalid.status_code == 422
