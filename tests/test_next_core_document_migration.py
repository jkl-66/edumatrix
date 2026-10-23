from __future__ import annotations

from io import BytesIO

import nbformat
import pytest
from fastapi.testclient import TestClient

import knowledge_api
from app.database import DBChapter, DBSourceDocument, DBSourceRef, SessionLocal
from app.main import app
from document_parser import DocumentParseError, parse_document
from tests.api_test_helpers import auth_headers


def _pdf_bytes() -> bytes:
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    output = canvas.Canvas(buffer)
    output.drawString(72, 720, "PDF provenance page one")
    output.showPage()
    output.drawString(72, 720, "PDF provenance page two")
    output.save()
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("DOCX migration", level=1)
    document.add_paragraph("Paragraph with canonical provenance")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "metric"
    table.cell(0, 1).text = "value"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "PPTX migration"
    slide.placeholders[1].text = "Slide-level canonical location"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _notebook_bytes() -> bytes:
    notebook = nbformat.v4.new_notebook()
    notebook.cells = [
        nbformat.v4.new_markdown_cell("# Notebook migration"),
        nbformat.v4.new_code_cell(
            "print('safe source')",
            execution_count=17,
            outputs=[nbformat.v4.new_output("stream", name="stdout", text="SECRET_OUTPUT\n")],
        ),
    ]
    return nbformat.writes(notebook).encode("utf-8")


@pytest.mark.parametrize(
    ("filename", "factory", "location_kind", "needle"),
    [
        ("source.pdf", _pdf_bytes, "page", "PDF provenance page one"),
        ("source.docx", _docx_bytes, "paragraph", "canonical provenance"),
        ("source.pptx", _pptx_bytes, "slide", "Slide-level canonical location"),
        ("source.ipynb", _notebook_bytes, "cell", "print('safe source')"),
    ],
)
def test_four_formats_share_versioned_dual_hash_location_contract(
    filename, factory, location_kind, needle
):
    parsed = parse_document(factory(), filename)
    assert parsed.schema_version == "edumatrix.document-parse.v1"
    assert parsed.parser_version.startswith("edumatrix-")
    assert len(parsed.source_sha256) == len(parsed.processed_sha256) == 64
    assert parsed.source_sha256 != parsed.processed_sha256
    assert needle in parsed.text
    assert any(location.kind == location_kind for location in parsed.locations)
    for location in parsed.locations:
        assert 0 <= location.char_start < location.char_end <= len(parsed.text)
        assert parsed.text[location.char_start:location.char_end].strip()


def test_notebook_outputs_and_execution_counts_are_not_migrated():
    parsed = parse_document(_notebook_bytes(), "training.ipynb")
    assert "SECRET_OUTPUT" not in parsed.text
    assert "execution_count" not in parsed.text
    assert parsed.metadata["stripped_output_count"] == 1
    assert parsed.metadata["outputs_included"] is False
    assert [location.cell_index for location in parsed.locations] == [0, 1]


@pytest.mark.parametrize("filename", ["broken.pdf", "broken.docx", "broken.pptx", "broken.ipynb"])
def test_corrupt_binary_formats_fail_closed(filename):
    with pytest.raises(DocumentParseError) as exc:
        parse_document(b"not-a-valid-document", filename)
    assert exc.value.code in {
        "document_limit_rejected",
        "pdf_parse_failed",
        "docx_parse_failed",
        "pptx_parse_failed",
        "notebook_parse_failed",
    }


def test_batch_upload_isolates_failure_and_persists_authorized_canonical_source(monkeypatch):
    monkeypatch.setattr(knowledge_api, "build_graph_after_upload", lambda *args, **kwargs: None)

    def unavailable_index(*args, **kwargs):
        raise RuntimeError("synthetic index outage")

    monkeypatch.setattr(knowledge_api.hybrid_rag, "ingest_user_documents", unavailable_index)

    async def no_guide(**kwargs):
        return None

    monkeypatch.setattr(knowledge_api, "_generate_doc_guide_for_document", no_guide)

    with TestClient(app) as client:
        owner_headers = auth_headers(client, "nc_core_05_owner")
        other_headers = auth_headers(client, "nc_core_05_other")
        response = client.post(
            "/api/knowledge/upload-batch?student_id=nc_core_05_owner",
            headers=owner_headers,
            files=[
                ("files", ("training.ipynb", _notebook_bytes(), "application/x-ipynb+json")),
                ("files", ("broken.pdf", b"broken-pdf", "application/pdf")),
            ],
        )
        assert response.status_code == 200
        payload = response.json()
        assert (payload["total"], payload["succeeded"], payload["failed"]) == (2, 1, 1)
        assert [item["status"] for item in payload["results"]] == ["ready", "failed"]
        document_id = payload["results"][0]["document"]["id"]
        source_document_id = payload["results"][0]["document"]["source_document_id"]
        assert payload["results"][0]["document"]["indexing_warnings"] == ["rag_index_deferred"]

        detail = client.get(
            f"/api/knowledge/{document_id}?student_id=nc_core_05_owner",
            headers=owner_headers,
        )
        assert detail.status_code == 200
        source = detail.json()
        assert source["source_document_id"] == source_document_id
        assert source["source_processing_status"] == "ready"
        assert source["source_parser_version"] == "edumatrix-notebook-to-markdown-v2"
        assert source["source_sha256"] != source["processed_sha256"]
        assert len(source["source_locations"]) == 2
        assert source["source_parser_metadata"]["stripped_output_count"] == 1
        source_ref_id = source["source_locations"][0]["source_ref_id"]

        allowed = client.get(f"/api/v1/source-refs/{source_ref_id}", headers=owner_headers)
        assert allowed.status_code == 200
        assert allowed.json()["location_kind"] == "cell"
        denied = client.get(f"/api/v1/source-refs/{source_ref_id}", headers=other_headers)
        assert denied.status_code == 403

        with SessionLocal() as session:
            canonical = session.get(DBSourceDocument, source_document_id)
            assert canonical is not None
            assert canonical.processing_status == "ready"
            assert session.query(DBSourceRef).filter_by(source_document_id=source_document_id).count() == 2
            assert session.query(DBChapter).filter_by(source_document_id=source_document_id).count() == 2
            assert session.query(DBSourceDocument).filter_by(filename="broken.pdf").count() == 0

        deleted = client.delete(
            f"/api/knowledge/{document_id}?student_id=nc_core_05_owner",
            headers=owner_headers,
        )
        assert deleted.status_code == 200
        with SessionLocal() as session:
            assert session.get(DBSourceDocument, source_document_id) is None
            assert session.query(DBSourceRef).filter_by(source_document_id=source_document_id).count() == 0
            assert session.query(DBChapter).filter_by(source_document_id=source_document_id).count() == 0
