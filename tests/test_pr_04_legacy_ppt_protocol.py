from __future__ import annotations

import hashlib
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

import document_parser
from app.database import DBChatAttachment, SessionLocal
from app.main import app
from document_parser import DocumentParseError, parse_document
from tests.test_f7_t04_chat_api import _conversation, _key, _upload
from tests.api_test_helpers import auth_headers


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Legacy conversion contract"
    slide.placeholders[1].text = "The original .ppt remains the source attachment."
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_legacy_ppt_conversion_keeps_original_hash_and_slide_locations(monkeypatch: pytest.MonkeyPatch) -> None:
    converted = _pptx_bytes()
    original = b"legacy-ppt-source"
    monkeypatch.setattr(
        document_parser,
        "_convert_legacy_ppt",
        lambda raw, filename: (converted, {"source_format": "ppt", "converted_format": "pptx", "conversion_mode": "test_converter"}),
    )

    parsed = parse_document(original, "lesson.ppt")

    assert parsed.file_format == "ppt"
    assert parsed.parser_version == "edumatrix-ppt-legacy-v1"
    assert parsed.source_sha256 == hashlib.sha256(original).hexdigest()
    assert parsed.metadata["source_format"] == "ppt"
    assert parsed.metadata["converted_format"] == "pptx"
    assert parsed.metadata["conversion_mode"] == "test_converter"
    assert parsed.metadata["slide_count"] == 1
    assert parsed.locations[0].slide_number == 1
    assert "Legacy conversion contract" in parsed.text


def test_legacy_ppt_converter_unavailable_is_explicit_and_retryable(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(document_parser.shutil, "which", lambda _name: None)

    with pytest.raises(DocumentParseError) as caught:
        parse_document(b"legacy-ppt-source", "lesson.ppt")

    assert caught.value.code == "legacy_ppt_converter_unavailable"
    assert caught.value.file_format == "ppt"
    assert "另存为 PPTX" in str(caught.value)


def test_chat_keeps_legacy_ppt_bytes_when_conversion_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    def fail_conversion(_raw: bytes, _filename: str):
        raise DocumentParseError(
            "legacy_ppt_conversion_failed",
            "旧版 PPT 转换暂时失败，原文件已保留，请重试或先另存为 PPTX。",
            file_format="ppt",
        )

    monkeypatch.setattr(document_parser, "_convert_legacy_ppt", fail_conversion)
    username = _key("pr04-chat")
    original = b"legacy-ppt-that-must-remain"
    with TestClient(app) as client:
        headers = auth_headers(client, username)
        conversation = _conversation(client, headers)
        attachment = _upload(
            client,
            headers,
            conversation["id"],
            filename="lesson.ppt",
            raw=original,
            mime_type="application/vnd.ms-powerpoint",
        )
        parsed = client.post(
            f"/api/v1/chat/attachments/{attachment['id']}/parse",
            headers=headers,
            json={"expected_lock_version": attachment["lockVersion"]},
        )
        assert parsed.status_code == 200, parsed.text
        payload = parsed.json()
        assert payload["status"] == "failed_retryable"
        assert payload["failure"]["code"] == "legacy_ppt_conversion_failed"
        assert payload["failure"]["retryable"] is True
        assert payload["sourceSha256"] == hashlib.sha256(original).hexdigest()

        blocked = client.post(
            f"/api/v1/chat/conversations/{conversation['id']}/messages",
            headers=headers,
            json={
                "content": "解释这个附件",
                "attachment_ids": [attachment["id"]],
                "context_refs": [],
                "idempotency_key": _key("message"),
            },
        )
        assert blocked.status_code == 409
        assert blocked.json()["detail"]["code"] == "chat_attachment_not_ready"

    with SessionLocal() as session:
        row = session.get(DBChatAttachment, attachment["id"])
        assert row is not None
        assert row.raw_bytes == original
        assert row.source_sha256 == hashlib.sha256(original).hexdigest()


def test_knowledge_upload_accepts_ppt_format_for_parser_protocol() -> None:
    from knowledge_api import SUPPORTED_EXTENSIONS

    assert SUPPORTED_EXTENSIONS[".ppt"] == "presentation"

